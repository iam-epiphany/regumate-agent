"""Banking Workbench audit, projection, validation, freeze and verify.

Evaluation-only.  Production code never imports this package.

The pipeline treats the private gold file as the single source of truth:
- ``audit`` verifies every case against the official corpus (spreadsheet
  coordinates recomputed from the original files with Decimal comparison,
  text evidence matched against the indexed corpus text, calculations
  recomputed with Decimal).
- ``project`` derives the public question projection (five fields only).
- ``validate-batch`` enforces batch-level quotas, public/private identity and
  zero private-field leakage.
- ``freeze`` writes lock.json with four SHA256 digests
  (questions / gold / scorer / corpus) and refuses to overwrite.
- ``verify`` recomputes the digests and compares.
"""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from difflib import SequenceMatcher
import json
import re
import sqlite3
from pathlib import Path
from typing import Any

import sys
from pathlib import Path as _Path

sys.path.insert(0, str(_Path(__file__).resolve().parents[3]))

try:
    from banking_workbench.workbench_common import (
    ANSWERABLE_QUOTA,
    DIFFICULTY_QUOTA,
    MIN_SOURCE_DOCUMENTS,
    PRIVATE_FIELDS,
    PUBLIC_FIELDS,
    QUESTION_TYPE_QUOTA,
    cell_value_numeric,
    corpus_manifest_hash,
    normalise,
    normalise_title,
    read_jsonl,
    sha256_file,
    write_jsonl,
    )
except ModuleNotFoundError:
    from scripts.evaluation.banking_workbench.workbench_common import (
    ANSWERABLE_QUOTA,
    DIFFICULTY_QUOTA,
    MIN_SOURCE_DOCUMENTS,
    PRIVATE_FIELDS,
    PUBLIC_FIELDS,
    QUESTION_TYPE_QUOTA,
    cell_value_numeric,
    corpus_manifest_hash,
    normalise,
    normalise_title,
    read_jsonl,
    sha256_file,
    write_jsonl,
)

PROJECT_ROOT = Path(__file__).resolve().parents[3]
DEFAULT_CORPUS = PROJECT_ROOT / "data" / "contest_dataset" / "dataset" / "nfra_page_attachments_500"
DEFAULT_DB = PROJECT_ROOT / "data" / "evaluation" / "final_runtime" / "app.db"
SCORER_PATH = PROJECT_ROOT / "scripts" / "evaluation" / "banking_workbench" / "workbench_scorer.py"


# ---------------------------------------------------------------------------
# corpus text access (from the indexed corpus; coordinates are recomputed
# from the original spreadsheet files independently)
# ---------------------------------------------------------------------------


def _corpus_text_by_document(db_path: Path) -> dict[str, list[str]]:
    """Map normalised filename -> list of normalised chunk texts."""

    mapping: dict[str, list[str]] = {}
    if not db_path.exists():
        return mapping
    conn = sqlite3.connect(str(db_path))
    try:
        rows = conn.execute(
            "SELECT c.source_file, c.text FROM document_chunks c "
            "JOIN documents d ON d.document_id = c.document_id "
            "WHERE d.status = 'indexed'"
        ).fetchall()
        for source_file, text in rows:
            key = normalise_title(source_file)
            mapping.setdefault(key, []).append(normalise(text))
    finally:
        conn.close()
    return mapping


def _document_filename(relative_path: str) -> str:
    return normalise_title(Path(str(relative_path or "")).name)


def _spreadsheet_cell_value(corpus_root: Path, relative_path: str, sheet: str, row: int, column: int) -> Any:
    """Read one cell from the original spreadsheet file.

    ``row``/``column`` are 1-based (matching the SpreadsheetCell index):
    openpyxl consumes 1-based coordinates directly; xlrd needs 0-based.
    """

    path = corpus_root / str(relative_path)
    if not path.exists():
        raise FileNotFoundError(f"corpus file missing: {relative_path}")
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            worksheet = workbook[sheet] if sheet in workbook.sheetnames else workbook.worksheets[0]
            return worksheet.cell(row=row, column=column).value
        finally:
            workbook.close()
    if suffix == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(str(path))
        try:
            worksheet = workbook.sheet_by_name(sheet) if sheet in workbook.sheet_names() else workbook.sheet_by_index(0)
            return worksheet.cell_value(row - 1, column - 1)
        finally:
            workbook.release_resources()
    raise ValueError(f"unsupported spreadsheet type: {suffix}")


def _audit_sources(
    gold: dict[str, Any],
    corpus_root: Path,
    corpus_text: dict[str, list[str]],
    errors: list[str],
) -> None:
    sources = gold.get("required_sources") or []
    if not sources and gold.get("answerable"):
        errors.append("answerable case has no required_sources")
        return
    for index, source in enumerate(sources):
        relative_path = str(source.get("relative_path") or "")
        if not relative_path:
            errors.append(f"source[{index}] missing relative_path")
            continue
        if not (corpus_root / relative_path).exists():
            errors.append(f"source[{index}] not in corpus: {relative_path}")
            continue
        if "coordinate" in source or "expected_value" in source:
            sheet = str(source.get("sheet") or "")
            try:
                row = int(source["row"])
                column = int(source["column"])
                raw = _spreadsheet_cell_value(corpus_root, relative_path, sheet, row, column)
            except Exception as exc:  # noqa: BLE001
                errors.append(f"source[{index}] spreadsheet read failed: {exc}")
                continue
            expected = cell_value_numeric(source.get("expected_value"))
            actual = cell_value_numeric(raw)
            if expected is None:
                errors.append(f"source[{index}] expected_value not numeric")
            elif actual is None:
                errors.append(f"source[{index}] file cell not numeric: {raw!r}")
            elif abs(actual - expected) / max(abs(expected), Decimal("1")) > Decimal("1e-6"):
                errors.append(
                    f"source[{index}] coordinate mismatch: expected {expected}, file has {raw!r}"
                )
        else:
            evidence = normalise(source.get("evidence_text") or "")
            if not evidence:
                errors.append(f"source[{index}] text source missing evidence_text")
                continue
            filename = _document_filename(relative_path)
            chunks = corpus_text.get(filename, [])
            if not any(evidence in chunk for chunk in chunks):
                errors.append(f"source[{index}] evidence_text not found in corpus text of {relative_path}")


def _audit_calculation(gold: dict[str, Any], errors: list[str]) -> None:
    calculation = gold.get("calculation")
    if not calculation or not gold.get("answerable"):
        return
    operator = str(calculation.get("operator") or "")
    try:
        left = Decimal(str(calculation["left"]))
        right = Decimal(str(calculation["right"]))
        result = Decimal(str(calculation.get("result") or ""))
    except (KeyError, InvalidOperation):
        errors.append("calculation missing left/right/result numeric fields")
        return
    computed = {
        "difference": left - right,
        "sum": left + right,
        "ratio": left / right if right != 0 else None,
    }.get(operator)
    if computed is None:
        errors.append(f"calculation unsupported operator: {operator}")
    elif computed != result:
        errors.append(f"calculation mismatch: {left} {operator} {right} = {computed}, gold says {result}")


# ---------------------------------------------------------------------------
# per-case audit
# ---------------------------------------------------------------------------


def _question_text_errors(question: str) -> list[str]:
    errors = []
    if "�" in question:
        errors.append("replacement character")
    if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", question):
        errors.append("control characters")
    if len(re.findall(r"\?", question)) >= 3:
        errors.append("question-mark density suggests code-page damage")
    if not re.search(r"[一-鿿]", question):
        errors.append("no Chinese content")
    return errors


def audit_gold(gold: dict[str, Any], corpus_root: Path, corpus_text: dict[str, list[str]]) -> list[str]:
    errors: list[str] = []
    for field in ("id", "question", "question_type", "difficulty", "business_relevance"):
        if not str(gold.get(field) or "").strip():
            errors.append(f"missing field: {field}")
    question = str(gold.get("question") or "")
    errors.extend(f"question: {item}" for item in _question_text_errors(question))
    answerable = bool(gold.get("answerable"))
    if answerable:
        if not str(gold.get("canonical_answer") or "").strip():
            errors.append("answerable case missing canonical_answer")
        conclusions = gold.get("required_conclusions") or []
        if not conclusions:
            errors.append("answerable case missing required_conclusions")
        for index, conclusion in enumerate(conclusions):
            if not str(conclusion or "").strip():
                errors.append(f"required_conclusions[{index}] empty")
        _audit_sources(gold, corpus_root, corpus_text, errors)
        _audit_calculation(gold, errors)
    else:
        if not str(gold.get("expected_refusal_code") or "").strip():
            errors.append("refusal case missing expected_refusal_code")
        if not str(gold.get("refusal_rationale") or "").strip():
            errors.append("refusal case missing refusal_rationale")
    return errors


def _duplicate_reason(left: dict[str, Any], right: dict[str, Any]) -> str | None:
    """Five-way deduplication: exact string, token overlap, sequence
    similarity, same evidence+conclusion, same reasoning structure."""

    left_q = normalise(left.get("question") or "")
    right_q = normalise(right.get("question") or "")
    if left_q == right_q:
        return "identical_question"
    left_tokens = set(re.findall(r"[一-鿿A-Za-z0-9]{2,}", left_q))
    right_tokens = set(re.findall(r"[一-鿿A-Za-z0-9]{2,}", right_q))
    if left_tokens and right_tokens:
        overlap = len(left_tokens & right_tokens) / min(len(left_tokens), len(right_tokens))
        if overlap >= 0.9 and min(len(left_tokens), len(right_tokens)) >= 5:
            return "token_overlap"
    if SequenceMatcher(None, left_q, right_q).ratio() >= 0.98:
        return "sequence_similar"
    left_evidence = normalise(
        "".join(str(s.get("evidence_text") or "") for s in (left.get("required_sources") or []))
    )
    right_evidence = normalise(
        "".join(str(s.get("evidence_text") or "") for s in (right.get("required_sources") or []))
    )
    if left_evidence and left_evidence == right_evidence:
        left_conclusions = normalise("".join(left.get("required_conclusions") or []))
        right_conclusions = normalise("".join(right.get("required_conclusions") or []))
        if left_conclusions and left_conclusions == right_conclusions:
            return "same_evidence_same_conclusion"
    left_frames = normalise(str(left.get("required_evidence_aspects") or ""))
    right_frames = normalise(str(right.get("required_evidence_aspects") or ""))
    if (
        left_frames
        and left_frames == right_frames
        and left.get("question_type") == right.get("question_type")
        and normalise(str(left.get("calculation") or "")) == normalise(str(right.get("calculation") or ""))
    ):
        return "same_reasoning_structure"
    return None


def audit(gold_path: Path, corpus_root: Path, db_path: Path, output: Path) -> int:
    rows = read_jsonl(gold_path)
    corpus_text = _corpus_text_by_document(db_path)
    per_case: list[dict[str, Any]] = []
    duplicate_pairs: list[str] = []
    for index, gold in enumerate(rows):
        errors = audit_gold(gold, corpus_root, corpus_text)
        per_case.append({"id": gold.get("id"), "passed": not errors, "errors": errors})
        for other in rows[:index]:
            reason = _duplicate_reason(other, gold)
            if reason:
                duplicate_pairs.append(f"{other.get('id')} ~ {gold.get('id')}: {reason}")
    report = {
        "schema_version": 1,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "case_count": len(rows),
        "passed_count": sum(1 for item in per_case if item["passed"]),
        "duplicate_pairs": duplicate_pairs,
        "cases": per_case,
        "passed": all(item["passed"] for item in per_case) and not duplicate_pairs,
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: v for k, v in report.items() if k != "cases"}, ensure_ascii=False, indent=2))
    return 0 if report["passed"] else 1


def project(gold_path: Path, output: Path) -> int:
    rows = read_jsonl(gold_path)
    public = [{field: row.get(field) for field in PUBLIC_FIELDS} for row in rows]
    write_jsonl(output, public)
    print(f"projected {len(public)} public questions -> {output}")
    return 0


def validate_batch(gold_path: Path, public_path: Path, report_path: Path) -> int:
    gold = read_jsonl(gold_path)
    public = read_jsonl(public_path)
    errors: list[str] = []
    if len(gold) != 100:
        errors.append(f"expected 100 gold cases, found {len(gold)}")
    if len(public) != len(gold):
        errors.append("public/gold length mismatch")
    gold_ids = [str(item.get("id") or "") for item in gold]
    if len(gold_ids) != len(set(gold_ids)):
        errors.append("duplicate gold ids")
    if gold_ids and public and [str(item.get("id") or "") for item in public] != gold_ids:
        errors.append("public/gold id order mismatch")
    for index, item in enumerate(public):
        for field in PUBLIC_FIELDS:
            if field == "answerable":
                if not isinstance(item.get(field), bool):
                    errors.append(f"public[{index}] missing boolean answerable")
            elif str(item.get(field) or "").strip() == "":
                errors.append(f"public[{index}] missing {field}")
        serialized = json.dumps(item, ensure_ascii=False)
        for field in PRIVATE_FIELDS:
            if re.search(rf'"{field}"\s*:', serialized):
                errors.append(f"public[{index}] leaks private field {field}")
        if gold[index].get("question") != item.get("question"):
            errors.append(f"public[{index}] question differs from gold")
    by_type: dict[str, int] = {}
    by_difficulty: dict[str, int] = {}
    answerable_count = 0
    source_documents: set[str] = set()
    for item in gold:
        by_type[str(item.get("question_type"))] = by_type.get(str(item.get("question_type")), 0) + 1
        by_difficulty[str(item.get("difficulty"))] = by_difficulty.get(str(item.get("difficulty")), 0) + 1
        if item.get("answerable"):
            answerable_count += 1
            for source in item.get("required_sources") or []:
                source_documents.add(normalise_title(str(source.get("relative_path") or "")))
        else:
            by_type["refusal"] = by_type.get("refusal", 0)
    for kind, quota in QUESTION_TYPE_QUOTA.items():
        actual = by_type.get(kind, 0)
        if actual != quota:
            errors.append(f"question_type {kind}: expected {quota}, found {actual}")
    for level, quota in DIFFICULTY_QUOTA.items():
        actual = by_difficulty.get(level, 0)
        if actual != quota:
            errors.append(f"difficulty {level}: expected {quota}, found {actual}")
    if answerable_count != ANSWERABLE_QUOTA["answerable"]:
        errors.append(f"answerable: expected {ANSWERABLE_QUOTA['answerable']}, found {answerable_count}")
    if len(source_documents) < MIN_SOURCE_DOCUMENTS:
        errors.append(f"source documents: expected >= {MIN_SOURCE_DOCUMENTS}, found {len(source_documents)}")
    report = {
        "schema_version": 1,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "passed": not errors,
        "errors": errors,
        "by_question_type": by_type,
        "by_difficulty": by_difficulty,
        "source_document_count": len(source_documents),
    }
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["passed"] else 1


def _freeze_digests(public_root: Path, private_root: Path, corpus_root: Path) -> dict[str, str]:
    return {
        "questions_sha256": sha256_file(public_root / "public" / "questions.jsonl"),
        "gold_sha256": sha256_file(private_root / "gold.jsonl"),
        "scorer_sha256": sha256_file(SCORER_PATH),
        "corpus_sha256": corpus_manifest_hash(corpus_root),
    }


def freeze(public_root: Path, private_root: Path, corpus_root: Path) -> int:
    lock_path = public_root / "lock.json"
    if lock_path.exists():
        print(f"refusing to overwrite existing lock: {lock_path}")
        return 1
    digests = _freeze_digests(public_root, private_root, corpus_root)
    lock = {
        "schema_version": 1,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "first_run_immutable": True,
        **digests,
    }
    lock_path.parent.mkdir(parents=True, exist_ok=True)
    lock_path.write_text(json.dumps(lock, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(lock, ensure_ascii=False, indent=2))
    return 0


def verify(public_root: Path, private_root: Path, corpus_root: Path, report_path: Path | None = None) -> int:
    lock_path = public_root / "lock.json"
    if not lock_path.exists():
        print(f"lock not found: {lock_path}")
        return 2
    expected = json.loads(lock_path.read_text(encoding="utf-8"))
    actual = _freeze_digests(public_root, private_root, corpus_root)
    checks = {key: actual.get(key) == expected.get(key) for key in expected if key.endswith("_sha256")}
    passed = all(checks.values())
    report = {
        "schema_version": 1,
        "passed": passed,
        "checks": checks,
        "lock_path": str(lock_path),
    }
    if report_path:
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if passed else 2


def main() -> int:
    parser = argparse.ArgumentParser(description="Banking Workbench audit/project/validate/freeze/verify")
    sub = parser.add_subparsers(dest="command", required=True)
    a = sub.add_parser("audit")
    a.add_argument("--gold", type=Path, required=True)
    a.add_argument("--corpus", type=Path, default=DEFAULT_CORPUS)
    a.add_argument("--db", type=Path, default=DEFAULT_DB)
    a.add_argument("--output", type=Path, required=True)
    a.set_defaults(func=audit)
    p = sub.add_parser("project")
    p.add_argument("--gold", type=Path, required=True)
    p.add_argument("--output", type=Path, required=True)
    p.set_defaults(func=project)
    v = sub.add_parser("validate-batch")
    v.add_argument("--gold", type=Path, required=True)
    v.add_argument("--public-questions", type=Path, required=True)
    v.add_argument("--report", type=Path, required=True)
    v.set_defaults(func=validate_batch)
    f = sub.add_parser("freeze")
    f.add_argument("--public-root", type=Path, required=True)
    f.add_argument("--private-root", type=Path, required=True)
    f.add_argument("--corpus", type=Path, default=DEFAULT_CORPUS)
    f.set_defaults(func=freeze)
    vf = sub.add_parser("verify")
    vf.add_argument("--public-root", type=Path, required=True)
    vf.add_argument("--private-root", type=Path, required=True)
    vf.add_argument("--corpus", type=Path, default=DEFAULT_CORPUS)
    vf.add_argument("--report", type=Path)
    vf.set_defaults(func=verify)
    args = parser.parse_args()
    if args.command == "audit":
        return args.func(args.gold, args.corpus, args.db, args.output)
    if args.command == "project":
        return args.func(args.gold, args.output)
    if args.command == "validate-batch":
        return args.func(args.gold, args.public_questions, args.report)
    if args.command == "freeze":
        return args.func(args.public_root, args.private_root, args.corpus)
    if args.command == "verify":
        return args.func(args.public_root, args.private_root, args.corpus, args.report)
    return 2


if __name__ == "__main__":
    raise SystemExit(main())
