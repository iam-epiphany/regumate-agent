from __future__ import annotations

import argparse
from datetime import datetime, timezone
import json
import math
from pathlib import Path
import re
import sqlite3
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_QA = ROOT / "data" / "contest dataset" / "QA数据.xlsx"
DEFAULT_DATABASE = ROOT / "data" / "evaluation" / "final_runtime" / "app.db"
DEFAULT_OUTPUT = ROOT / "data" / "evaluation" / "final" / "official_excel_cells.json"


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Validate all official Excel evidence cells against the normalized SQLite index."
    )
    parser.add_argument("--qa", type=Path, default=DEFAULT_QA)
    parser.add_argument("--database", type=Path, default=DEFAULT_DATABASE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    expected = read_expected_cells(args.qa)
    results = validate_cells(args.database, expected)
    located = sum(item["located"] for item in results)
    values_matched = sum(item["value_matched"] for item in results)
    artifact = {
        "run": {
            "created_at": datetime.now(timezone.utc).isoformat(),
            "qa": str(args.qa),
            "database": str(args.database),
        },
        "summary": {
            "expected": len(results),
            "located": located,
            "location_rate": located / len(results) if results else 0.0,
            "values_matched": values_matched,
            "value_match_rate": values_matched / len(results) if results else 0.0,
        },
        "failures": [item for item in results if not item["located"] or not item["value_matched"]],
        "results": results,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(artifact, ensure_ascii=False, indent=2), encoding="utf-8")
    args.output.with_suffix(".md").write_text(render_report(artifact), encoding="utf-8")
    print(json.dumps(artifact["summary"], ensure_ascii=False, indent=2))
    return 0 if located == len(results) and values_matched == len(results) else 1


def read_expected_cells(path: Path) -> list[dict[str, str]]:
    worksheet = load_workbook(path, read_only=True, data_only=True).active
    rows = worksheet.iter_rows(values_only=True)
    headers = [str(value or "") for value in next(rows)]
    expected: list[dict[str, str]] = []
    for row in rows:
        record = dict(zip(headers, row, strict=False))
        if str(record.get("source_type")) != "excel":
            continue
        evidence = str(record.get("evidence") or "")
        filename = Path(re.split(r"[；;]", evidence, maxsplit=1)[0]).name
        sheet_match = re.search(r"工作表[：:]([^；;]+)", evidence)
        sheet_name = sheet_match.group(1).strip() if sheet_match else ""
        pair_by_coordinate = {
            coordinate: (label.strip(), value.strip())
            for label, value, coordinate in re.findall(
                r"([^；;=]+)=([^；;()]+)\(([A-Z]+\d+)\)", evidence
            )
        }
        lookup = re.search(
            r"单元格[：:]([A-Z]+\d+).*?原始值[：:]([^；;。]+)", evidence
        )
        if lookup:
            pair_by_coordinate.setdefault(lookup.group(1), ("", lookup.group(2).strip()))
        for coordinate in re.findall(r"\b[A-Z]+\d+\b", evidence):
            label, value = pair_by_coordinate.get(coordinate, ("", ""))
            expected.append(
                {
                    "question_id": str(record.get("id")),
                    "filename": filename,
                    "sheet_name": sheet_name,
                    "coordinate": coordinate,
                    "label": label,
                    "expected_value": value,
                }
            )
    return expected


def validate_cells(database: Path, expected: list[dict[str, str]]) -> list[dict[str, Any]]:
    connection = sqlite3.connect(database)
    connection.row_factory = sqlite3.Row
    try:
        results: list[dict[str, Any]] = []
        for item in expected:
            rows = connection.execute(
                """
                SELECT sc.value, sc.numeric_value, sc.row_label, sc.column_label
                FROM spreadsheet_cells AS sc
                JOIN documents AS d ON d.document_id = sc.document_id
                WHERE d.filename = ? AND TRIM(sc.sheet_name) = TRIM(?) AND sc.coordinate = ?
                """,
                (item["filename"], item["sheet_name"], item["coordinate"]),
            ).fetchall()
            matched = next(
                (row for row in rows if values_equal(item["expected_value"], row["value"], row["numeric_value"])),
                None,
            )
            results.append(
                {
                    **item,
                    "located": bool(rows),
                    "candidate_count": len(rows),
                    "value_matched": matched is not None,
                    "actual_values": [str(row["value"]) for row in rows],
                }
            )
        return results
    finally:
        connection.close()


def values_equal(expected: str, actual: Any, numeric_actual: Any) -> bool:
    normalized_expected = str(expected).replace(",", "").strip()
    normalized_actual = str(actual or "").replace(",", "").strip()
    if normalized_expected == normalized_actual:
        return True
    try:
        expected_number = float(normalized_expected.rstrip("%"))
        actual_number = float(numeric_actual if numeric_actual is not None else normalized_actual.rstrip("%"))
    except (TypeError, ValueError):
        return False
    decimal_match = re.search(r"\.(\d+)", normalized_expected.rstrip("%"))
    decimal_places = len(decimal_match.group(1)) if decimal_match else 0
    tolerance = 0.5 * (10 ** -decimal_places) + 1e-9
    return math.isclose(expected_number, actual_number, rel_tol=0.0, abs_tol=tolerance)


def render_report(artifact: dict[str, Any]) -> str:
    summary = artifact["summary"]
    lines = [
        "# 官方 Excel 标准单元格审计",
        "",
        f"- 标准引用：{summary['expected']}",
        f"- 成功定位：{summary['located']}（{summary['location_rate']:.2%}）",
        f"- 值一致：{summary['values_matched']}（{summary['value_match_rate']:.2%}）",
        "",
        "## 失败明细",
        "",
    ]
    failures = artifact["failures"]
    if failures:
        lines.extend(
            f"- {item['question_id']} {item['filename']} / {item['sheet_name']} / "
            f"{item['coordinate']}：expected={item['expected_value']} actual={item['actual_values']}"
            for item in failures
        )
    else:
        lines.append("- 无")
    return "\n".join(lines) + "\n"


if __name__ == "__main__":
    raise SystemExit(main())
