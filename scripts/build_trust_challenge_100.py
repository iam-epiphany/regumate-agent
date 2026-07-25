from __future__ import annotations

import argparse
import hashlib
import json
import re
import sqlite3
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))
DEFAULT_QA_RESULT = PROJECT_ROOT / "data/evaluation/final/gpu_full_with_snapshot_cache/contest_qa_all_results.json"
DEFAULT_OOD_RESULT = PROJECT_ROOT / "data/evaluation/final/gpu_ood_with_snapshot_cache/contest_qa_ood_results.json"
DEFAULT_DB = PROJECT_ROOT / "data/evaluation/final_runtime/app.db"
DEFAULT_QA_WORKBOOK = PROJECT_ROOT / "data/contest_dataset/QA数据.xlsx"
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "data/evaluation/trust_challenge_100"

SCENARIO_SPECS = {
    "制度事实与定义": {"medium": 3, "hard": 7, "dev_medium": 1, "dev_hard": 3},
    "条款语义与监管约束": {"medium": 3, "hard": 9, "dev_medium": 1, "dev_hard": 4},
    "表格精确取数": {"medium": 8, "hard": 7, "dev_medium": 3, "dev_hard": 3},
    "表格比较与计算": {"medium": 5, "hard": 10, "dev_medium": 2, "dev_hard": 4},
    "跨制度文件推理": {"medium": 2, "hard": 8, "dev_medium": 1, "dev_hard": 3},
    "制度与报表联合判断": {"medium": 2, "hard": 10, "dev_medium": 1, "dev_hard": 4},
    "Word/PDF 公式": {"medium": 2, "hard": 6, "dev_medium": 1, "dev_hard": 2},
    "时效、版本与冲突": {"medium": 1, "hard": 7, "dev_medium": 0, "dev_hard": 3},
    "可信拒答与歧义": {"medium": 4, "hard": 6, "dev_medium": 2, "dev_hard": 2},
}


@dataclass(frozen=True)
class Evidence:
    document_id: str
    filename: str
    file_sha256: str
    chunk_id: str
    chunk_sha256: str
    excerpt: str
    evidence_text_sha256: str
    page_number: int | None
    section_title: str | None
    section_number: str | None
    sheet_name: str | None
    cells: list[str]
    anchor_type: str = "sqlite_chunk"

    def as_dict(self) -> dict[str, Any]:
        return {
            "document_id": self.document_id,
            "filename": self.filename,
            "file_sha256": self.file_sha256,
            "chunk_id": self.chunk_id,
            "chunk_sha256": self.chunk_sha256,
            "page_number": self.page_number,
            "section_title": self.section_title,
            "section_number": self.section_number,
            "sheet_name": self.sheet_name,
            "cells": self.cells,
            "evidence_text": self.excerpt,
            "evidence_text_sha256": self.evidence_text_sha256,
            "anchor_type": self.anchor_type,
        }


def main() -> int:
    parser = argparse.ArgumentParser(description="Build the official-document ReguMate 100-case trust challenge.")
    parser.add_argument("--qa-result", type=Path, default=DEFAULT_QA_RESULT)
    parser.add_argument("--ood-result", type=Path, default=DEFAULT_OOD_RESULT)
    parser.add_argument("--db", type=Path, default=DEFAULT_DB)
    parser.add_argument("--qa-workbook", type=Path, default=DEFAULT_QA_WORKBOOK)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    args = parser.parse_args()

    qa_rows = _load_results(args.qa_result)
    _attach_workbook_options(qa_rows, args.qa_workbook)
    ood_rows = _load_results(args.ood_result)
    connection = sqlite3.connect(args.db)
    connection.row_factory = sqlite3.Row
    corpus = Corpus(connection)

    cases = ChallengeBuilder(corpus, qa_rows, ood_rows).build()
    summary = validate_cases(cases, corpus)
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    questions_path = output_dir / "questions.jsonl"
    gold_path = output_dir / "gold.jsonl"
    audit_path = output_dir / "build_audit.json"
    _write_jsonl(questions_path, [case["question_record"] for case in cases])
    _write_jsonl(gold_path, [case["gold_record"] for case in cases])
    audit = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "status": "codex_verified",
        "verification_method": [
            "generation_check",
            "evidence_reverification_without_answer_generation",
            "programmatic_consistency_validation",
        ],
        "expert_reviewed": False,
        "limitations": [
            "未经过银行监管专家人工复核。",
            "60题为封存回归集，不是第三方独立盲测。",
            "题目由已有官方500文档及其可定位证据构造，不使用虚构制度或报表。",
        ],
        "source_artifacts": {
            "qa_result": _relative(args.qa_result),
            "qa_result_sha256": _sha256_file(args.qa_result),
            "ood_result": _relative(args.ood_result),
            "ood_result_sha256": _sha256_file(args.ood_result),
            "sqlite": _relative(args.db),
        },
        "summary": summary,
        "output_hashes": {
            "questions_sha256": _sha256_file(questions_path),
            "gold_sha256": _sha256_file(gold_path),
        },
    }
    audit_path.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f"wrote {_relative(questions_path)}")
    print(f"wrote {_relative(gold_path)}")
    print(f"wrote {_relative(audit_path)}")
    return 0


class Corpus:
    def __init__(self, connection: sqlite3.Connection) -> None:
        self.connection = connection
        self.documents = {
            row["document_id"]: dict(row)
            for row in connection.execute(
                "SELECT document_id, filename, file_sha256, title, version_status, identity_review_status FROM documents"
            )
        }
        self.documents_by_filename = {row["filename"]: row for row in self.documents.values()}
        self.chunks = {
            row["chunk_id"]: dict(row)
            for row in connection.execute(
                "SELECT chunk_id, document_id, text, chunk_metadata, page_number, section_title FROM document_chunks"
            )
        }

    def evidence(self, citation: dict[str, Any]) -> Evidence:
        chunk_id = str(citation.get("chunk_id") or "")
        chunk = self.chunks.get(chunk_id)
        if not chunk:
            raise ValueError(f"Citation chunk is absent from SQLite: {chunk_id}")
        document = self.documents.get(str(chunk["document_id"]))
        if not document:
            raise ValueError(f"Citation document is absent from SQLite: {chunk['document_id']}")
        if document["filename"] != citation.get("filename"):
            raise ValueError(f"Citation filename mismatch for {chunk_id}")
        metadata = citation.get("metadata") or {}
        cells = [
            str(cell.get("coordinate"))
            for cell in metadata.get("cells") or []
            if cell.get("coordinate")
        ]
        excerpt = str(citation.get("excerpt") or chunk["text"] or "").strip()
        return Evidence(
            document_id=str(document["document_id"]),
            filename=str(document["filename"]),
            file_sha256=str(document["file_sha256"]),
            chunk_id=chunk_id,
            chunk_sha256=_sha256_text(str(chunk["text"] or "")),
            excerpt=excerpt,
            evidence_text_sha256=_sha256_text(excerpt),
            page_number=citation.get("page_number") or chunk.get("page_number"),
            section_title=citation.get("section_title") or chunk.get("section_title"),
            section_number=citation.get("section_number"),
            sheet_name=metadata.get("sheet_name"),
            cells=cells,
        )

    def by_prefix(self, prefix: str) -> dict[str, Any]:
        matches = [row for row in self.documents.values() if str(row["filename"]).startswith(prefix + "_")]
        if len(matches) != 1:
            raise ValueError(f"Expected one official document with prefix {prefix}, found {len(matches)}")
        return matches[0]

    def formula_evidence(self, prefix: str, needles: Iterable[str]) -> Evidence:
        document = self.by_prefix(prefix)
        candidates = [chunk for chunk in self.chunks.values() if chunk["document_id"] == document["document_id"]]
        normalized_needles = [_normalize_formula(value) for value in needles]
        for chunk in candidates:
            # A chunk may inherit document-level formula metadata that was extracted
            # elsewhere in the file.  Only anchor gold evidence when the formula is
            # present in the indexed chunk text itself.
            haystack = _normalize_formula(str(chunk["text"] or ""))
            if any(needle and needle in haystack for needle in normalized_needles):
                metadata = _safe_json(chunk["chunk_metadata"])
                excerpt = str(chunk["text"] or "").strip()
                return Evidence(
                    document_id=str(document["document_id"]),
                    filename=str(document["filename"]),
                    file_sha256=str(document["file_sha256"]),
                    chunk_id=str(chunk["chunk_id"]),
                    chunk_sha256=_sha256_text(str(chunk["text"] or "")),
                    excerpt=excerpt,
                    evidence_text_sha256=_sha256_text(excerpt),
                    page_number=chunk.get("page_number"),
                    section_title=chunk.get("section_title"),
                    section_number=metadata.get("section_number"),
                    sheet_name=None,
                    cells=[],
                )
        # The frozen 500-document index predates the formula parser upgrade. Reparse the
        # byte-identical official staging file and keep an explicit parser-block anchor;
        # this makes the current indexing gap visible instead of fabricating a chunk hit.
        from backend.app.services.document_parser import parse_document

        staging_matches = list((PROJECT_ROOT / "data/contest_staging").glob(prefix + "_*"))
        if len(staging_matches) != 1:
            raise ValueError(f"Expected one staging source for formula prefix {prefix}, found {len(staging_matches)}")
        staging_path = staging_matches[0]
        if _sha256_file(staging_path) != document["file_sha256"]:
            raise ValueError(f"Formula staging source hash differs from official indexed source: {staging_path.name}")
        parsed = parse_document(staging_path)
        for index, block in enumerate(parsed.blocks):
            formulas = block.metadata.get("formulas") or []
            haystack = _normalize_formula(block.text + json.dumps(formulas, ensure_ascii=False))
            if any(needle and needle in haystack for needle in normalized_needles):
                excerpt = str(block.text or "").strip()
                synthetic_id = f"{document['document_id']}-FORMULA-BLOCK-{index:04d}"
                return Evidence(
                    document_id=str(document["document_id"]),
                    filename=str(document["filename"]),
                    file_sha256=str(document["file_sha256"]),
                    chunk_id=synthetic_id,
                    chunk_sha256=_sha256_text(excerpt),
                    excerpt=excerpt,
                    evidence_text_sha256=_sha256_text(excerpt),
                    page_number=block.page_number,
                    section_title=block.section_title,
                    section_number=block.metadata.get("section_number"),
                    sheet_name=None,
                    cells=[],
                    anchor_type="source_parser_block",
                )
        raise ValueError(f"Formula evidence not found for {document['filename']}: {list(needles)}")


class ChallengeBuilder:
    def __init__(self, corpus: Corpus, qa_rows: list[dict[str, Any]], ood_rows: list[dict[str, Any]]) -> None:
        self.corpus = corpus
        self.qa_rows = qa_rows
        self.ood_rows = ood_rows
        self.case_number = 0
        self.primary_counts: Counter[str] = Counter()
        self.cases: list[dict[str, Any]] = []
        self.text_rows = [row for row in qa_rows if row.get("source_type") in {"word", "pdf"}]
        self.table_rows = [row for row in qa_rows if row.get("source_type") == "excel"]

    def build(self) -> list[dict[str, Any]]:
        table_take, table_compute = self._select_table_rows()
        text_seed = self._select_text_rows(52)

        self._add_single_rows("制度事实与定义", text_seed[:10], "text_fact")
        self._add_single_rows("条款语义与监管约束", text_seed[10:22], "multi_assertion")
        self._add_single_rows("表格精确取数", table_take, "numeric")
        self._add_single_rows("表格比较与计算", table_compute, "numeric")
        self._add_cross_file(text_seed[22:42])
        self._add_joint(text_seed[40:52], self._match_joint_tables(text_seed[40:52], table_take + table_compute))
        self._add_formula_cases()
        self._add_version_cases()
        self._add_refusal_cases()
        self._assign_difficulty_and_split()
        return self.cases

    def _select_table_rows(self) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
        selected_files: set[str] = set()
        take: list[dict[str, Any]] = []
        compute: list[dict[str, Any]] = []
        for row in self.table_rows:
            filename = _primary_filename(row)
            if row.get("qa_type") == "表格取数" and filename and filename not in selected_files:
                take.append(row)
                selected_files.add(filename)
                if len(take) == 15:
                    break
        for row in self.table_rows:
            filename = _primary_filename(row)
            if row.get("qa_type") in {"表格比较", "表格计算"} and filename and filename not in selected_files:
                compute.append(row)
                selected_files.add(filename)
                if len(compute) == 15:
                    break
        if len(take) != 15 or len(compute) != 15:
            raise ValueError(f"Unable to select 30 distinct spreadsheet sources: take={len(take)}, compute={len(compute)}")
        return take, compute

    def _select_text_rows(self, count: int) -> list[dict[str, Any]]:
        by_file: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in self.text_rows:
            filename = _primary_filename(row)
            quality, evidence_length = _text_row_evidence_quality(row)
            if (
                filename
                and row.get("answer_correct") is True
                and row.get("source_hit") is True
                and quality >= 0.4
                and evidence_length >= 40
            ):
                by_file[filename].append(row)
        for rows in by_file.values():
            rows.sort(key=lambda row: (-_text_row_evidence_quality(row)[0], str(row.get("id") or "")))
        ordered_files = sorted(by_file)
        selected: list[dict[str, Any]] = []
        round_index = 0
        while len(selected) < count:
            added = False
            for filename in ordered_files:
                rows = by_file[filename]
                if round_index < len(rows):
                    selected.append(rows[round_index])
                    added = True
                    if len(selected) == count:
                        break
            if not added:
                break
            round_index += 1
        if len(selected) != count:
            raise ValueError(f"Unable to select {count} official text cases")
        return selected

    def _add_single_rows(self, scenario: str, rows: list[dict[str, Any]], scoring_type: str) -> None:
        for row in rows:
            evidence = [self.corpus.evidence(item) for item in row.get("citations") or []]
            canonical = _canonical_answer(row)
            question = _open_question(row, canonical)
            calculation = _calculation_gold(row, evidence) if scoring_type == "numeric" else None
            self._add_case(
                scenario=scenario,
                question=question,
                scoring_type=scoring_type,
                canonical_answer=canonical,
                required_conclusions=_conclusions(canonical),
                forbidden_conclusions=[],
                evidence=evidence,
                source_case_ids=[str(row.get("id"))],
                calculation=calculation,
            )

    def _add_cross_file(self, rows: list[dict[str, Any]]) -> None:
        for index in range(0, 20, 2):
            left, right = rows[index], rows[index + 1]
            if _primary_filename(left) == _primary_filename(right):
                raise ValueError("Cross-file pair unexpectedly uses the same source")
            left_answer, right_answer = _canonical_answer(left), _canonical_answer(right)
            question = (
                "请跨文件分别回答以下两个事项，并为每项给出对应来源。\n"
                f"关于前一份文件：{_open_question(left, left_answer)}\n"
                f"关于后一份文件：{_open_question(right, right_answer)}"
            )
            evidence = [self.corpus.evidence(item) for row in (left, right) for item in row.get("citations") or []]
            self._add_case(
                scenario="跨制度文件推理",
                question=question,
                scoring_type="multi_assertion",
                canonical_answer=f"事项一：{left_answer}\n事项二：{right_answer}",
                required_conclusions=[left_answer, right_answer],
                forbidden_conclusions=[],
                evidence=evidence,
                source_case_ids=[str(left.get("id")), str(right.get("id"))],
            )

    def _add_joint(self, text_rows: list[dict[str, Any]], table_rows: list[dict[str, Any]]) -> None:
        for text_row, table_row in zip(text_rows, table_rows, strict=True):
            text_answer = _canonical_answer(text_row)
            table_answer = _canonical_answer(table_row)
            question = (
                "请联合核验制度文件与统计报表，但不要把行业汇总统计直接当作单一机构的合规结论。\n"
                f"制度侧：{_open_question(text_row, text_answer)}\n"
                f"报表侧：{_open_table_question(table_row, table_answer)}\n"
                "证据边界：仅凭上述两项材料能否直接认定某一家银行或保险机构合规？请说明理由。"
            )
            evidence = [
                self.corpus.evidence(item)
                for row in (text_row, table_row)
                for item in row.get("citations") or []
            ]
            boundary = "上述材料只能支持各自的制度事实与行业统计事实，不能据此直接认定某一家机构合规。"
            self._add_case(
                scenario="制度与报表联合判断",
                question=question,
                scoring_type="multi_assertion",
                canonical_answer=f"制度侧：{text_answer}\n报表侧：{table_answer}\n证据边界：{boundary}",
                required_conclusions=[text_answer, table_answer, boundary],
                forbidden_conclusions=["可以直接认定该机构合规", "能够证明该机构合规"],
                evidence=evidence,
                source_case_ids=[str(text_row.get("id")), str(table_row.get("id"))],
            )

    def _match_joint_tables(
        self,
        text_rows: list[dict[str, Any]],
        table_rows: list[dict[str, Any]],
    ) -> list[dict[str, Any]]:
        selected: list[dict[str, Any]] = []
        use_counts: Counter[str] = Counter()
        for text_row in text_rows:
            text = (_primary_filename(text_row) or "") + str(text_row.get("question") or "")
            if any(term in text for term in ("商业银行", "银行版", "银行函证", "资本管理", "消费金融")):
                domain_terms = ("银行", "商业银行", "贷款", "资产", "负债")
            elif any(term in text for term in ("保险", "寿险", "偿付能力", "意外伤害")):
                domain_terms = ("保险", "险公司", "保费")
            else:
                domain_terms = ("银行", "商业银行", "贷款", "资产", "负债")
            candidates = [
                row
                for row in table_rows
                if any(term in ((_primary_filename(row) or "") + str(row.get("question") or "")) for term in domain_terms)
            ] or table_rows
            def table_load(row: dict[str, Any]) -> tuple[int, int, str]:
                filename = _primary_filename(row) or ""
                document = self.corpus.documents_by_filename.get(filename) or {}
                return (
                    self.primary_counts[str(document.get("document_id") or "")],
                    use_counts[filename],
                    str(row.get("id")),
                )

            selected.append(min(candidates, key=table_load))
            use_counts[_primary_filename(selected[-1]) or ""] += 1
        return selected

    def _add_formula_cases(self) -> None:
        specs = [
            ("422", ["RWA=E*R*12.5", "RWA＝E×R×12.5"], "请依据该 Word 公式计算 RWA：E=80，R=0.5。", "500", "formula_numeric", {"operands": {"E": "80", "R": "0.5"}, "unit": None, "period": None, "formula": "RWA=E×R×12.5", "rounding": "精确值", "result": "500"}),
            ("420", ["d=SD*Notional", "d＝SD×Notional"], "请依据该 Word 公式计算 d：SD=0.4，Notional=250。", "100", "formula_numeric", {"operands": {"SD": "0.4", "Notional": "250"}, "unit": None, "period": None, "formula": "d=SD×Notional", "rounding": "精确值", "result": "100"}),
            ("421", ["LGD^*=LGD_s*(E_s)/(E*(1+H_e))+LGD_u*(E_u)/(E*(1+H_e))"], "请依据该 Word 公式计算 LGD*：LGD_s=0.4，E_s=60，E=100，H_e=0.2，LGD_u=0.6，E_u=40。", "0.4", "formula_numeric", {"operands": {"LGD_s": "0.4", "E_s": "60", "E": "100", "H_e": "0.2", "LGD_u": "0.6", "E_u": "40"}, "unit": None, "period": None, "formula": "LGD^*=LGD_s*(E_s)/(E*(1+H_e))+LGD_u*(E_u)/(E*(1+H_e))", "rounding": "精确值", "result": "0.4"}),
            ("424", ["杠杆率=", "杠杆率＝"], "请依据该 Word 公式计算杠杆率：核心一级资本=900，核心一级资本扣除项=0，调整后表内外资产余额=10000。", "9%", "formula_numeric", {"operands": {"核心一级资本": "900", "核心一级资本扣除项": "0", "调整后表内外资产余额": "10000"}, "unit": "%", "period": None, "formula": "杠杆率=(核心一级资本-核心一级资本扣除项)/调整后表内外资产余额", "rounding": "百分比", "result": "9%"}),
            ("406", ["K_CM_i=", "KCMi="], "请依据该 Word 公式说明并计算 K_CM_i：K_CCP=10，DF_i=20，pref=1，DF_CCP=30，DF_CM=5；若公式运算能力不支持，必须明确拒答。", "unsupported_operation", "formula_refusal", None),
            ("407", ["l=max", "l＝max"], "请依据该 Word 公式计算 l：A=12，K_IRB=5；若 max 运算未被可信执行器支持，必须明确拒答。", "unsupported_operation", "formula_refusal", None),
            ("408", ["RWA_investment=min", "RWAinvestment=min"], "请依据该 Word 公式计算 RWA_investment：Avg=2，RW_fund=3，Lvg=4，Equity=100，Investment=1；若 min 运算未被可信执行器支持，必须明确拒答。", "unsupported_operation", "formula_refusal", None),
            ("410", ["SES=sqrt", "SES＝sqrt"], "请依据该 Word 公式计算 SES：ISES_NM=1，SES_NM=2，ρ=0.5；若平方根运算未被可信执行器支持，必须明确拒答。", "unsupported_operation", "formula_refusal", None),
        ]
        for prefix, needles, question, expected, scoring_type, calculation in specs:
            evidence = [self.corpus.formula_evidence(prefix, needles)]
            document = self.corpus.by_prefix(prefix)
            title = str(document.get("title") or _display_title(evidence[0].filename))
            question = question.replace("该 Word 公式", f"《{title}》中的 Word 公式")
            self._add_case(
                scenario="Word/PDF 公式",
                question=question,
                scoring_type=scoring_type,
                canonical_answer=expected,
                required_conclusions=[expected],
                forbidden_conclusions=[] if scoring_type == "formula_numeric" else ["已完成确定性计算"],
                evidence=evidence,
                source_case_ids=[f"FORMULA-{prefix}"],
                calculation=calculation,
                expected_refusal_code=expected if scoring_type == "formula_refusal" else None,
            )

    def _add_version_cases(self) -> None:
        prefixes = ["387", "394", "395", "425", "441", "443", "444", "445"]
        for prefix in prefixes:
            document = self.corpus.by_prefix(prefix)
            if document.get("version_status") != "unknown":
                raise ValueError(f"Version case source is not unknown: {document['filename']}")
            filename = str(document["filename"])
            title = _display_title(filename)
            question = (
                f"根据知识库中的文档身份信息，能否确认《{title}》截至当前为现行有效、已失效或已被替代？"
                "请勿根据文件名或日期推断；身份字段不足时应明确说明未知，并拒绝作法律效力断言。"
            )
            self._add_case(
                scenario="时效、版本与冲突",
                question=question,
                scoring_type="version_uncertainty",
                canonical_answer="该文档的版本状态为未知，无法仅凭现有身份信息确认其现行、失效或替代关系。",
                required_conclusions=["版本状态未知", "无法确认法律效力"],
                forbidden_conclusions=["确认现行有效", "确认已经失效", "确认已被替代"],
                evidence=[],
                source_case_ids=[f"VERSION-{prefix}"],
                identity_source={
                    "document_id": document["document_id"],
                    "filename": filename,
                    "file_sha256": document["file_sha256"],
                    "version_status": document["version_status"],
                    "identity_review_status": document["identity_review_status"],
                },
                expected_refusal_code="version_status_unknown",
            )

    def _add_refusal_cases(self) -> None:
        selected: list[dict[str, Any]] = []
        seen_titles: set[str] = set()
        for row in self.ood_rows:
            title = _title_from_question(str(row.get("question") or ""))
            intended = _find_document_for_title(self.corpus, title)
            if intended and self.primary_counts[str(intended["document_id"])] >= 4:
                continue
            if title and title not in seen_titles:
                selected.append(row)
                seen_titles.add(title)
            if len(selected) == 10:
                break
        if len(selected) != 10:
            raise ValueError("Unable to select ten distinct OOD refusal cases")
        for row in selected:
            question = _open_ood_question(str(row["question"]))
            refusal_code = str(row.get("refusal_reason") or "insufficient_context")
            search_term = _ood_search_term(question)
            corpus_hits = 0
            if search_term:
                corpus_hits = self.corpus.connection.execute(
                    "SELECT COUNT(*) FROM document_chunks WHERE text LIKE ?",
                    (f"%{search_term}%",),
                ).fetchone()[0]
            if corpus_hits:
                raise ValueError(f"OOD marker unexpectedly occurs in corpus: {search_term}")
            intended = _find_document_for_title(self.corpus, _title_from_question(question))
            identity_source = None
            if intended:
                identity_source = {
                    "document_id": intended["document_id"],
                    "filename": intended["filename"],
                    "file_sha256": intended["file_sha256"],
                }
            self._add_case(
                scenario="可信拒答与歧义",
                question=question,
                scoring_type="refusal",
                canonical_answer="当前知识库依据不足，无法给出确定答案。",
                required_conclusions=["依据不足", "无法给出确定答案"],
                forbidden_conclusions=["已核实存在该条款", "给出确定数值"],
                evidence=[],
                source_case_ids=[str(row.get("id"))],
                identity_source=identity_source,
                expected_refusal_code=refusal_code,
                corpus_search={
                    "scope": "all_500_document_chunks",
                    "query": search_term,
                    "match_count": corpus_hits,
                    "insufficiency_type": "nonexistent_period_or_clause",
                },
            )

    def _add_case(
        self,
        *,
        scenario: str,
        question: str,
        scoring_type: str,
        canonical_answer: str,
        required_conclusions: list[str],
        forbidden_conclusions: list[str],
        evidence: list[Evidence],
        source_case_ids: list[str],
        calculation: dict[str, Any] | None = None,
        expected_refusal_code: str | None = None,
        identity_source: dict[str, Any] | None = None,
        corpus_search: dict[str, Any] | None = None,
    ) -> None:
        self.case_number += 1
        case_id = f"TC{self.case_number:03d}"
        documents = _unique_documents(evidence, identity_source)
        primary = self._choose_primary(documents)
        question_record = {
            "id": case_id,
            "scenario": scenario,
            "difficulty": None,
            "split": None,
            "question": question,
            "scoring_type": scoring_type,
        }
        gold_record = {
            **question_record,
            "review_status": "codex_verified",
            "expert_reviewed": False,
            "canonical_answer": canonical_answer,
            "allowed_expressions": _allowed_expressions(canonical_answer),
            "required_conclusions": required_conclusions,
            "forbidden_conclusions": forbidden_conclusions,
            "expected_refusal_code": expected_refusal_code,
            "primary_document_id": primary.get("document_id") if primary else None,
            "required_documents": documents,
            "evidence": [item.as_dict() for item in evidence],
            "calculation": calculation,
            "identity_source": identity_source,
            "corpus_search": corpus_search,
            "source_case_ids": source_case_ids,
        }
        self.cases.append({"question_record": question_record, "gold_record": gold_record})

    def _choose_primary(self, documents: list[dict[str, Any]]) -> dict[str, Any] | None:
        if not documents:
            return None
        eligible = [item for item in documents if self.primary_counts[item["document_id"]] < 4]
        if not eligible:
            raise ValueError(f"All required documents already have four primary cases: {[item['document_id'] for item in documents]}")
        document = min(eligible, key=lambda item: (self.primary_counts[item["document_id"]], item["document_id"]))
        self.primary_counts[document["document_id"]] += 1
        return document

    def _assign_difficulty_and_split(self) -> None:
        by_scenario: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for case in self.cases:
            by_scenario[case["question_record"]["scenario"]].append(case)
        for scenario, spec in SCENARIO_SPECS.items():
            cases = by_scenario.get(scenario) or []
            expected = spec["medium"] + spec["hard"]
            if len(cases) != expected:
                raise ValueError(f"{scenario}: expected {expected}, got {len(cases)}")
            assignments = (
                [("medium", "dev")] * spec["dev_medium"]
                + [("medium", "holdout")] * (spec["medium"] - spec["dev_medium"])
                + [("hard", "dev")] * spec["dev_hard"]
                + [("hard", "holdout")] * (spec["hard"] - spec["dev_hard"])
            )
            for case, (difficulty, split) in zip(cases, assignments, strict=True):
                case["question_record"]["difficulty"] = difficulty
                case["question_record"]["split"] = split
                case["gold_record"]["difficulty"] = difficulty
                case["gold_record"]["split"] = split


def validate_cases(cases: list[dict[str, Any]], corpus: Corpus) -> dict[str, Any]:
    questions = [case["question_record"] for case in cases]
    gold = [case["gold_record"] for case in cases]
    ids = [row["id"] for row in questions]
    if len(cases) != 100 or len(set(ids)) != 100:
        raise ValueError("Challenge must contain exactly 100 unique cases")
    if any(set(row) - {"id", "scenario", "difficulty", "split", "question", "scoring_type"} for row in questions):
        raise ValueError("Online question file contains fields outside the public contract")
    scenario_counts = Counter((row["scenario"], row["difficulty"]) for row in questions)
    split_counts = Counter((row["scenario"], row["split"], row["difficulty"]) for row in questions)
    for scenario, spec in SCENARIO_SPECS.items():
        if scenario_counts[(scenario, "medium")] != spec["medium"]:
            raise ValueError(f"Medium distribution mismatch for {scenario}")
        if scenario_counts[(scenario, "hard")] != spec["hard"]:
            raise ValueError(f"Hard distribution mismatch for {scenario}")
        if split_counts[(scenario, "dev", "medium")] != spec["dev_medium"]:
            raise ValueError(f"Dev medium distribution mismatch for {scenario}")
        if split_counts[(scenario, "dev", "hard")] != spec["dev_hard"]:
            raise ValueError(f"Dev hard distribution mismatch for {scenario}")
    if Counter(row["split"] for row in questions) != Counter({"holdout": 60, "dev": 40}):
        raise ValueError("Challenge split must be 40 development / 60 holdout")
    if Counter(row["difficulty"] for row in questions) != Counter({"hard": 70, "medium": 30}):
        raise ValueError("Challenge difficulty must be 30 medium / 70 hard")

    primary_counts = Counter(row["primary_document_id"] for row in gold if row.get("primary_document_id"))
    if primary_counts and max(primary_counts.values()) > 4:
        raise ValueError(f"A primary source exceeds four cases: {primary_counts.most_common(3)}")
    source_documents = {
        document["document_id"]
        for row in gold
        for document in row.get("required_documents") or []
    }
    if len(source_documents) < 50:
        raise ValueError(f"Challenge covers only {len(source_documents)} distinct official documents")
    for row in gold:
        for item in row.get("evidence") or []:
            chunk = corpus.chunks.get(item["chunk_id"])
            if item.get("anchor_type") == "sqlite_chunk":
                if not chunk or _sha256_text(str(chunk["text"] or "")) != item["chunk_sha256"]:
                    raise ValueError(f"Evidence chunk hash mismatch: {item['chunk_id']}")
            elif item.get("anchor_type") == "source_parser_block":
                if item["chunk_sha256"] != _sha256_text(item["evidence_text"]):
                    raise ValueError(f"Formula parser block hash mismatch: {item['chunk_id']}")
            else:
                raise ValueError(f"Unknown evidence anchor type: {item.get('anchor_type')}")
            if _sha256_text(item["evidence_text"]) != item["evidence_text_sha256"]:
                raise ValueError(f"Evidence excerpt hash mismatch: {item['chunk_id']}")
    return {
        "case_count": len(cases),
        "difficulty": dict(Counter(row["difficulty"] for row in questions)),
        "split": dict(Counter(row["split"] for row in questions)),
        "scenario": dict(Counter(row["scenario"] for row in questions)),
        "scenario_difficulty": {f"{key[0]}::{key[1]}": value for key, value in sorted(scenario_counts.items())},
        "distinct_required_documents": len(source_documents),
        "distinct_primary_documents": len(primary_counts),
        "max_cases_per_primary_document": max(primary_counts.values()) if primary_counts else 0,
        "evidence_anchor_count": sum(len(row.get("evidence") or []) for row in gold),
        "question_gold_isolated": True,
    }


def _load_results(path: Path) -> list[dict[str, Any]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    rows = payload.get("results") or []
    if not isinstance(rows, list) or not rows:
        raise ValueError(f"No result rows found in {path}")
    return rows


def _primary_filename(row: dict[str, Any]) -> str | None:
    citations = row.get("citations") or []
    return str(citations[0].get("filename")) if citations and citations[0].get("filename") else None


def _canonical_answer(row: dict[str, Any]) -> str:
    answer = re.sub(r"\s*\[\d+\]", "", str(row.get("answer") or "")).strip()
    answer = re.sub(
        r"^(?:根据检索到的知识片段，正确表述是：|"
        r"根据知识库，(?:正确的?表述(?:是|为)|与材料内容一致的选项是)：|"
        r"与材料内容一致的是：|答案为\s*)",
        "",
        answer,
    )
    answer = answer.replace("依据表格证据，核验值为", "数值为")
    answer = re.sub(r"^根据知识库，选项[“\"](.+?)[”\"]表述正确(?:。|$)", r"\1。", answer, flags=re.S)
    answer = re.sub(r"^根据知识库，", "", answer)
    if row.get("source_type") == "excel":
        duplicate = re.fullmatch(r"\s*([-+]?\d+(?:\.\d+)?)。数值为\s*\1。?", answer)
        if duplicate:
            answer = duplicate.group(1)
        else:
            answer = re.sub(r"^(.+?)。数值为\s*([-+]?\d+(?:\.\d+)?)。?$", r"\1，数值为\2", answer)
    answer = answer.replace("。；", "；")
    return answer.strip("。； ") + "。"


def _text_row_evidence_quality(row: dict[str, Any]) -> tuple[float, int]:
    answer = _normalized_evidence_text(_canonical_answer(row))
    evidence = _normalized_evidence_text(
        "\n".join(str(item.get("excerpt") or "") for item in row.get("citations") or [])
    )
    if len(answer) < 2 or len(evidence) < 2:
        return 0.0, len(evidence)
    answer_pairs = {answer[index:index + 2] for index in range(len(answer) - 1)}
    evidence_pairs = {evidence[index:index + 2] for index in range(len(evidence) - 1)}
    return len(answer_pairs & evidence_pairs) / len(answer_pairs), len(evidence)


def _normalized_evidence_text(value: str) -> str:
    return re.sub(r"[^0-9A-Za-z\u4e00-\u9fff]", "", str(value or "")).casefold()


def _open_question(row: dict[str, Any], canonical: str) -> str:
    if row.get("source_type") == "excel":
        return _open_table_question(row, canonical)
    title = _title_from_question(str(row.get("question") or "")) or _display_title(_primary_filename(row) or "该文件")
    cues = _question_cues(canonical)
    cue_text = "、".join(f"“{cue}”" for cue in cues)
    return (
        f"根据《{title}》，请说明与{cue_text}相关的明确规定。"
        "请完整给出材料明确写明的条件、范围、期限或例外。"
    )


def _question_cue(answer: str) -> str:
    clause = re.sub(r"[。；].*$", "", answer).strip()
    clause = re.sub(r"^(?:根据《[^》]+》(?:第[^，。；]+)?，?)", "", clause)
    clause = re.sub(r"\s+", "", clause)
    for marker in ("应当", "不得", "属于", "包括", "是", "在", "由", "可以", "需要", "采用", "按照", "为"):
        position = clause.find(marker)
        if 5 <= position <= 30:
            return clause[:position].rstrip("，：")
    return clause[:36].rstrip("，：")


def _question_cues(answer: str) -> list[str]:
    clauses = _conclusions(answer)
    cues = [_question_cue(clause) for clause in clauses[:2]]
    return list(dict.fromkeys(cue for cue in cues if cue)) or [_question_cue(answer)]


def _open_table_question(row: dict[str, Any], canonical: str | None = None) -> str:
    question = str(row.get("question") or "").strip()
    options = [str(value) for value in row.get("_options") or [] if str(value).strip()]
    options = _comparable_table_options(canonical or _canonical_answer(row), options)
    candidate_text = "、".join(f"“{value}”" for value in options)
    highest = f"请比较{candidate_text}，找出数值最高的项目，并给出该项目名称、数值和单位。" if candidate_text else "请找出数值最高的项目，并给出该项目名称、数值和单位。"
    lowest = f"请比较{candidate_text}，找出数值最低的项目，并给出该项目名称、数值和单位。" if candidate_text else "请找出数值最低的项目，并给出该项目名称、数值和单位。"
    question = question.replace("以下哪一项数值最高？", highest)
    question = question.replace("以下哪一项数值最低？", lowest)
    question = question.replace("下列哪一项数值最高？", highest)
    question = question.replace("下列哪一项数值最低？", lowest)
    return question


def _comparable_table_options(canonical: str, options: list[str]) -> list[str]:
    """Remove candidates that are not comparable with the locked metric.

    Insurance workbooks place premium flows, insured amounts and balance-sheet
    stocks in one numeric column. A mathematical max across those dimensions is
    meaningless even when every cell uses 亿元, so retain only the premium
    total and its business-line components.
    """

    if "原保险保费收入" not in canonical:
        return options
    comparable = [
        option
        for option in options
        if not (
            "资产" in option
            or ("保险金额" in option and "保费收入" not in option)
        )
    ]
    return comparable if len(comparable) >= 2 else options


def _calculation_gold(row: dict[str, Any], evidence: list[Evidence]) -> dict[str, Any]:
    canonical = _canonical_answer(row)
    numbers = re.findall(r"(?<![A-Za-z_])[-+]?\d+(?:\.\d+)?%?", canonical)
    metadata = (row.get("citations") or [{}])[0].get("metadata") or {}
    operation = metadata.get("operation") or metadata.get("calculation") or {}
    unit = _table_unit(metadata)
    period = metadata.get("period")
    return {
        "operands": operation.get("operands") if isinstance(operation, dict) else None,
        "unit": unit,
        "period": period,
        "formula": operation.get("expression") if isinstance(operation, dict) else None,
        "rounding": "按题目‘约为’或原表精度；未指定时保持证据精度",
        "result": numbers[-1] if numbers else None,
        "evidence_cells": sorted({cell for item in evidence for cell in item.cells}),
    }


def _conclusions(answer: str) -> list[str]:
    parts = [part.strip() for part in re.split(r"[；。\n]+", answer) if len(part.strip()) >= 2]
    return parts or [answer.strip()]


def _allowed_expressions(answer: str) -> list[str]:
    variants = {answer.strip(), answer.strip("。")}
    variants.add(answer.replace("必须", "应当").strip())
    variants.add(answer.replace("应当", "必须").strip())
    return sorted(value for value in variants if value)


def _unique_documents(evidence: list[Evidence], identity_source: dict[str, Any] | None) -> list[dict[str, Any]]:
    records: dict[str, dict[str, Any]] = {}
    for item in evidence:
        records[item.document_id] = {
            "document_id": item.document_id,
            "filename": item.filename,
            "file_sha256": item.file_sha256,
        }
    if identity_source and identity_source.get("document_id"):
        records[str(identity_source["document_id"])] = {
            "document_id": identity_source["document_id"],
            "filename": identity_source["filename"],
            "file_sha256": identity_source["file_sha256"],
        }
    return sorted(records.values(), key=lambda item: item["document_id"])


def _find_document_for_title(corpus: Corpus, title: str | None) -> dict[str, Any] | None:
    if not title:
        return None
    compact = re.sub(r"\s+", "", title)
    candidates = [row for row in corpus.documents.values() if compact in re.sub(r"\s+", "", str(row["filename"]))]
    return min(candidates, key=lambda row: len(str(row["filename"]))) if candidates else None


def _title_from_question(question: str) -> str | None:
    match = re.search(r"《([^》]+)》", question)
    return match.group(1).strip() if match else None


def _display_title(filename: str) -> str:
    name = Path(filename).stem
    name = re.sub(r"^\d+_", "", name)
    parts = name.split("_")
    return parts[-1] if parts else name


def _ood_search_term(question: str) -> str:
    marker = re.search(r"(Z-\d{3})", question)
    if marker:
        return marker.group(1)
    if "2099年" in question:
        return "2099年"
    return ""


def _open_ood_question(question: str) -> str:
    question = question.replace("以下哪一项数值最高？", "请找出数值最高的项目，并给出项目名称、数值和单位。")
    question = question.replace("以下哪一项数值最低？", "请找出数值最低的项目，并给出项目名称、数值和单位。")
    question = question.replace("下列哪一项数值最高？", "请找出数值最高的项目，并给出项目名称、数值和单位。")
    question = question.replace("下列哪一项数值最低？", "请找出数值最低的项目，并给出项目名称、数值和单位。")
    return question


def _safe_json(value: str | None) -> dict[str, Any]:
    try:
        payload = json.loads(value or "{}")
        return payload if isinstance(payload, dict) else {}
    except json.JSONDecodeError:
        return {}


def _table_unit(metadata: dict[str, Any]) -> str | None:
    unit = str(metadata.get("unit") or "").strip()
    if not unit:
        return None
    alternatives = [item.strip() for item in re.split(r"[、,/，]", unit) if item.strip()]
    if len(alternatives) <= 1:
        return unit
    label = " ".join(str(metadata.get(key) or "") for key in ("row_label", "column_label", "metric", "indicator"))
    if any(term in label for term in ("率", "比例", "占比", "比率")) and "%" in alternatives:
        return "%"
    if any(term in label for term in ("件数", "保单", "数量", "户数", "家数")):
        count_unit = next((item for item in alternatives if "件" in item or "户" in item or "家" in item), None)
        if count_unit:
            return count_unit
    return next((item for item in alternatives if "元" in item), alternatives[0])


def _attach_workbook_options(rows: list[dict[str, Any]], path: Path) -> None:
    worksheet = load_workbook(path, read_only=True, data_only=True).active
    headers = [str(cell.value or "") for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
    options: dict[str, list[str]] = {}
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        record = dict(zip(headers, values, strict=False))
        options[str(record.get("id"))] = [str(record.get(f"option_{letter}") or "") for letter in "abcd"]
    for row in rows:
        row["_options"] = options.get(str(row.get("id")), [])


def _normalize_formula(value: str) -> str:
    normalized = value.translate(str.maketrans({"×": "*", "＝": "=", "（": "(", "）": ")"}))
    return re.sub(r"[\s{}\\]", "", normalized).lower()


def _sha256_text(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _write_jsonl(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    path.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows), encoding="utf-8")


def _relative(path: Path) -> str:
    try:
        return str(path.resolve().relative_to(PROJECT_ROOT)).replace("\\", "/")
    except ValueError:
        return str(path.resolve())


if __name__ == "__main__":
    raise SystemExit(main())
