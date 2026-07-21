from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import random
import re
import statistics
import time
import urllib.error
import urllib.request
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_QA = PROJECT_ROOT / "data" / "contest_dataset" / "QA数据.xlsx"
DEFAULT_OUTPUT = PROJECT_ROOT / "data" / "evaluation" / "contest_qa"
SPLIT_SEED = 20260713


@dataclass(frozen=True)
class Case:
    id: str
    source_type: str
    difficulty: str
    qa_type: str
    question: str
    options: tuple[str, ...]
    answer: str
    answer_text: str
    evidence: str
    file_label: str
    source_title: str = ""
    split: str = ""


def main() -> int:
    parser = argparse.ArgumentParser(description="Evaluate ReguMate against the official 300-question dataset.")
    parser.add_argument("--qa", type=Path, default=DEFAULT_QA)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--base-url", default="http://127.0.0.1:8000")
    parser.add_argument("--split", choices=["dev", "holdout", "all", "ood"], default="all")
    parser.add_argument("--source-type", choices=["excel", "word", "pdf"], help="Optional diagnostic subset.")
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--start", type=int, default=0, help="Zero-based offset after split selection.")
    parser.add_argument("--timeout", type=float, default=30.0)
    parser.add_argument("--prepare-only", action="store_true")
    parser.add_argument("--resume", action="store_true", help="Resume from the per-split checkpoint.")
    parser.add_argument("--retries", type=int, default=2)
    parser.add_argument("--request-delay", type=float, default=0.15)
    parser.add_argument(
        "--capture-performance",
        action="store_true",
        help="Capture the latest request trace, resource sample, and cache counters after every case.",
    )
    parser.add_argument(
        "--evidence-manifest",
        type=Path,
        help="Optional reviewed evidence manifest for Recall/rerank/citation metrics.",
    )
    args = parser.parse_args()

    cases = assign_splits(read_cases(args.qa))
    ood_cases = build_ood_cases(cases)
    args.output.mkdir(parents=True, exist_ok=True)
    write_split_manifest(args.output / "contest_qa_splits.json", cases, ood_cases)
    if args.prepare_only:
        print_split_summary(cases, ood_cases)
        return 0

    selected = ood_cases if args.split == "ood" else [
        case for case in cases if args.split == "all" or case.split == args.split
    ]
    if args.source_type:
        selected = [case for case in selected if case.source_type == args.source_type]
    selected = selected[args.start:]
    if args.limit > 0:
        selected = selected[: args.limit]
    checkpoint_path = args.output / f"contest_qa_{args.split}_checkpoint.json"
    prior = _load_checkpoint(checkpoint_path) if args.resume else {}
    evidence_manifest = _load_evidence_manifest(args.evidence_manifest)
    results: list[dict[str, Any]] = []
    for position, case in enumerate(selected, start=1):
        result = prior.get(case.id)
        if result is None:
            result = evaluate_case(
                case,
                args.base_url,
                args.timeout,
                args.retries,
                capture_performance=args.capture_performance,
            )
            prior[case.id] = result
            _write_checkpoint(checkpoint_path, prior)
            if args.request_delay > 0 and position < len(selected):
                time.sleep(args.request_delay)
        _reconcile_grounded_prediction(result)
        if case.id in evidence_manifest:
            result["manifest_evidence"] = score_manifest_evidence(
                result, evidence_manifest[case.id]
            )
        results.append(result)
        print(
            f"[{position}/{len(selected)}] {case.id} correct={result['answer_correct']} "
            f"source={result['source_hit']} elapsed={result['elapsed_ms']}ms",
            flush=True,
        )
    artifact = {
        "run": {
            "created_at": datetime.now(timezone.utc).isoformat(),
            "split": args.split,
            "case_count": len(selected),
            "base_url": args.base_url,
            "backend": fetch_backend_identity(args.base_url),
            "qa_sha256": sha256_file(args.qa),
            "evaluator_sha256": sha256_file(Path(__file__)),
            "split_seed": SPLIT_SEED,
            "source_type_filter": args.source_type,
            "evidence_manifest_sha256": (
                sha256_file(args.evidence_manifest) if args.evidence_manifest else None
            ),
        },
        "summary": summarize(results),
        "results": results,
    }
    result_path = args.output / f"contest_qa_{args.split}_results.json"
    result_path.write_text(json.dumps(artifact, ensure_ascii=False, indent=2), encoding="utf-8")
    report_path = args.output / f"contest_qa_{args.split}_report.md"
    report_path.write_text(render_report(artifact), encoding="utf-8")
    print(json.dumps(artifact["summary"], ensure_ascii=False, indent=2))
    print(f"results: {result_path}\nreport: {report_path}")
    return 0


def read_cases(path: Path) -> list[Case]:
    worksheet = load_workbook(path, read_only=True, data_only=True).active
    headers = [str(cell.value or "") for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
    cases: list[Case] = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        record = dict(zip(headers, row, strict=False))
        cases.append(
            Case(
                id=str(record["id"]),
                source_type=str(record["source_type"]),
                difficulty=str(record.get("difficulty_cn") or record.get("difficulty") or ""),
                qa_type=str(record["qa_type"]),
                question=str(record["question"]),
                options=tuple(str(record[f"option_{letter}"]) for letter in "abcd"),
                answer=str(record["answer"]).upper(),
                answer_text=str(record["answer_text"]),
                evidence=str(record["evidence"]),
                file_label=str(record["file_label"]),
                source_title=str(record.get("source_title") or ""),
            )
        )
    return cases


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def fetch_backend_identity(base_url: str) -> dict[str, Any]:
    try:
        with urllib.request.urlopen(
            f"{base_url.rstrip('/')}/api/health/rag", timeout=10
        ) as response:
            body = json.loads(response.read().decode("utf-8"))
        return {
            "build_id": body.get("build_id"),
            "collection": body.get("qdrant", {}).get("collection"),
            "ready": body.get("ready"),
            "performance": body.get("performance") or {},
            "models": body.get("models") or {},
        }
    except (OSError, urllib.error.URLError, json.JSONDecodeError) as exc:
        return {"build_id": None, "ready": None, "identity_error": str(exc)}


def assign_splits(cases: list[Case]) -> list[Case]:
    groups: dict[tuple[str, str, str], list[Case]] = defaultdict(list)
    for case in cases:
        groups[(case.source_type, case.qa_type, case.difficulty)].append(case)
    holdout_ids: set[str] = set()
    for key, items in groups.items():
        seeded = random.Random(f"{SPLIT_SEED}:{key}")
        ordered = list(items)
        seeded.shuffle(ordered)
        count = max(1, round(len(ordered) * 0.2))
        holdout_ids.update(case.id for case in ordered[:count])
    # Rebalance to exactly 60 without disturbing deterministic ordering.
    stable = sorted(cases, key=lambda case: hashlib.sha256(f"{SPLIT_SEED}:{case.id}".encode()).hexdigest())
    if len(holdout_ids) > 60:
        for case in stable:
            if len(holdout_ids) <= 60:
                break
            holdout_ids.discard(case.id)
    elif len(holdout_ids) < 60:
        for case in stable:
            if len(holdout_ids) >= 60:
                break
            holdout_ids.add(case.id)
    return [Case(**{**case.__dict__, "split": "holdout" if case.id in holdout_ids else "dev"}) for case in cases]


def build_ood_cases(cases: list[Case]) -> list[Case]:
    selected = sorted(cases, key=lambda case: case.id)[::10][:30]
    result: list[Case] = []
    for index, case in enumerate(selected, start=1):
        if case.source_type == "excel":
            if index % 2:
                question = re.sub(r"20\d{2}年", "2099年", case.question, count=1)
                if question == case.question:
                    question = case.question + "（限定为不存在的2099年数据）"
            else:
                question = (
                    f"请在《{Path(case.file_label).stem}》中查询不存在的指标"
                    f"“虚构监管指标Z-{index:03d}”的数值。"
                )
        else:
            question = f"根据《{Path(case.file_label).stem}》，不存在的监管条款Z-{index:03d}具体规定了什么？"
        result.append(
            Case(
                id=f"OOD{index:03d}", source_type=case.source_type, difficulty="越界",
                qa_type="无依据拒答", question=question, options=(), answer="REFUSE",
                answer_text="", evidence="", file_label=case.file_label,
                source_title=case.source_title, split="ood",
            )
        )
    return result


def evaluate_case(
    case: Case,
    base_url: str,
    timeout: float,
    retries: int = 2,
    *,
    capture_performance: bool = False,
) -> dict[str, Any]:
    payload = {"question": case.question, "options": list(case.options), "include_debug": True}
    started = datetime.now(timezone.utc)
    body: dict[str, Any] = {}
    error = None
    error_type = None
    attempts = 0
    for attempt in range(retries + 1):
        attempts = attempt + 1
        try:
            request = urllib.request.Request(
                f"{base_url.rstrip('/')}/api/qa/ask",
                data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
                headers={"Content-Type": "application/json"}, method="POST",
            )
            with urllib.request.urlopen(request, timeout=timeout) as response:
                body = json.loads(response.read().decode("utf-8"))
            error = None
            error_type = None
            break
        except urllib.error.HTTPError as exc:
            error, error_type = f"HTTP {exc.code}: {exc.reason}", "http"
        except TimeoutError as exc:
            error, error_type = str(exc), "timeout"
        except urllib.error.URLError as exc:
            error, error_type = str(exc), "connection"
        except json.JSONDecodeError as exc:
            error, error_type = str(exc), "invalid_json"
        if attempt < retries:
            time.sleep(min(2 ** attempt, 4))
    elapsed_ms = (datetime.now(timezone.utc) - started).total_seconds() * 1000
    answer = str(body.get("answer") or "")
    predicted = grounded_option(body) or extract_option(answer, case.options)
    expected_refusal = case.answer == "REFUSE"
    answer_correct = bool(body.get("refused")) if expected_refusal else predicted == case.answer
    citations = body.get("citations") or []
    context_package = body.get("context_package") or {}
    source_hit = any(source_matches(case, str(item.get("filename") or "")) for item in citations)
    excel_evidence = score_excel_evidence(case, citations) if case.source_type == "excel" else None
    text_evidence = score_text_evidence(case, citations) if case.source_type != "excel" else None
    performance = fetch_performance_snapshot(base_url) if capture_performance else None
    return {
        "id": case.id, "split": case.split, "source_type": case.source_type,
        "qa_type": case.qa_type, "difficulty": case.difficulty, "question": case.question,
        "expected": case.answer, "predicted": predicted, "answer": answer,
        "answer_correct": answer_correct, "refused": bool(body.get("refused")),
        "source_hit": source_hit, "excel_evidence": excel_evidence,
        "text_evidence": text_evidence,
        "answer_type": body.get("answer_type"), "generation_status": body.get("generation_status"),
        "grounding_validation": body.get("grounding_validation") or {},
        "citations": citations,
        "context_package": context_package,
        "claims": body.get("claims") or [],
        "refusal_reason": body.get("refusal_reason"),
        "degraded": bool(body.get("degraded")),
        "elapsed_ms": round(elapsed_ms, 2), "error": error,
        "error_type": error_type, "attempts": attempts,
        "citation_count": len(citations),
        "performance": performance,
    }


def grounded_option(body: dict[str, Any]) -> str | None:
    validation = body.get("grounding_validation") or {}
    if not (
        validation.get("passed")
        and validation.get("selected_option_supported")
        and validation.get("selected_option_format_valid")
    ):
        return None
    raw = validation.get("selected_option")
    if isinstance(raw, int) or (isinstance(raw, str) and raw.strip().isdigit()):
        index = int(raw)
        return chr(65 + index) if 0 <= index < 4 else None
    label = str(raw or "").strip().upper()
    return label if label in {"A", "B", "C", "D"} else None


def _reconcile_grounded_prediction(result: dict[str, Any]) -> None:
    predicted = grounded_option({"grounding_validation": result.get("grounding_validation")})
    if predicted is None or result.get("expected") == "REFUSE":
        return
    result["predicted"] = predicted
    result["answer_correct"] = predicted == result.get("expected")


def fetch_performance_snapshot(base_url: str) -> dict[str, Any] | None:
    try:
        with urllib.request.urlopen(
            f"{base_url.rstrip('/')}/api/health/rag", timeout=10
        ) as response:
            body = json.loads(response.read().decode("utf-8"))
        performance = body.get("performance") or {}
        traces = performance.get("recent_traces") or []
        qa_traces = [item for item in traces if item.get("kind") == "qa"]
        return {
            "trace": qa_traces[-1] if qa_traces else None,
            "resource": (performance.get("resources") or {}).get("current") or {},
            "resource_peaks": (performance.get("resources") or {}).get("peaks") or {},
            "model_runtime": body.get("model_runtime") or {},
        }
    except (OSError, urllib.error.URLError, json.JSONDecodeError):
        return None


def _load_evidence_manifest(path: Path | None) -> dict[str, dict[str, Any]]:
    if path is None:
        return {}
    payload = json.loads(path.read_text(encoding="utf-8"))
    return {str(item["case_id"]): item for item in payload.get("cases") or []}


def score_manifest_evidence(
    result: dict[str, Any],
    manifest_case: dict[str, Any],
) -> dict[str, Any]:
    citations = result.get("citations") or []
    cited_ids = {str(item.get("chunk_id")) for item in citations}
    context_package = result.get("context_package") or {}
    final_ids = {
        str(item.get("chunk_id"))
        for item in context_package.get("context_chunks") or []
    }
    reranked_ids: list[str] = []
    for aspect in (context_package.get("retrieval_summary") or {}).get("aspect_retrievals") or []:
        for diagnostic in aspect.get("diagnostics") or []:
            if diagnostic.get("query_type") == "aspect_fused":
                reranked_ids.extend(
                    str(item.get("chunk_id"))
                    for item in diagnostic.get("reranked_ranking") or []
                )
    required = [
        {
            **aspect,
            "acceptable_chunk_ids": {
                str(chunk_id) for chunk_id in aspect.get("acceptable_chunk_ids") or []
            },
        }
        for aspect in manifest_case.get("required_aspects") or []
    ]
    evaluable = [aspect for aspect in required if aspect["acceptable_chunk_ids"]]
    final_covered = sum(bool(aspect["acceptable_chunk_ids"] & final_ids) for aspect in evaluable)
    cited_covered = sum(bool(aspect["acceptable_chunk_ids"] & cited_ids) for aspect in evaluable)
    reranked_set = set(reranked_ids)
    rerank_covered = sum(bool(aspect["acceptable_chunk_ids"] & reranked_set) for aspect in evaluable)
    acceptable_union = (
        set().union(*(aspect["acceptable_chunk_ids"] for aspect in evaluable))
        if evaluable
        else set()
    )
    supported_citations = sum(
        str(item.get("chunk_id")) in acceptable_union for item in citations
    )
    ranks = [
        index
        for index, chunk_id in enumerate(reranked_ids, start=1)
        if chunk_id in acceptable_union
    ]
    return {
        "review_status": manifest_case.get("review_status"),
        "required_aspect_count": len(required),
        "evaluable_aspect_count": len(evaluable),
        "final_recall": final_covered / len(evaluable) if evaluable else None,
        "rerank_retention": (
            rerank_covered / len(evaluable) if evaluable and reranked_ids else None
        ),
        "citation_completeness": cited_covered / len(evaluable) if evaluable else None,
        "citation_accuracy": (
            supported_citations / len(citations)
            if citations and acceptable_union
            else None
        ),
        "best_acceptable_rerank_rank": min(ranks) if ranks else None,
    }


def extract_option(answer: str, options: tuple[str, ...]) -> str | None:
    match = re.search(r"(?:答案(?:为|是)?\s*|^)([A-D])(?:[：:。\s]|$)", answer.upper())
    if match:
        return match.group(1)
    answer_without_citations = re.sub(r"\[\d+\]", " ", answer)
    numeric_options = [_parse_number(option) for option in options]
    if options and all(value is not None for value in numeric_options):
        answer_numbers = [
            value
            for token in re.findall(r"[-+]?\d[\d,]*(?:\.\d+)?%?", answer_without_citations)
            for value in [_parse_number(token)]
            if value is not None
        ]
        matches: list[tuple[float, int]] = []
        for index, option_value in enumerate(numeric_options):
            assert option_value is not None
            for answer_value in answer_numbers:
                if math.isclose(answer_value, option_value, rel_tol=1e-4, abs_tol=0.02):
                    matches.append((abs(answer_value - option_value), index))
        if matches:
            _, best_index = min(matches)
            return chr(65 + best_index)

    normalized = normalize_option_text(answer_without_citations)
    matched_options = [
        (len(normalized_option), index)
        for index, option in enumerate(options)
        for normalized_option in [normalize_option_text(option)]
        if normalized_option and normalized_option in normalized
    ]
    if matched_options:
        _, best_index = max(matched_options)
        return chr(65 + best_index)
    return None


def _parse_number(value: Any) -> float | None:
    text = str(value or "").strip().replace(",", "")
    percent = text.endswith("%")
    if percent:
        text = text[:-1]
    try:
        number = float(text)
    except ValueError:
        return None
    return number / 100 if percent else number


def normalize_option_text(value: Any) -> str:
    return re.sub(r"[^0-9a-zA-Z%\u4e00-\u9fff]+", "", str(value or "")).lower()


def source_matches(case: Case, actual_filename: str) -> bool:
    actual = _normalize_source_name(actual_filename)
    if not actual:
        return False
    expected_names = [case.file_label, case.source_title]
    for expected_name in expected_names:
        expected = _normalize_source_name(expected_name)
        if len(expected) >= 6 and (expected in actual or actual in expected):
            return True
    return False


def _normalize_source_name(value: str) -> str:
    stem = Path(str(value or "")).stem.lower()
    stem = re.sub(r"^\d+_", "", stem)
    stem = re.sub(r"[（(]\s*(?:pdf|word|docx?)\s*[）)]$", "", stem, flags=re.IGNORECASE)
    stem = stem.replace("附件", "").replace("版", "")
    return re.sub(r"[^0-9a-z\u4e00-\u9fff]+", "", stem)


def score_excel_evidence(case: Case, citations: list[dict[str, Any]]) -> dict[str, Any]:
    expected_cells = set(re.findall(r"\b[A-Z]+\d+\b", case.evidence))
    if case.qa_type == "表格比较":
        option_labels = {normalize(option) for option in case.options}
        first_cell_by_label: dict[str, str] = {}
        for label, coordinate in re.findall(
            r"([^；;=]+)=.*?\(([A-Z]+\d+)\)", case.evidence
        ):
            normalized_label = normalize(label)
            if normalized_label in option_labels:
                # Official evidence may repeat labels under a later business
                # group (for example premium income vs. claim payments). The
                # first occurrence is the comparison group named by the item.
                first_cell_by_label.setdefault(normalized_label, coordinate)
        option_cells = set(first_cell_by_label.values())
        if option_cells:
            expected_cells = option_cells
    actual_cells: set[str] = set()
    for citation in citations:
        metadata = citation.get("metadata") or {}
        if metadata.get("cell"):
            actual_cells.add(str(metadata["cell"]))
        for item in metadata.get("calculation_cells") or []:
            if item.get("cell"):
                actual_cells.add(str(item["cell"]))
        for item in metadata.get("comparison_cells") or []:
            if item.get("cell"):
                actual_cells.add(str(item["cell"]))
    return {
        "expected_cells": sorted(expected_cells), "actual_cells": sorted(actual_cells),
        "cell_recall": len(expected_cells & actual_cells) / len(expected_cells) if expected_cells else 0.0,
    }


def score_text_evidence(case: Case, citations: list[dict[str, Any]]) -> dict[str, Any]:
    """Measure whether each official fact is present in the cited excerpts.

    Character bigrams are robust to whitespace and punctuation while remaining
    strict enough that citing an unrelated paragraph from the correct file does
    not count as evidence coverage.
    """

    aspects = [value for value in re.split(r"[；;]", case.evidence) if normalize_evidence(value)]
    excerpts = [normalize_evidence(item.get("excerpt")) for item in citations]
    recalls: list[float] = []
    for aspect in aspects:
        expected = _character_ngrams(normalize_evidence(aspect), 2)
        if not expected:
            recalls.append(0.0)
            continue
        best = max(
            (len(expected & _character_ngrams(excerpt, 2)) / len(expected) for excerpt in excerpts),
            default=0.0,
        )
        recalls.append(best)
    covered = sum(recall >= 0.65 for recall in recalls)
    return {
        "aspect_count": len(aspects),
        "covered_aspect_count": covered,
        "aspect_coverage": covered / len(aspects) if aspects else 0.0,
        "evidence_bigram_recall": sum(recalls) / len(recalls) if recalls else 0.0,
        "aspect_recalls": [round(value, 4) for value in recalls],
    }


def summarize(results: list[dict[str, Any]]) -> dict[str, Any]:
    completed = [item for item in results if not item["error"]]
    timings = [float(item["elapsed_ms"]) for item in completed]
    summary: dict[str, Any] = {
        "case_count": len(results), "completed_count": len(completed),
        "error_count": len(results) - len(completed),
        "answer_accuracy": mean(item["answer_correct"] for item in completed),
        "source_hit_rate": mean(item["source_hit"] for item in completed if item["expected"] != "REFUSE"),
        "elapsed_ms": timing_summary(timings),
        # This is only the first case observed by the evaluator. A real cold-start
        # metric requires a fresh container plus a text request and is collected by
        # the dedicated performance benchmark.
        "first_case_ms": timings[0] if timings else 0.0,
        "cold_start_ms": None,
        "cold_start_definition": "not_measured_use_fresh_container_text_benchmark",
        "warm_elapsed_ms": timing_summary(timings[1:] if len(timings) > 1 else timings),
    }
    by_slice: dict[str, Any] = {}
    for field in ("source_type", "qa_type", "difficulty", "answer_type", "generation_status"):
        values: dict[str, Any] = {}
        for value in sorted({str(item[field]) for item in completed}):
            subset = [item for item in completed if str(item[field]) == value]
            values[value] = {
                "count": len(subset),
                "accuracy": mean(item["answer_correct"] for item in subset),
                "elapsed_ms": timing_summary([float(item["elapsed_ms"]) for item in subset]),
            }
        by_slice[field] = values
    summary["by_slice"] = by_slice
    excel = [item for item in completed if item.get("excel_evidence")]
    summary["excel_cell_recall"] = mean(item["excel_evidence"]["cell_recall"] for item in excel)
    text = [item for item in completed if item.get("text_evidence")]
    summary["text_evidence_coverage_rate"] = mean(
        item["text_evidence"]["aspect_coverage"] for item in text
    )
    summary["text_evidence_bigram_recall"] = mean(
        item["text_evidence"]["evidence_bigram_recall"] for item in text
    )
    grounded = [item for item in completed if item.get("grounding_validation")]
    summary["grounding_pass_rate"] = mean(
        bool(item["grounding_validation"].get("passed")) for item in grounded
    )
    summary["key_entity_error_rate"] = mean(
        bool(item["grounding_validation"].get("unsupported_entities"))
        or bool(item["grounding_validation"].get("unsupported_claim_entities"))
        or bool(item["grounding_validation"].get("invalid_citation_ids"))
        or bool(item["grounding_validation"].get("invalid_answer_citation_ids"))
        or bool(item["grounding_validation"].get("missing_inline_citation_ids"))
        for item in grounded
    )
    ood = [item for item in completed if item.get("expected") == "REFUSE"]
    summary["ood_refusal_rate"] = mean(bool(item.get("refused")) for item in ood)
    summary["citation_coverage_rate"] = mean(
        int(item.get("citation_count") or 0) > 0
        for item in completed
        if item.get("expected") != "REFUSE"
    )
    summary["failure_stage_counts"] = dict(Counter(item.get("error_type") or "success" for item in results))
    summary["generation_status_counts"] = dict(
        Counter(str(item.get("generation_status") or "unknown") for item in completed)
    )
    summary["answer_type_counts"] = dict(
        Counter(str(item.get("answer_type") or "unknown") for item in completed)
    )
    summary["incorrect_generation_status_counts"] = dict(
        Counter(
            str(item.get("generation_status") or "unknown")
            for item in completed
            if not item.get("answer_correct")
        )
    )
    manifested = [item["manifest_evidence"] for item in completed if item.get("manifest_evidence")]
    reviewed = [item for item in manifested if item.get("review_status") == "reviewed"]
    seeded = [
        item
        for item in manifested
        if item.get("review_status") in {"reviewed", "seeded"}
    ]
    summary["evidence_manifest"] = {
        "case_count": len(manifested),
        "reviewed_case_count": len(reviewed),
        "seeded_case_count": len(seeded),
        "reviewed_final_recall": _mean_optional(item.get("final_recall") for item in reviewed),
        "reviewed_rerank_retention": _mean_optional(
            item.get("rerank_retention") for item in reviewed
        ),
        "reviewed_citation_accuracy": _mean_optional(
            item.get("citation_accuracy") for item in reviewed
        ),
        "reviewed_citation_completeness": _mean_optional(
            item.get("citation_completeness") for item in reviewed
        ),
        "seeded_final_recall": _mean_optional(item.get("final_recall") for item in seeded),
    }
    return summary


def timing_summary(values: list[float]) -> dict[str, float]:
    if not values:
        return {"avg": 0.0, "p50": 0.0, "p95": 0.0, "max": 0.0}
    ordered = sorted(values)
    return {
        "avg": round(statistics.mean(values), 2),
        "p50": round(statistics.median(values), 2),
        "p95": round(ordered[min(len(ordered) - 1, math.ceil(len(ordered) * .95) - 1)], 2),
        "max": round(max(values), 2),
    }


def mean(values: Any) -> float:
    items = [float(value) for value in values]
    return round(sum(items) / len(items), 4) if items else 0.0


def _mean_optional(values: Any) -> float | None:
    items = [float(value) for value in values if value is not None]
    return round(sum(items) / len(items), 4) if items else None


def write_split_manifest(path: Path, cases: list[Case], ood: list[Case]) -> None:
    payload = {"seed": SPLIT_SEED, "dev": [c.id for c in cases if c.split == "dev"], "holdout": [c.id for c in cases if c.split == "holdout"], "ood": [c.id for c in ood]}
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def print_split_summary(cases: list[Case], ood: list[Case]) -> None:
    print(Counter(case.split for case in cases), "ood", len(ood))


def _load_checkpoint(path: Path) -> dict[str, dict[str, Any]]:
    if not path.exists():
        return {}
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}
    return payload if isinstance(payload, dict) else {}


def _write_checkpoint(path: Path, results: dict[str, dict[str, Any]]) -> None:
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def render_report(artifact: dict[str, Any]) -> str:
    summary = artifact["summary"]
    return "\n".join([
        "# ReguMate 官方 QA 端到端评测", "",
        f"- 评测集：{artifact['run']['split']}", f"- 题目数：{summary['case_count']}",
        f"- 答案准确率：{summary['answer_accuracy']:.2%}",
        f"- 来源命中率：{summary['source_hit_rate']:.2%}",
        f"- Excel 单元格召回：{summary['excel_cell_recall']:.2%}",
        f"- Word/PDF 标准证据覆盖率：{summary['text_evidence_coverage_rate']:.2%}",
        f"- Word/PDF 标准证据二元组召回：{summary['text_evidence_bigram_recall']:.2%}",
        f"- 引用覆盖率：{summary['citation_coverage_rate']:.2%}",
        f"- Grounding 校验通过率：{summary['grounding_pass_rate']:.2%}",
        f"- 关键实体错误率：{summary['key_entity_error_rate']:.2%}",
        f"- 无答案题拒答率：{summary['ood_refusal_rate']:.2%}",
        f"- P95 响应时间：{summary['elapsed_ms']['p95']:.2f} ms", "",
        "## 分项结果", "", "```json", json.dumps(summary["by_slice"], ensure_ascii=False, indent=2), "```", "",
        "## 失败样例", "",
        *[f"- {item['id']}：expected={item['expected']} predicted={item['predicted']} error={item['error']}" for item in artifact["results"] if not item["answer_correct"]][:30],
    ])


def normalize(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).lower()


def normalize_evidence(value: Any) -> str:
    return re.sub(r"[^0-9a-z\u4e00-\u9fff]+", "", str(value or "").lower())


def _character_ngrams(value: str, size: int) -> set[str]:
    if len(value) < size:
        return {value} if value else set()
    return {value[index : index + size] for index in range(len(value) - size + 1)}


if __name__ == "__main__":
    raise SystemExit(main())
