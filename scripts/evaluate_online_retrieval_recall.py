from __future__ import annotations

import argparse
from datetime import datetime, timezone
import json
from pathlib import Path
from time import perf_counter, sleep
from typing import Any
import urllib.request
import urllib.error

from openpyxl import load_workbook

from evaluation_common import (
    DEFAULT_OUTPUT_DIR,
    describe_numbers,
    project_path,
    read_json,
    rel,
    safe_div,
    sha256_file,
    write_json,
    write_jsonl,
)


DEFAULT_QA = Path("data/contest_dataset/QA数据.xlsx")
DEFAULT_EVIDENCE_MANIFEST = Path("data/evaluation/performance_experiments/evidence_manifest_seed.json")


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Run online retrieval Recall@K through /api/qa/retrieve and score against evidence manifest."
    )
    parser.add_argument("--qa", type=Path, default=DEFAULT_QA)
    parser.add_argument("--evidence-manifest", type=Path, default=DEFAULT_EVIDENCE_MANIFEST)
    parser.add_argument("--base-url", default="http://127.0.0.1:8000")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR / "online_retrieval")
    parser.add_argument("--split", choices=["all", "dev", "holdout"], default="all")
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--start", type=int, default=0)
    parser.add_argument("--timeout", type=float, default=240.0)
    parser.add_argument("--request-delay", type=float, default=0.05)
    parser.add_argument("--resume", action="store_true")
    args = parser.parse_args()

    output_dir = project_path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    cases = _read_cases(project_path(args.qa))
    if args.split != "all":
        cases = [case for case in cases if case.get("split") == args.split]
    cases = cases[args.start :]
    if args.limit > 0:
        cases = cases[: args.limit]
    manifest_cases = _manifest_cases(project_path(args.evidence_manifest))
    checkpoint_path = output_dir / "online_retrieval_checkpoint.json"
    prior = _read_json_if_exists(checkpoint_path) if args.resume else {}
    results: list[dict[str, Any]] = []
    for idx, case in enumerate(cases, start=1):
        case_id = str(case["id"])
        if case_id in prior:
            result = prior[case_id]
        else:
            result = _evaluate_case(case, manifest_cases.get(case_id), args.base_url, args.timeout)
            prior[case_id] = result
            write_json(checkpoint_path, prior)
            if args.request_delay > 0 and idx < len(cases):
                sleep(args.request_delay)
        results.append(result)
        print(
            f"[{idx}/{len(cases)}] {case_id} evaluable={result.get('evaluable')} "
            f"r@5={result.get('case_hit_at_5')} elapsed={result.get('elapsed_s')}s",
            flush=True,
        )

    health = _get_json(f"{args.base_url.rstrip('/')}/api/health/ready", timeout=15)
    artifact = {
        "run": {
            "created_at": datetime.now(timezone.utc).isoformat(),
            "base_url": args.base_url,
            "qa_file": rel(project_path(args.qa)),
            "qa_sha256": sha256_file(project_path(args.qa)),
            "evidence_manifest": rel(project_path(args.evidence_manifest)),
            "evidence_manifest_sha256": sha256_file(project_path(args.evidence_manifest)),
            "split": args.split,
            "case_count": len(cases),
            "limit": args.limit,
            "start": args.start,
            "health": _compact_health(health),
        },
        "summary": _summarize(results),
        "results": results,
    }
    write_json(output_dir / "online_retrieval_results.json", artifact)
    write_jsonl(output_dir / "online_retrieval_details.jsonl", results)
    write_json(DEFAULT_OUTPUT_DIR / "online_retrieval_metrics.json", artifact)
    print(json.dumps(artifact["summary"], ensure_ascii=False, indent=2))
    return 0


def _read_cases(path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = [str(cell.value or "") for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    cases = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        record = dict(zip(headers, row, strict=False))
        cases.append(
            {
                "id": str(record.get("id") or ""),
                "question": str(record.get("question") or ""),
                "source_type": str(record.get("source_type") or ""),
                "qa_type": str(record.get("qa_type") or ""),
                "difficulty": str(record.get("difficulty_cn") or record.get("difficulty") or ""),
                "split": str(record.get("split") or ""),
            }
        )
    return cases


def _manifest_cases(path: Path) -> dict[str, dict[str, Any]]:
    data = read_json(path)
    return {
        str(case.get("case_id")): case
        for case in data.get("cases", [])
        if isinstance(case, dict) and case.get("case_id")
    }


def _evaluate_case(
    case: dict[str, Any],
    manifest_case: dict[str, Any] | None,
    base_url: str,
    timeout: float,
) -> dict[str, Any]:
    started = perf_counter()
    try:
        package = _post_json(
            f"{base_url.rstrip('/')}/api/qa/retrieve",
            {"question": case["question"]},
            timeout=timeout,
        )
        elapsed_s = round(perf_counter() - started, 3)
    except Exception as exc:
        return {
            **case,
            "elapsed_s": round(perf_counter() - started, 3),
            "evaluable": False,
            "error": str(exc),
        }
    retrieval_summary = package.get("retrieval_summary") or {}
    ranked_ids = _ranked_chunk_ids(retrieval_summary, package)
    selected_ids = _selected_chunk_ids(retrieval_summary, package)
    required = [
        aspect
        for aspect in (manifest_case or {}).get("required_aspects", [])
        if isinstance(aspect, dict) and aspect.get("acceptable_chunk_ids")
    ]
    needs_review = bool(manifest_case and manifest_case.get("review_status") == "needs_review")
    if not required:
        return {
            **case,
            "elapsed_s": elapsed_s,
            "evaluable": False,
            "review_status": (manifest_case or {}).get("review_status"),
            "needs_review": needs_review,
            "reason": "no_acceptable_chunk_ids",
            "ranked_candidate_count": len(ranked_ids),
            "selected_chunk_count": len(selected_ids),
            "retrieval_summary_compact": _compact_retrieval(retrieval_summary),
        }
    aspect_results = []
    case_hit = {1: True, 3: True, 5: True}
    prompt_hit_all = True
    best_ranks = []
    for aspect in required:
        acceptable = {str(item) for item in aspect.get("acceptable_chunk_ids") or []}
        rank = _first_hit_rank(ranked_ids, acceptable)
        prompt_hit = bool(acceptable.intersection(selected_ids))
        if rank is not None:
            best_ranks.append(rank)
        for k in case_hit:
            if rank is None or rank > k:
                case_hit[k] = False
        if not prompt_hit:
            prompt_hit_all = False
        aspect_results.append(
            {
                "aspect_id": aspect.get("aspect_id"),
                "acceptable_chunk_ids": sorted(acceptable),
                "best_hit_rank": rank,
                "hit_at_1": rank is not None and rank <= 1,
                "hit_at_3": rank is not None and rank <= 3,
                "hit_at_5": rank is not None and rank <= 5,
                "selected_for_prompt_hit": prompt_hit,
            }
        )
    return {
        **case,
        "elapsed_s": elapsed_s,
        "evaluable": True,
        "review_status": (manifest_case or {}).get("review_status"),
        "required_aspect_count": len(required),
        "ranked_candidate_count": len(ranked_ids),
        "selected_chunk_count": len(selected_ids),
        "case_hit_at_1": case_hit[1],
        "case_hit_at_3": case_hit[3],
        "case_hit_at_5": case_hit[5],
        "case_prompt_hit_all": prompt_hit_all,
        "best_hit_rank": min(best_ranks) if best_ranks else None,
        "retrieval_summary_compact": _compact_retrieval(retrieval_summary),
        "aspects": aspect_results,
    }


def _ranked_chunk_ids(retrieval_summary: dict[str, Any], package: dict[str, Any]) -> list[str]:
    ids: list[str] = []
    for aspect in retrieval_summary.get("aspect_retrievals") or []:
        for chunk in aspect.get("retrieved_chunks") or []:
            chunk_id = chunk.get("chunk_id") if isinstance(chunk, dict) else None
            if chunk_id and str(chunk_id) not in ids:
                ids.append(str(chunk_id))
    if not ids:
        for chunk in package.get("context_chunks") or []:
            chunk_id = chunk.get("chunk_id") if isinstance(chunk, dict) else None
            if chunk_id and str(chunk_id) not in ids:
                ids.append(str(chunk_id))
    return ids


def _selected_chunk_ids(retrieval_summary: dict[str, Any], package: dict[str, Any]) -> set[str]:
    selected: set[str] = set()
    prompt = retrieval_summary.get("prompt_selection") or {}
    for chunk_id in prompt.get("final_prompt_chunk_ids") or []:
        selected.add(str(chunk_id))
    for chunk in package.get("context_chunks") or []:
        if isinstance(chunk, dict) and chunk.get("chunk_id"):
            selected.add(str(chunk["chunk_id"]))
    return selected


def _first_hit_rank(ranked_ids: list[str], acceptable: set[str]) -> int | None:
    for idx, chunk_id in enumerate(ranked_ids, start=1):
        if chunk_id in acceptable:
            return idx
    return None


def _summarize(results: list[dict[str, Any]]) -> dict[str, Any]:
    evaluable = [row for row in results if row.get("evaluable") is True]
    aspect_rows = [aspect for row in evaluable for aspect in row.get("aspects", [])]
    needs_review = [row for row in results if row.get("needs_review") or row.get("reason") == "no_acceptable_chunk_ids"]
    return {
        "case_count": len(results),
        "completed_count": sum(1 for row in results if not row.get("error")),
        "error_count": sum(1 for row in results if row.get("error")),
        "evaluable_case_count": len(evaluable),
        "needs_review_or_no_acceptable_count": len(needs_review),
        "aspect_count": len(aspect_rows),
        "case_recall_at_1": _rate(evaluable, "case_hit_at_1"),
        "case_recall_at_3": _rate(evaluable, "case_hit_at_3"),
        "case_recall_at_5": _rate(evaluable, "case_hit_at_5"),
        "aspect_recall_at_1": _aspect_rate(aspect_rows, "hit_at_1"),
        "aspect_recall_at_3": _aspect_rate(aspect_rows, "hit_at_3"),
        "aspect_recall_at_5": _aspect_rate(aspect_rows, "hit_at_5"),
        "prompt_context_case_recall": _rate(evaluable, "case_prompt_hit_all"),
        "elapsed_s": describe_numbers([float(row["elapsed_s"]) for row in results if row.get("elapsed_s") is not None]),
        "ranked_candidate_count": describe_numbers([
            float(row["ranked_candidate_count"]) for row in results if isinstance(row.get("ranked_candidate_count"), int)
        ]),
        "metric_formula": {
            "case_recall_at_k": "全部标准 aspect 均在在线检索返回候选 Top-K 内命中 acceptable_chunk_ids 的题数 / 可评估题数",
            "aspect_recall_at_k": "在线检索返回候选 Top-K 内命中 acceptable_chunk_ids 的标准 aspect 数 / 可评估 aspect 数",
        },
        "limitations": [
            "This script evaluates the system online retrieval endpoint /api/qa/retrieve.",
            "It records current backend health and collection settings; CPU/GPU performance depends on the running service.",
            "Cases without acceptable_chunk_ids are listed separately and are not counted in Recall@K denominators.",
        ],
    }


def _compact_retrieval(summary: dict[str, Any]) -> dict[str, Any]:
    return {
        "top_k": summary.get("top_k"),
        "used_chunks": summary.get("used_chunks"),
        "query_count": summary.get("query_count"),
        "raw_candidate_count": summary.get("raw_candidate_count"),
        "candidate_count": summary.get("candidate_count"),
        "rerank_input_count": summary.get("rerank_input_count"),
        "reranked_count": summary.get("reranked_count"),
        "filtered_count": summary.get("filtered_count"),
        "has_sufficient_context": summary.get("has_sufficient_context"),
        "retrieval_covered_aspect_count": summary.get("retrieval_covered_aspect_count"),
        "prompt_covered_aspect_count": summary.get("prompt_covered_aspect_count"),
        "timings_ms": summary.get("timings_ms") or {},
    }


def _compact_health(health: dict[str, Any]) -> dict[str, Any]:
    return {
        "ready": health.get("ready"),
        "build_id": health.get("build_id"),
        "qdrant_collection": health.get("qdrant_collection"),
        "qdrant_ready": health.get("qdrant_ready"),
        "model_device": health.get("model_device"),
        "performance": {
            key: (health.get("performance") or {}).get(key)
            for key in [
                "selected_mode",
                "backend",
                "embedding_batch_size",
                "rerank_batch_size",
                "rerank_max_length",
                "torch_num_threads",
            ]
        },
    }


def _rate(rows: list[dict[str, Any]], key: str) -> float | None:
    return safe_div(sum(1 for row in rows if row.get(key) is True), len(rows))


def _aspect_rate(rows: list[dict[str, Any]], key: str) -> float | None:
    return safe_div(sum(1 for row in rows if row.get(key) is True), len(rows))


def _read_json_if_exists(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}


def _post_json(url: str, body: dict[str, Any], timeout: float) -> dict[str, Any]:
    data = json.dumps(body, ensure_ascii=False).encode("utf-8")
    request = urllib.request.Request(
        url,
        data=data,
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    with urllib.request.urlopen(request, timeout=timeout) as response:
        return json.loads(response.read().decode("utf-8"))


def _get_json(url: str, timeout: float) -> dict[str, Any]:
    try:
        with urllib.request.urlopen(url, timeout=timeout) as response:
            return json.loads(response.read().decode("utf-8"))
    except urllib.error.URLError:
        return {}


if __name__ == "__main__":
    raise SystemExit(main())
