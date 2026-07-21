from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import shutil
import sys
from time import perf_counter
from statistics import mean, median

from sqlalchemy import select

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))


def _bootstrap_environment() -> None:
    parser = argparse.ArgumentParser(add_help=False)
    parser.add_argument("--data-dir", type=Path)
    parser.add_argument("--collection")
    known, _ = parser.parse_known_args()
    if known.data_dir:
        os.environ["REGUMATE_DATA_DIR"] = str(known.data_dir.resolve())
    if known.collection:
        os.environ["QDRANT_COLLECTION"] = known.collection
        os.environ["QDRANT_AUTO_CREATE_COLLECTION"] = "true"


_bootstrap_environment()

from backend.app.core.config import DOCUMENT_DIR, INDEX_VERSION, SUPPORTED_DOCUMENT_EXTENSIONS  # noqa: E402
from backend.app.core.database import SessionLocal, init_db  # noqa: E402
from backend.app.models.document import Document, DocumentChunk, SpreadsheetCell  # noqa: E402
from backend.app.services.chunk_service import build_chunks_from_parsed  # noqa: E402
from backend.app.services.document_indexing_service import DocumentIndexingError, index_document  # noqa: E402
from backend.app.services.document_parser import DocumentParseError, parse_document  # noqa: E402
from backend.app.services.spreadsheet_cell_index_service import rebuild_spreadsheet_cell_index  # noqa: E402
from backend.app.services.qdrant_admin_service import (  # noqa: E402
    collection_status,
    delete_collection_if_exists,
    promote_collection_alias,
)
from backend.app.services.performance_metrics import (  # noqa: E402
    measure,
    resource_metrics_snapshot,
    start_resource_sampling,
    stop_resource_sampling,
    timing_metrics_snapshot,
)
from backend.app.services.vector_store_service import count_document_vectors, delete_document_vectors  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


DEFAULT_SOURCE = PROJECT_ROOT / "data" / "contest_dataset" / "dataset" / "nfra_page_attachments_500"


def main() -> int:
    parser = argparse.ArgumentParser(description="Idempotently ingest the official ReguMate contest corpus.")
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--source-manifest", type=Path, help="Map ASCII staged names back to official filenames.")
    parser.add_argument("--parse-only", action="store_true", help="Build SQLite metadata/cells without Qdrant vectors.")
    parser.add_argument("--index-existing", action="store_true", help="Index already parsed SQLite chunks without reparsing source files.")
    parser.add_argument("--retry-failed", action="store_true")
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--qa-only", action="store_true", help="Only ingest files referenced by QA数据.xlsx.")
    parser.add_argument("--extensions", default="", help="Comma-separated suffixes, for example .xls,.xlsx")
    parser.add_argument("--data-dir", type=Path, help="Isolated SQLite/document runtime directory.")
    parser.add_argument("--collection", default=os.getenv("QDRANT_COLLECTION", "regumate_contest_v3_build"))
    parser.add_argument("--alias", default="regumate_contest_v3")
    parser.add_argument("--promote-alias", action="store_true")
    parser.add_argument("--force-rebuild", action="store_true")
    parser.add_argument("--reset-collection", action="store_true")
    parser.add_argument(
        "--manifest",
        type=Path,
        default=PROJECT_ROOT / "data" / "evaluation" / "final" / "ingest_manifest.json",
    )
    args = parser.parse_args()
    if args.parse_only and args.promote_alias:
        parser.error("--parse-only cannot be combined with --promote-alias")
    if args.parse_only and args.index_existing:
        parser.error("--parse-only cannot be combined with --index-existing")
    if args.reset_collection:
        delete_collection_if_exists(args.collection)
    run_started = perf_counter()
    start_resource_sampling(interval_seconds=1.0)
    init_db()
    source_names = _source_name_mapping(args.source_manifest)
    files = sorted(path for path in args.source.rglob("*") if path.is_file() and path.suffix.lower() in SUPPORTED_DOCUMENT_EXTENSIONS)
    if source_names:
        files = [path for path in files if path.name in source_names]
    if args.qa_only:
        labels = _qa_file_labels(PROJECT_ROOT / "data" / "contest_dataset" / "QA数据.xlsx")
        files = [path for path in files if any(label in source_names.get(path.name, path.name) for label in labels)]
    if args.extensions:
        suffixes = {value.strip().lower() for value in args.extensions.split(",") if value.strip()}
        files = [path for path in files if path.suffix.lower() in suffixes]
    if args.limit:
        files = files[: args.limit]
    failures: list[dict[str, str]] = []
    records: list[dict[str, object]] = []
    completed = skipped = 0
    with SessionLocal() as db:
        for position, source in enumerate(files, start=1):
            started = perf_counter()
            source_name = source_names.get(source.name, source.name)
            document_id = _document_id_from_name(source_name)
            existing = db.scalar(select(Document).where(Document.document_id == document_id))
            if args.index_existing:
                if existing is None:
                    failures.append({"file": str(source), "error": "document has not been parsed"})
                    records.append(
                        {
                            "file": str(source), "sha256": _file_sha256(source),
                            "extension": source.suffix.lower(), "status": "failed",
                            "error": "document has not been parsed",
                            "elapsed_ms": round((perf_counter() - started) * 1000, 2),
                        }
                    )
                    _write_manifest(args.manifest, args, files, records, failures)
                    continue
                if existing.status != "indexed":
                    try:
                        index_document(db, existing)
                        completed += 1
                    except DocumentIndexingError as exc:
                        db.rollback()
                        failures.append({"file": str(source), "error": str(exc)})
                        records.append(
                            {
                                "file": str(source), "sha256": _file_sha256(source),
                                "extension": source.suffix.lower(), "status": "failed",
                                "error": str(exc),
                                "elapsed_ms": round((perf_counter() - started) * 1000, 2),
                            }
                        )
                        _write_manifest(args.manifest, args, files, records, failures)
                        continue
                else:
                    skipped += 1
                records.append(_existing_record(db, source, source_name, existing, perf_counter() - started))
                _write_manifest(args.manifest, args, files, records, failures)
                print(
                    f"[{position}/{len(files)}] indexed {source.name} vectors={records[-1]['vector_count']} "
                    f"elapsed={perf_counter()-started:.2f}s",
                    flush=True,
                )
                continue
            if existing and args.force_rebuild:
                if not args.parse_only:
                    delete_document_vectors(document_id)
                db.delete(existing)
                db.commit()
                existing = None
            if existing and existing.status == "indexed":
                skipped += 1
                records.append(_existing_record(db, source, source_name, existing, perf_counter() - started))
                _write_manifest(args.manifest, args, files, records, failures)
                continue
            if existing and not args.retry_failed:
                skipped += 1
                continue
            try:
                if existing:
                    if not args.parse_only:
                        delete_document_vectors(document_id)
                    db.delete(existing)
                    db.commit()
                stored = DOCUMENT_DIR / f"{document_id}{source.suffix.lower()}"
                stored.parent.mkdir(parents=True, exist_ok=True)
                if not stored.exists():
                    shutil.copy2(source, stored)
                with measure("index.parse"):
                    parsed = parse_document(stored, source_name=source_name)
                with measure("index.chunk_build"):
                    drafts = build_chunks_from_parsed(
                        document_id=document_id,
                        parsed=parsed,
                        source_file=source_name,
                    )
                if not drafts:
                    raise DocumentParseError("文档没有可入库内容")
                document = Document(
                    document_id=document_id, filename=source_name, content_type=None,
                    file_type=source.suffix.lower().lstrip("."), size=source.stat().st_size,
                    storage_path=str(stored), document_metadata=json.dumps(parsed.metadata or {}, ensure_ascii=False),
                    status="uploaded", index_version=INDEX_VERSION, chunk_count=len(drafts),
                )
                with measure("index.sqlite_chunk_persist"):
                    db.add(document)
                    chunks: list[DocumentChunk] = []
                    for draft in drafts:
                        chunk = DocumentChunk(
                            chunk_id=draft.chunk_id, document_id=document_id, text=draft.text,
                            embedding_text=draft.embedding_text,
                            chunk_metadata=json.dumps(draft.metadata or {}, ensure_ascii=False),
                            token_count=draft.token_count, index_status="uploaded", index_version=INDEX_VERSION,
                            title=draft.title, section_title=draft.section_title, page_number=draft.page_number,
                            source_file=source_name,
                        )
                        chunks.append(chunk)
                        db.add(chunk)
                    db.flush()
                    rebuild_spreadsheet_cell_index(db, document, chunks)
                    if args.parse_only and source.suffix.lower() in {".xls", ".xlsx"}:
                        document.status = "table_indexed"
                        for chunk in chunks:
                            chunk.index_status = "table_indexed"
                    db.commit()
                if not args.parse_only:
                    index_document(db, document)
                completed += 1
                records.append(_document_record(db, source, source_name, document, parsed.metadata, perf_counter() - started))
                _write_manifest(args.manifest, args, files, records, failures)
                print(f"[{position}/{len(files)}] ok {source.name} chunks={len(chunks)} elapsed={perf_counter()-started:.2f}s", flush=True)
            except (DocumentParseError, DocumentIndexingError, OSError, ValueError) as exc:
                db.rollback()
                failures.append({"file": str(source), "error": str(exc)})
                records.append(
                    {
                        "file": str(source),
                        "sha256": _file_sha256(source),
                        "extension": source.suffix.lower(),
                        "status": "failed",
                        "error": str(exc),
                        "elapsed_ms": round((perf_counter() - started) * 1000, 2),
                    }
                )
                _write_manifest(args.manifest, args, files, records, failures)
                print(f"[{position}/{len(files)}] failed {source.name}: {exc}", flush=True)
    output = PROJECT_ROOT / "data" / "evaluation" / "contest_ingest_failures.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(failures, ensure_ascii=False, indent=2), encoding="utf-8")
    stop_resource_sampling()
    _write_manifest(
        args.manifest,
        args,
        files,
        records,
        failures,
        completed=True,
        run_elapsed_ms=(perf_counter() - run_started) * 1000,
    )
    if args.promote_alias:
        if failures or len(files) != 500:
            print("alias not promoted: full 500-file ingestion without failures is required")
            return 1
        status = collection_status(args.collection)
        expected_vectors = sum(int(record.get("chunk_count") or 0) for record in records)
        if status.points_count != expected_vectors:
            print(f"alias not promoted: vectors expected={expected_vectors} actual={status.points_count}")
            return 1
        promote_collection_alias(collection_name=args.collection, alias_name=args.alias)
        print(f"promoted alias {args.alias} -> {args.collection}")
    print(f"completed={completed} skipped={skipped} failed={len(failures)} failures={output}")
    return 1 if failures else 0


def _document_id(path: Path) -> str:
    return _document_id_from_name(path.name)


def _document_id_from_name(name: str) -> str:
    digest = hashlib.sha256(name.encode("utf-8")).hexdigest()[:16].upper()
    return f"DOC-C-{digest}"


def _qa_file_labels(path: Path) -> set[str]:
    worksheet = load_workbook(path, read_only=True, data_only=True).active
    headers = [str(cell.value or "") for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
    file_index = headers.index("file_label")
    return {Path(str(row[file_index])).stem for row in worksheet.iter_rows(min_row=2, values_only=True)}


def _document_record(db, source: Path, source_name: str, document: Document, metadata: dict, elapsed: float) -> dict[str, object]:
    cell_count = db.scalar(
        select(SpreadsheetCell.id).where(SpreadsheetCell.document_id == document.document_id).limit(1)
    )
    total_cells = len(
        db.scalars(select(SpreadsheetCell.id).where(SpreadsheetCell.document_id == document.document_id)).all()
    ) if cell_count is not None else 0
    vector_count = count_document_vectors(document.document_id) if document.status == "indexed" else 0
    return {
        "file": str(source),
        "original_name": source_name,
        "sha256": _file_sha256(source),
        "extension": source.suffix.lower(),
        "status": document.status,
        "document_id": document.document_id,
        "parser_backend": metadata.get("parser_backend") or metadata.get("loader_name"),
        "loader_name": metadata.get("loader_name"),
        "degraded": bool(metadata.get("degraded", False)),
        "degradation_reason": metadata.get("degradation_reason"),
        "block_count": int(metadata.get("block_count") or 0),
        "chunk_count": document.chunk_count,
        "cell_count": total_cells,
        "vector_count": vector_count,
        "elapsed_ms": round(elapsed * 1000, 2),
        "error": None,
    }


def _existing_record(db, source: Path, source_name: str, document: Document, elapsed: float) -> dict[str, object]:
    try:
        metadata = json.loads(document.document_metadata or "{}")
    except json.JSONDecodeError:
        metadata = {}
    record = _document_record(db, source, source_name, document, metadata, elapsed)
    record["resumed"] = True
    return record


def _write_manifest(
    path: Path,
    args,
    files: list[Path],
    records: list[dict[str, object]],
    failures: list[dict[str, str]],
    *,
    completed: bool = False,
    run_elapsed_ms: float | None = None,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    success_statuses = {"indexed", "table_indexed", "uploaded"} if args.parse_only else {"indexed"}
    payload = {
        "run": {
            "source": str(args.source),
            "data_dir": str(args.data_dir or os.getenv("REGUMATE_DATA_DIR", "")),
            "collection": args.collection,
            "alias": args.alias,
            "file_count": len(files),
            "completed": completed,
        },
        "summary": {
            "processed": len(records),
            "success": sum(
                record.get("status") in success_statuses
                for record in records
            ),
            "degraded": sum(bool(record.get("degraded")) for record in records),
            "failed": len(failures),
            "chunks": sum(int(record.get("chunk_count") or 0) for record in records),
            "cells": sum(int(record.get("cell_count") or 0) for record in records),
            "vectors": sum(int(record.get("vector_count") or 0) for record in records),
            "by_extension": _group_manifest_records(records, "extension", success_statuses),
            "by_parser": _group_manifest_records(records, "parser_backend", success_statuses),
        },
        "documents": records,
        "failures": failures,
    }
    if completed:
        elapsed_values = sorted(float(record.get("elapsed_ms") or 0.0) for record in records)
        payload["performance"] = {
            "total_elapsed_ms": round(run_elapsed_ms or sum(elapsed_values), 2),
            "per_document_ms": _distribution(elapsed_values),
            "stages": timing_metrics_snapshot(),
            "resources": resource_metrics_snapshot(),
        }
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def _distribution(values: list[float]) -> dict[str, float | int]:
    if not values:
        return {"count": 0, "average": 0.0, "p50": 0.0, "p95": 0.0, "maximum": 0.0}
    p95_index = min(len(values) - 1, max(0, int(len(values) * 0.95 + 0.999999) - 1))
    return {
        "count": len(values),
        "average": round(mean(values), 2),
        "p50": round(median(values), 2),
        "p95": round(values[p95_index], 2),
        "maximum": round(values[-1], 2),
    }


def _group_manifest_records(
    records: list[dict[str, object]], field: str, success_statuses: set[str]
) -> dict[str, dict[str, int]]:
    grouped: dict[str, dict[str, int]] = {}
    for record in records:
        key = str(record.get(field) or "unknown")
        bucket = grouped.setdefault(
            key,
            {"processed": 0, "success": 0, "degraded": 0, "failed": 0, "chunks": 0, "cells": 0},
        )
        bucket["processed"] += 1
        bucket["success"] += int(record.get("status") in success_statuses)
        bucket["degraded"] += int(bool(record.get("degraded")))
        bucket["failed"] += int(record.get("status") == "failed")
        bucket["chunks"] += int(record.get("chunk_count") or 0)
        bucket["cells"] += int(record.get("cell_count") or 0)
    return dict(sorted(grouped.items()))


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _source_name_mapping(path: Path | None) -> dict[str, str]:
    if path is None:
        return {}
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    files = payload.get("files") if isinstance(payload, dict) else None
    if not isinstance(files, list):
        raise ValueError("source manifest does not contain a files list")
    return {
        str(item["staged_name"]): str(item["original_name"])
        for item in files
        if isinstance(item, dict) and item.get("staged_name") and item.get("original_name")
    }


if __name__ == "__main__":
    raise SystemExit(main())
