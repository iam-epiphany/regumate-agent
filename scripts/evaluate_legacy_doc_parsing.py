from __future__ import annotations

import argparse
from datetime import datetime, timezone
import json
from pathlib import Path
import sys
from time import perf_counter

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from backend.app.services.document_parser import DocumentParseError, parse_document  # noqa: E402


def main() -> int:
    parser = argparse.ArgumentParser(description="Validate all official legacy .doc files.")
    parser.add_argument(
        "--source",
        type=Path,
        default=ROOT / "data" / "contest dataset" / "dataset" / "nfra_page_attachments_500",
    )
    parser.add_argument(
        "--qa",
        type=Path,
        default=ROOT / "data" / "contest dataset" / "QA数据.xlsx",
    )
    parser.add_argument("--source-manifest", type=Path)
    parser.add_argument(
        "--output",
        type=Path,
        default=ROOT / "data" / "evaluation" / "final" / "legacy_doc_report.json",
    )
    args = parser.parse_args()
    names = _source_name_mapping(args.source_manifest)
    files = sorted(
        path for path in args.source.rglob("*.doc")
        if names.get(path.name, path.name).lower().endswith(".doc")
    )
    qa_labels = _qa_labels(args.qa)
    records = []
    for position, path in enumerate(files, start=1):
        source_name = names.get(path.name, path.name)
        qa_referenced = any(Path(label).stem in Path(source_name).stem for label in qa_labels)
        record = {"file": source_name, "staged_file": path.name, "qa_referenced": qa_referenced}
        for mode, loader in (("primary", None), ("antiword_fallback", "antiword-doc")):
            started = perf_counter()
            try:
                parsed = parse_document(path, loader_name=loader, source_name=source_name)
                record[mode] = {
                    "ok": True,
                    "loader": parsed.metadata.get("loader_name"),
                    "degraded": parsed.metadata.get("degraded", False),
                    "blocks": len(parsed.blocks),
                    "characters": len(parsed.text),
                    "elapsed_ms": round((perf_counter() - started) * 1000, 2),
                    "error": None,
                }
            except DocumentParseError as exc:
                record[mode] = {
                    "ok": False,
                    "elapsed_ms": round((perf_counter() - started) * 1000, 2),
                    "error": str(exc),
                }
        records.append(record)
        print(
            f"[{position}/{len(files)}] {source_name} primary={record['primary']['ok']} "
            f"fallback={record['antiword_fallback']['ok']}",
            flush=True,
        )
    summary = {
        "file_count": len(files),
        "primary_success": sum(item["primary"]["ok"] for item in records),
        "fallback_success": sum(item["antiword_fallback"]["ok"] for item in records),
        "qa_referenced": sum(item["qa_referenced"] for item in records),
        "qa_primary_coverage": sum(item["qa_referenced"] and item["primary"]["ok"] for item in records),
        "qa_fallback_coverage": sum(item["qa_referenced"] and item["antiword_fallback"]["ok"] for item in records),
    }
    artifact = {
        "created_at": datetime.now(timezone.utc).isoformat(),
        "summary": summary,
        "documents": records,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(artifact, ensure_ascii=False, indent=2), encoding="utf-8")
    args.output.with_suffix(".md").write_text(_render(summary, records), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0 if summary["primary_success"] == len(files) else 1


def _qa_labels(path: Path) -> set[str]:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    headers = [str(cell.value or "") for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    index = headers.index("file_label")
    return {Path(str(row[index])).name for row in sheet.iter_rows(min_row=2, values_only=True)}


def _source_name_mapping(path: Path | None) -> dict[str, str]:
    if path is None:
        return {}
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    return {
        str(item["staged_name"]): str(item["original_name"])
        for item in payload.get("files", [])
        if item.get("staged_name") and item.get("original_name")
    }


def _render(summary: dict, records: list[dict]) -> str:
    primary_failures = [
        f"- {item['file']}：{item['primary'].get('error')}"
        for item in records
        if not item["primary"]["ok"]
    ]
    fallback_unavailable = [
        f"- {item['file']}：{item['antiword_fallback'].get('error')}"
        for item in records
        if not item["antiword_fallback"]["ok"]
    ]
    return "\n".join(
        [
            "# 官方旧DOC解析报告", "",
            f"- 文件数：{summary['file_count']}",
            f"- LibreOffice主路径成功：{summary['primary_success']}",
            f"- antiword降级成功：{summary['fallback_success']}",
            f"- QA关联DOC：{summary['qa_referenced']}",
            f"- QA主路径覆盖：{summary['qa_primary_coverage']}",
            f"- QA降级覆盖：{summary['qa_fallback_coverage']}", "",
            "## LibreOffice主路径失败", "", *(primary_failures or ["- 无"]), "",
            "## antiword不支持（主路径仍成功）", "", *(fallback_unavailable or ["- 无"]), "",
        ]
    )


if __name__ == "__main__":
    raise SystemExit(main())
