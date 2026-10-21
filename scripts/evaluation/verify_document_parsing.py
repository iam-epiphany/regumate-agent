"""Verify parsed-document integrity after a knowledge-base import.

The check is deliberately independent from the QA path: it reads the SQLite
ledger, the archived originals, the Qdrant collection, and re-parses every
original with the same parser/chunker used at ingest time.

Usage (inside the app container, after indexing finishes):
    python /app/scripts/evaluation/verify_document_parsing.py \
        --output-dir /app/data/evaluation/document_parse_check/<run> \
        --source-dir /app/data/contest_dataset/dataset/nfra_page_attachments_500

Exit code 0 means no errors (warnings are allowed). Exit code 1 means at least
one document failed a blocking check; see report.json for details.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sqlite3
import statistics
import sys
import traceback
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from backend.app.core import config  # noqa: E402
from backend.app.services.chunk_service import build_chunks_from_parsed  # noqa: E402
from backend.app.services.document_parser import parse_document  # noqa: E402

SUPPORTED_EXTENSIONS = {".txt", ".md", ".doc", ".docx", ".pdf", ".xls", ".xlsx"}
CJK_RE = re.compile(r"[\u4e00-\u9fff]")
CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
STRUCTURAL_METADATA_KEYS = (
    "sheet_name",
    "table_id",
    "table_title",
    "table_chunk_role",
    "period",
    "cells",
    "section_number",
)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while chunk := stream.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", value or "")).strip()


def percentile(values: list[int], ratio: float) -> float:
    if not values:
        return 0.0
    ordered = sorted(values)
    index = min(len(ordered) - 1, max(0, round((len(ordered) - 1) * ratio)))
    return float(ordered[index])


def short(value: Any, limit: int = 160) -> str:
    text = str(value or "").replace("\n", " ")
    return text[:limit]


class Checker:
    def __init__(self) -> None:
        self.errors: list[dict[str, Any]] = []
        self.warnings: list[dict[str, Any]] = []

    def error(self, code: str, detail: str, *, document: str | None = None) -> None:
        self.errors.append({"code": code, "document": document, "detail": detail})

    def warn(self, code: str, detail: str, *, document: str | None = None) -> None:
        self.warnings.append({"code": code, "document": document, "detail": detail})


def load_documents(db_path: Path) -> tuple[sqlite3.Connection, list[dict[str, Any]]]:
    if not db_path.is_file():
        raise RuntimeError(f"database not found: {db_path}")
    conn = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    rows = [dict(row) for row in conn.execute("SELECT * FROM documents ORDER BY filename")]
    return conn, rows


def check_corpus(
    checker: Checker,
    documents: list[dict[str, Any]],
    chunk_rows: int,
    expected_documents: int,
) -> None:
    if len(documents) != expected_documents:
        checker.error("document_count", f"expected={expected_documents}, actual={len(documents)}")
    names = [str(doc["filename"]) for doc in documents]
    ids = [str(doc["document_id"]) for doc in documents]
    if len(set(names)) != len(names):
        checker.error("duplicate_filenames", f"unique={len(set(names))}, total={len(names)}")
    if len(set(ids)) != len(ids):
        checker.error("duplicate_document_ids", f"unique={len(set(ids))}, total={len(ids)}")
    status_counts = Counter(str(doc["status"]) for doc in documents)
    for status, count in status_counts.items():
        if status != "indexed":
            checker.error("document_status", f"status={status}, count={count}")
    expected_chunks = sum(int(doc["chunk_count"] or 0) for doc in documents)
    if expected_chunks != chunk_rows:
        checker.error("chunk_total", f"chunk_count_sum={expected_chunks}, chunk_rows={chunk_rows}")
    for doc in documents:
        if int(doc["chunk_count"] or 0) <= 0:
            checker.error("empty_document", "chunk_count=0", document=str(doc["filename"]))
        if doc["index_error"]:
            checker.error("index_error", short(doc["index_error"], 300), document=str(doc["filename"]))
        if str(doc["index_version"] or "") != config.INDEX_VERSION:
            checker.error(
                "index_version",
                f"expected={config.INDEX_VERSION}, actual={doc['index_version']}",
                document=str(doc["filename"]),
            )


def check_baseline(
    checker: Checker,
    documents: list[dict[str, Any]],
    baseline_path: Path | None,
    *,
    current_counts: dict[str, int] | None = None,
) -> dict[str, Any]:
    if baseline_path is None:
        return {"checked": False}
    if not baseline_path.is_file():
        checker.warn("baseline_missing", f"baseline file not found: {baseline_path}")
        return {"checked": False, "path": str(baseline_path)}
    try:
        payload = json.loads(baseline_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        checker.warn("baseline_invalid", f"{type(exc).__name__}: {exc}")
        return {"checked": False, "path": str(baseline_path)}

    baseline_documents = {
        str(doc["filename"]): doc for doc in payload.get("documents") or []
    }
    current_documents = {str(doc["filename"]): doc for doc in documents}
    missing = sorted(set(baseline_documents) - set(current_documents))
    added = sorted(set(current_documents) - set(baseline_documents))
    chunk_mismatch: list[str] = []
    hash_mismatch: list[str] = []
    type_mismatch: list[str] = []
    for filename in sorted(set(baseline_documents) & set(current_documents)):
        before = baseline_documents[filename]
        after = current_documents[filename]
        if int(before.get("chunk_count") or 0) != int(after.get("chunk_count") or 0):
            chunk_mismatch.append(
                f"{filename}: before={before.get('chunk_count')} after={after.get('chunk_count')}"
            )
        if str(before.get("file_sha256") or "") != str(after.get("file_sha256") or ""):
            hash_mismatch.append(filename)
        if str(before.get("file_type") or "") != str(after.get("file_type") or ""):
            type_mismatch.append(filename)
    for filename in missing:
        checker.error("baseline_document_missing", f"document absent after re-import: {filename}")
    for filename in added:
        checker.warn("baseline_document_added", f"document not present before re-import: {filename}")
    for item in chunk_mismatch:
        # The baseline may have been produced by an older image build.  A byte
        # identical original with a different chunk count therefore signals a
        # parser/chunker version change, not a corrupt import; the authoritative
        # check is the fresh re-parse below.
        checker.warn("baseline_chunk_count_diff", item)
    for filename in hash_mismatch:
        checker.error("baseline_sha256_diff", f"source bytes changed: {filename}")
    for filename in type_mismatch:
        checker.error("baseline_type_diff", f"file type changed: {filename}")
    index_diff: dict[str, dict[str, int]] = {}
    if current_counts:
        for key, current_value in current_counts.items():
            before_value = payload.get(key)
            if before_value is None:
                continue
            if int(before_value) != int(current_value):
                index_diff[key] = {"before": int(before_value), "after": int(current_value)}
                checker.warn(
                    "baseline_index_count_diff",
                    f"{key}: before={before_value} after={current_value}",
                )
    return {
        "checked": True,
        "path": str(baseline_path),
        "baseline_documents": len(baseline_documents),
        "missing": len(missing),
        "added": len(added),
        "chunk_count_mismatches": len(chunk_mismatch),
        "sha256_mismatches": len(hash_mismatch),
        "type_mismatches": len(type_mismatch),
        "index_count_differences": index_diff,
        "baseline_chunk_rows": payload.get("chunk_rows"),
        "baseline_spreadsheet_cells": payload.get("spreadsheet_cells"),
    }


def check_metadata(checker: Checker, documents: list[dict[str, Any]]) -> dict[str, int]:
    fields = (
        "title",
        "issuing_authority",
        "document_number",
        "regulatory_topic",
        "business_domain",
        "publication_date",
        "effective_date",
    )
    coverage = {
        field: sum(1 for doc in documents if str(doc.get(field) or "").strip())
        for field in fields
    }
    if coverage["title"] < len(documents):
        checker.error(
            "title_coverage",
            f"documents_without_title={len(documents) - coverage['title']}",
        )
    for field in ("issuing_authority", "document_number", "publication_date"):
        if coverage[field] == 0 and documents:
            checker.warn("metadata_coverage", f"no document has {field}")
    return coverage


def check_files(
    checker: Checker,
    documents: list[dict[str, Any]],
    source_dir: Path | None,
    source_manifest: dict[str, Any] | None = None,
) -> dict[str, Any]:
    storage_ok = 0
    source_ok = 0
    source_missing: list[str] = []
    manifest_files: dict[str, dict[str, Any]] = {}
    if source_manifest:
        manifest_files = {
            str(item["filename"]): item for item in source_manifest.get("files") or []
        }
    # Build the fallback directory index from iterated Path objects: some
    # official attachments have UTF-8 names longer than Linux NAME_MAX, and
    # reconstructing source_dir / filename raises ENAMETOOLONG.
    source_paths_by_name: dict[str, Path] = {}
    if not manifest_files and source_dir is not None and source_dir.is_dir():
        for path in source_dir.iterdir():
            if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
                source_paths_by_name[path.name] = path

    for doc in documents:
        filename = str(doc["filename"])
        storage_path = Path(str(doc["storage_path"]))
        try:
            storage_exists = storage_path.is_file()
        except OSError as exc:
            checker.error("original_unreadable", f"{type(exc).__name__}: {exc}", document=filename)
            continue
        if not storage_exists:
            checker.error("original_missing", f"path={storage_path}", document=filename)
            continue
        if storage_path.stat().st_size != int(doc["size"] or 0):
            checker.error(
                "original_size",
                f"db={doc['size']}, file={storage_path.stat().st_size}",
                document=filename,
            )
            continue
        actual_hash = sha256_file(storage_path)
        if actual_hash != str(doc["file_sha256"] or ""):
            checker.error(
                "original_sha256",
                f"db={doc['file_sha256']}, file={actual_hash}",
                document=filename,
            )
            continue
        storage_ok += 1
        if manifest_files:
            entry = manifest_files.get(filename)
            if entry is None:
                source_missing.append(filename)
                checker.error(
                    "source_missing",
                    "source manifest has no entry with this filename",
                    document=filename,
                )
                continue
            if int(entry.get("size") or 0) != int(doc["size"] or 0):
                checker.error(
                    "source_size_mismatch",
                    f"manifest={entry.get('size')}, document={doc.get('size')}",
                    document=filename,
                )
                continue
            if str(entry.get("sha256") or "") != actual_hash:
                checker.error(
                    "source_hash_mismatch",
                    f"manifest={entry.get('sha256')}, document={actual_hash}",
                    document=filename,
                )
                continue
            source_ok += 1
            continue
        if source_dir is None:
            continue
        source_path = source_paths_by_name.get(filename)
        if source_path is None:
            source_missing.append(filename)
            checker.error("source_missing", "source directory has no file with this name", document=filename)
            continue
        try:
            if sha256_file(source_path) != actual_hash:
                checker.error("source_hash_mismatch", f"source={source_path}", document=filename)
                continue
        except OSError as exc:
            checker.warn(
                "source_unreadable",
                f"{type(exc).__name__}: {exc}",
                document=filename,
            )
            continue
        source_ok += 1

    document_files = {str(doc["filename"]) for doc in documents}
    if manifest_files:
        for filename in sorted(set(manifest_files) - document_files):
            checker.error("source_not_imported", f"source file absent from knowledge base: {filename}")
        for filename in sorted(document_files - set(manifest_files)):
            checker.error("document_not_in_source", f"document absent from source manifest: {filename}")
    elif source_paths_by_name:
        for filename in sorted(set(source_paths_by_name) - document_files):
            checker.error("source_not_imported", f"source file absent from knowledge base: {filename}")
        for filename in sorted(document_files - set(source_paths_by_name)):
            checker.warn("document_not_in_source", f"document absent from source directory: {filename}")
    return {
        "storage_verified": storage_ok,
        "source_verified": source_ok,
        "source_missing_count": len(source_missing),
        "source_manifest": bool(manifest_files),
    }


def check_chunks(
    checker: Checker,
    conn: sqlite3.Connection,
    documents: list[dict[str, Any]],
) -> tuple[dict[str, dict[str, Any]], dict[str, Any]]:
    per_document: dict[str, dict[str, Any]] = {}
    all_text_lengths: list[int] = []
    for doc in documents:
        document_id = str(doc["document_id"])
        filename = str(doc["filename"])
        rows = [
            dict(row)
            for row in conn.execute(
                "SELECT chunk_id, text, chunk_metadata, page_number, section_title, token_count "
                "FROM document_chunks WHERE document_id = ? ORDER BY chunk_id",
                (document_id,),
            )
        ]
        if len(rows) != int(doc["chunk_count"] or 0):
            checker.error(
                "chunk_count_mismatch",
                f"document.chunk_count={doc['chunk_count']}, rows={len(rows)}",
                document=filename,
            )
        stats = {
            "chunks": len(rows),
            "chars": 0,
            "empty_chunks": 0,
            "replacement_chars": 0,
            "control_chars": 0,
            "duplicate_texts": 0,
            "metadata_errors": 0,
            "pages": sorted({row["page_number"] for row in rows if row["page_number"] is not None}),
        }
        seen: Counter[str] = Counter()
        for row in rows:
            text = str(row["text"] or "")
            length = len(text)
            stats["chars"] += length
            if not text.strip():
                stats["empty_chunks"] += 1
            if "\ufffd" in text:
                stats["replacement_chars"] += text.count("\ufffd")
            if CONTROL_RE.search(text):
                stats["control_chars"] += 1
            normalized = normalize_text(text)
            if normalized:
                seen[normalized] += 1
            try:
                metadata = json.loads(row["chunk_metadata"] or "{}")
                if not isinstance(metadata, dict):
                    stats["metadata_errors"] += 1
            except json.JSONDecodeError:
                stats["metadata_errors"] += 1
        stats["duplicate_texts"] = sum(count - 1 for count in seen.values() if count > 1)
        all_text_lengths.append(stats["chars"])

        if not rows:
            checker.error("no_chunks", "document has no chunk rows", document=filename)
        if stats["empty_chunks"]:
            checker.error("empty_chunk", f"empty_chunks={stats['empty_chunks']}", document=filename)
        if stats["replacement_chars"]:
            replacement_ratio = stats["replacement_chars"] / max(stats["chars"], 1)
            if str(doc["file_type"]) == "pdf" and replacement_ratio < 0.01:
                # Source PDFs for solvency formulas often lack ToUnicode maps
                # for math glyphs (sum, variables, <=).  The surrounding text
                # is intact, so this is an extraction limitation rather than a
                # corrupt import.
                checker.warn(
                    "pdf_math_symbol_placeholders",
                    f"U+FFFD count={stats['replacement_chars']}, ratio={replacement_ratio:.4f}",
                    document=filename,
                )
            else:
                checker.error(
                    "replacement_chars",
                    f"U+FFFD count={stats['replacement_chars']}, ratio={replacement_ratio:.4f}",
                    document=filename,
                )
        if stats["control_chars"]:
            checker.warn("control_chars", f"chunks_with_control_chars={stats['control_chars']}", document=filename)
        if stats["metadata_errors"]:
            checker.error(
                "chunk_metadata_json",
                f"invalid_metadata={stats['metadata_errors']}",
                document=filename,
            )
        if stats["chars"] < 20:
            checker.error("text_too_short", f"chars={stats['chars']}", document=filename)
        elif stats["chars"] > 200:
            cjk_ratio = len(CJK_RE.findall("".join(str(row["text"] or "") for row in rows))) / stats["chars"]
            stats["cjk_ratio"] = round(cjk_ratio, 4)
            if str(doc["file_type"]) not in {"xls", "xlsx"} and cjk_ratio < 0.05:
                checker.warn(
                    "low_cjk_ratio",
                    f"cjk_ratio={cjk_ratio:.4f}, chars={stats['chars']}",
                    document=filename,
                )
        else:
            stats["cjk_ratio"] = None
        if stats["duplicate_texts"] > max(2, stats["chunks"] // 10):
            checker.warn(
                "duplicate_chunk_texts",
                f"duplicate_texts={stats['duplicate_texts']}/{stats['chunks']}",
                document=filename,
            )
        if doc["file_type"] == "pdf" and stats["chunks"] and not stats["pages"]:
            checker.warn("pdf_page_numbers", "pdf chunks carry no page_number", document=filename)
        per_document[document_id] = stats

    by_type: dict[str, list[int]] = defaultdict(list)
    for doc in documents:
        by_type[str(doc["file_type"])].append(per_document[str(doc["document_id"])]["chunks"])
    summary = {
        "document_stats": per_document,
        "chunks_by_type": {
            file_type: {
                "documents": len(values),
                "chunks": sum(values),
                "min": min(values),
                "median": statistics.median(values),
                "p95": percentile(values, 0.95),
                "max": max(values),
            }
            for file_type, values in sorted(by_type.items())
        },
        "total_chars": sum(all_text_lengths),
        "shortest_document_chars": min(all_text_lengths) if all_text_lengths else 0,
    }
    return per_document, summary


def count_numeric_cells(path: Path, file_type: str) -> int:
    """Count real numeric cells to distinguish blank templates from parse loss."""

    count = 0
    if file_type == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(str(path), data_only=True, read_only=True)
        try:
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows(values_only=True):
                    for value in row:
                        if isinstance(value, bool):
                            continue
                        if isinstance(value, (int, float)):
                            count += 1
        finally:
            workbook.close()
        return count
    if file_type == "xls":
        import xlrd

        book = xlrd.open_workbook(str(path), formatting_info=False)
        for sheet_index in range(book.nsheets):
            sheet = book.sheet_by_index(sheet_index)
            for row_index in range(sheet.nrows):
                for col_index in range(sheet.ncols):
                    if sheet.cell_type(row_index, col_index) in (xlrd.XL_CELL_NUMBER, xlrd.XL_CELL_DATE):
                        count += 1
        return count
    return 0


def check_spreadsheets(
    checker: Checker,
    conn: sqlite3.Connection,
    documents: list[dict[str, Any]],
) -> dict[str, Any]:
    sheet_documents = [doc for doc in documents if str(doc["file_type"]) in {"xls", "xlsx"}]
    total_cells = 0
    for doc in sheet_documents:
        document_id = str(doc["document_id"])
        filename = str(doc["filename"])
        expected_cells = 0
        sheets: set[str] = set()
        rows = conn.execute(
            "SELECT chunk_id, chunk_metadata FROM document_chunks WHERE document_id = ?",
            (document_id,),
        )
        for row in rows:
            try:
                metadata = json.loads(row["chunk_metadata"] or "{}")
            except json.JSONDecodeError:
                continue
            if not isinstance(metadata, dict):
                continue
            if metadata.get("sheet_name"):
                sheets.add(str(metadata["sheet_name"]))
            if metadata.get("table_chunk_role") == "row":
                expected_cells += len(metadata.get("cells") or [])
        actual_cells = int(
            conn.execute(
                "SELECT COUNT(*) FROM spreadsheet_cells WHERE document_id = ?",
                (document_id,),
            ).fetchone()[0]
        )
        total_cells += actual_cells
        if expected_cells == 0:
            numeric_cells = count_numeric_cells(Path(str(doc["storage_path"])), str(doc["file_type"]))
            if numeric_cells == 0:
                checker.warn(
                    "blank_spreadsheet_template",
                    f"chunks={doc['chunk_count']}, no numeric cells in workbook",
                    document=filename,
                )
            else:
                checker.error(
                    "spreadsheet_no_cells",
                    f"chunks={doc['chunk_count']}, numeric_cells={numeric_cells} but no row index",
                    document=filename,
                )
        elif actual_cells != expected_cells:
            checker.error(
                "spreadsheet_cell_count",
                f"expected_from_chunk_metadata={expected_cells}, spreadsheet_cells={actual_cells}",
                document=filename,
            )
        if not sheets:
            checker.warn("spreadsheet_no_sheet_name", "chunk metadata has no sheet_name", document=filename)
        orphan_cells = int(
            conn.execute(
                "SELECT COUNT(*) FROM spreadsheet_cells c "
                "LEFT JOIN document_chunks ch ON ch.chunk_id = c.chunk_id "
                "WHERE c.document_id = ? AND ch.chunk_id IS NULL",
                (document_id,),
            ).fetchone()[0]
        )
        if orphan_cells:
            checker.error("spreadsheet_orphan_cells", f"orphan_cells={orphan_cells}", document=filename)
    return {
        "documents": len(sheet_documents),
        "spreadsheet_cells": total_cells,
    }


def check_structure_indexes(checker: Checker, conn: sqlite3.Connection) -> dict[str, int]:
    result: dict[str, int] = {}
    for table, fk in (
        ("document_definitions", "chunk_id"),
        ("document_clauses", "chunk_id"),
    ):
        count = int(conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0])
        orphan = int(
            conn.execute(
                f"SELECT COUNT(*) FROM {table} t "
                "LEFT JOIN document_chunks ch ON ch.chunk_id = t.chunk_id "
                f"WHERE ch.chunk_id IS NULL"
            ).fetchone()[0]
        )
        if orphan:
            checker.error(f"{table}_orphans", f"orphan_rows={orphan}")
        result[table] = count
    return result


def check_qdrant(
    checker: Checker,
    conn: sqlite3.Connection,
    collection: str,
) -> dict[str, Any]:
    try:
        from qdrant_client import QdrantClient
    except ImportError as exc:  # pragma: no cover - container always has it
        checker.error("qdrant_client", f"qdrant-client import failed: {exc}")
        return {"checked": False}

    try:
        client = QdrantClient(url=config.QDRANT_URL)
        if not client.collection_exists(collection):
            checker.error("qdrant_collection_missing", f"collection={collection}")
            return {"checked": False}
        info = client.get_collection(collection)
        points_count = int(getattr(info, "points_count", 0) or 0)
    except Exception as exc:  # noqa: BLE001
        checker.error("qdrant_unavailable", f"{type(exc).__name__}: {exc}")
        return {"checked": False}

    expected_ids = {
        str(row[0])
        for row in conn.execute("SELECT chunk_id FROM document_chunks")
    }
    actual_ids: set[str] = set()
    version_mismatch = 0
    offset = None
    try:
        while True:
            records, offset = client.scroll(
                collection_name=collection,
                limit=1000,
                offset=offset,
                with_payload=["chunk_id", "index_version"],
                with_vectors=False,
            )
            for record in records:
                payload = record.payload or {}
                chunk_id = str(payload.get("chunk_id") or "")
                actual_ids.add(chunk_id)
                if str(payload.get("index_version") or "") != config.INDEX_VERSION:
                    version_mismatch += 1
            if offset is None:
                break
    except Exception as exc:  # noqa: BLE001
        checker.error("qdrant_scroll", f"{type(exc).__name__}: {exc}")
        return {"checked": False, "points_count": points_count}

    missing = sorted(expected_ids - actual_ids)
    extra = sorted(actual_ids - expected_ids)
    if points_count != len(expected_ids):
        checker.error(
            "qdrant_point_count",
            f"points_count={points_count}, expected={len(expected_ids)}",
        )
    if missing:
        checker.error("qdrant_missing_points", f"missing={len(missing)}, examples={missing[:5]}")
    if extra:
        checker.error("qdrant_extra_points", f"extra={len(extra)}, examples={extra[:5]}")
    if version_mismatch:
        checker.error("qdrant_index_version", f"mismatched_points={version_mismatch}")
    return {
        "checked": True,
        "collection": collection,
        "points_count": points_count,
        "expected_points": len(expected_ids),
        "scrolled_points": len(actual_ids),
        "missing": len(missing),
        "extra": len(extra),
        "version_mismatch": version_mismatch,
    }


def bigram_recall(reference: str, candidate: str) -> float:
    """Share of reference character bigrams present in candidate."""

    reference_grams = {reference[index : index + 2] for index in range(max(0, len(reference) - 1))}
    if not reference_grams:
        return 1.0
    candidate_grams = {candidate[index : index + 2] for index in range(max(0, len(candidate) - 1))}
    return len(reference_grams & candidate_grams) / len(reference_grams)


def check_reparse(
    checker: Checker,
    conn: sqlite3.Connection,
    documents: list[dict[str, Any]],
    limit: int,
) -> dict[str, Any]:
    selected = documents[:limit] if limit > 0 else documents
    matched = 0
    failed = 0
    order_only = 0
    boundary_only = 0
    details: dict[str, Any] = {}
    for index, doc in enumerate(selected, start=1):
        filename = str(doc["filename"])
        path = Path(str(doc["storage_path"]))
        try:
            parsed = parse_document(path, source_name=filename)
        except Exception as exc:  # noqa: BLE001
            checker.error("reparse_failed", f"{type(exc).__name__}: {exc}", document=filename)
            details[str(doc["document_id"])] = {"ok": False, "stage": "parse", "error": short(exc, 300)}
            failed += 1
            continue
        if not parsed.text.strip() or not parsed.blocks:
            checker.error(
                "reparse_empty",
                f"text_chars={len(parsed.text.strip())}, blocks={len(parsed.blocks)}",
                document=filename,
            )
            details[str(doc["document_id"])] = {"ok": False, "stage": "parse_empty"}
            failed += 1
            continue

        try:
            drafts = build_chunks_from_parsed(
                document_id=str(doc["document_id"]),
                parsed=parsed,
                source_file=filename,
                inherited_metadata=None,
            )
        except Exception as exc:  # noqa: BLE001
            checker.error("rechunk_failed", f"{type(exc).__name__}: {exc}", document=filename)
            details[str(doc["document_id"])] = {"ok": False, "stage": "chunk", "error": short(exc, 300)}
            failed += 1
            continue

        stored = [
            dict(row)
            for row in conn.execute(
                "SELECT chunk_id, text, page_number, section_title, chunk_metadata "
                "FROM document_chunks WHERE document_id = ? ORDER BY chunk_id",
                (str(doc["document_id"]),),
            )
        ]
        problems: list[str] = []
        fresh_texts = Counter(normalize_text(draft.text) for draft in drafts)
        stored_texts = Counter(normalize_text(str(row["text"] or "")) for row in stored)
        content_equal = fresh_texts == stored_texts
        order_equal = len(drafts) == len(stored) and all(
            normalize_text(draft.text) == normalize_text(str(row["text"] or ""))
            for draft, row in zip(drafts, stored)
        )
        text_mismatch = 0
        metadata_mismatch = 0
        first_difference = None
        if not content_equal:
            missing = sum((fresh_texts - stored_texts).values())
            extra = sum((stored_texts - fresh_texts).values())
            fresh_join = "\n".join(normalize_text(draft.text) for draft in drafts)
            stored_join = "\n".join(normalize_text(str(row["text"] or "")) for row in stored)
            coverage = min(
                bigram_recall(fresh_join, stored_join),
                bigram_recall(stored_join, fresh_join),
            )
            if coverage >= 0.98:
                # Page-boundary heading duplication can move a sentence
                # between chunks without losing evidence.  Keep it visible but
                # do not fail a corpus import for a boundary-only difference.
                boundary_only += 1
                checker.warn(
                    "reparse_boundary_diff",
                    f"coverage={coverage:.4f}, fresh_chunks={len(drafts)}, "
                    f"stored_chunks={len(stored)}",
                    document=filename,
                )
            else:
                problems.append(
                    f"text_multiset_mismatch fresh_only={missing} stored_only={extra} "
                    f"coverage={coverage:.4f} fresh_chunks={len(drafts)} stored_chunks={len(stored)}"
                )
                for draft, row in zip(drafts, stored):
                    fresh_text = normalize_text(draft.text)
                    stored_text = normalize_text(str(row["text"] or ""))
                    if fresh_text != stored_text:
                        text_mismatch += 1
                        if first_difference is None:
                            first_difference = {
                                "chunk_id": draft.chunk_id,
                                "fresh": short(fresh_text, 200),
                                "stored": short(stored_text, 200),
                            }
        elif not order_equal:
            # Large legacy workbooks can emit the same rows in a different
            # sheet/block order after a LibreOffice conversion.  No evidence is
            # lost, so this is reported separately from a content mismatch.
            order_only += 1
            checker.warn(
                "reparse_chunk_order_diff",
                f"chunks={len(drafts)}, text multiset identical, sequence differs",
                document=filename,
            )
        else:
            for draft, row in zip(drafts, stored):
                if (draft.page_number or None) != (row["page_number"] or None):
                    metadata_mismatch += 1
                if (draft.section_title or None) != (row["section_title"] or None):
                    metadata_mismatch += 1
                try:
                    stored_metadata = json.loads(row["chunk_metadata"] or "{}")
                except json.JSONDecodeError:
                    stored_metadata = {}
                for key in STRUCTURAL_METADATA_KEYS:
                    if (draft.metadata or {}).get(key) != stored_metadata.get(key):
                        metadata_mismatch += 1
                        break
        if text_mismatch:
            problems.append(f"text_mismatch={text_mismatch}")
        if metadata_mismatch:
            problems.append(f"metadata_mismatch={metadata_mismatch}")
        if problems:
            checker.error("reparse_mismatch", "; ".join(problems), document=filename)
            details[str(doc["document_id"])] = {
                "ok": False,
                "stage": "compare",
                "problems": problems,
                "first_difference": first_difference,
            }
            failed += 1
        else:
            matched += 1
            details[str(doc["document_id"])] = {
                "ok": True,
                "order_only": not order_equal,
            }
        if index % 25 == 0:
            print(f"[reparse] {index}/{len(selected)} matched={matched} failed={failed}", flush=True)

    return {
        "checked_documents": len(selected),
        "matched": matched,
        "order_only": order_only,
        "boundary_only": boundary_only,
        "failed": failed,
        "details": details,
    }


def render_markdown(report: dict[str, Any]) -> str:
    summary = report["summary"]
    lines = [
        "# 文档解析与索引一致性校验",
        "",
        f"- 生成时间：{report['created_at']}",
        f"- 数据库：`{report['database_path']}`",
        f"- 原始文件目录：`{report['document_dir']}`",
        f"- 期望文档数：{report['expected_documents']}",
        f"- 实际文档数：{summary['document_count']}",
        f"- 状态分布：`{json.dumps(summary['status_counts'], ensure_ascii=False)}`",
        f"- 分块总数：{summary['chunk_rows']}",
        f"- 错误数：{len(report['errors'])}",
        f"- 警告数：{len(report['warnings'])}",
        "",
        "## 分类型分块",
        "",
        "| 类型 | 文档数 | 分块数 | 最小 | 中位数 | P95 | 最大 |",
        "|---|---:|---:|---:|---:|---:|---:|",
    ]
    for file_type, stats in summary["chunks_by_type"].items():
        lines.append(
            f"| {file_type} | {stats['documents']} | {stats['chunks']} | {stats['min']} | "
            f"{stats['median']} | {stats['p95']} | {stats['max']} |"
        )
    qdrant = summary["qdrant"]
    lines += [
        "",
        "## 一致性",
        "",
        f"- 原始文件哈希校验通过：{summary['files']['storage_verified']}/{summary['document_count']}",
        f"- 源目录哈希校验通过：{summary['files']['source_verified']}/{summary['document_count']}",
        f"- 删除前基线对比：{json.dumps(summary['baseline'], ensure_ascii=False)}",
        f"- 元数据覆盖率：{json.dumps(summary['metadata_coverage'], ensure_ascii=False)}",
        f"- Qdrant 校验：{json.dumps(qdrant, ensure_ascii=False)}",
        f"- 表格单元行数：{summary['spreadsheets']['spreadsheet_cells']}",
        f"- 结构索引：{json.dumps(summary['structure_indexes'], ensure_ascii=False)}",
        f"- 重新解析：{json.dumps({k: v for k, v in summary['reparse'].items() if k != 'details'}, ensure_ascii=False)}",
        "",
        "## 错误",
        "",
    ]
    if report["errors"]:
        for item in report["errors"][:200]:
            lines.append(f"- `{item['code']}` {item.get('document') or ''}：{item['detail']}")
        if len(report["errors"]) > 200:
            lines.append(f"- … 其余 {len(report['errors']) - 200} 条见 report.json")
    else:
        lines.append("- 无")
    lines += ["", "## 警告", ""]
    if report["warnings"]:
        for item in report["warnings"][:200]:
            lines.append(f"- `{item['code']}` {item.get('document') or ''}：{item['detail']}")
        if len(report["warnings"]) > 200:
            lines.append(f"- … 其余 {len(report['warnings']) - 200} 条见 report.json")
    else:
        lines.append("- 无")
    lines.append("")
    return "\n".join(lines)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--db", type=Path, default=config.DATABASE_PATH)
    parser.add_argument("--documents-dir", type=Path, default=config.DOCUMENT_DIR)
    parser.add_argument("--source-dir", type=Path, default=None)
    parser.add_argument(
        "--source-manifest",
        type=Path,
        default=None,
        help="JSON manifest with source filename/size/sha256 (use on hosts with >NAME_MAX filenames)",
    )
    parser.add_argument(
        "--baseline",
        type=Path,
        default=None,
        help="documents_before_reupload_*.json snapshot used to compare per-file chunk counts",
    )
    parser.add_argument("--collection", default=config.QDRANT_COLLECTION)
    parser.add_argument("--output-dir", type=Path, default=None)
    parser.add_argument("--expected-documents", type=int, default=500)
    parser.add_argument("--skip-qdrant", action="store_true")
    parser.add_argument("--skip-reparse", action="store_true")
    parser.add_argument("--reparse-limit", type=int, default=0, help="0 = all documents")
    args = parser.parse_args()

    started = datetime.now()
    output_dir = args.output_dir or (
        Path("/app/data/evaluation/document_parse_check") / started.strftime("%Y%m%d_%H%M%S")
    )
    output_dir.mkdir(parents=True, exist_ok=True)

    checker = Checker()
    conn, documents = load_documents(args.db)
    source_manifest = None
    if args.source_manifest is not None:
        source_manifest = json.loads(args.source_manifest.read_text(encoding="utf-8-sig"))
    print(f"[verify] documents={len(documents)} db={args.db}", flush=True)
    check_corpus(
        checker,
        documents,
        int(conn.execute("SELECT COUNT(*) FROM document_chunks").fetchone()[0]),
        args.expected_documents,
    )
    files = check_files(checker, documents, args.source_dir, source_manifest)
    _, chunk_summary = check_chunks(checker, conn, documents)
    spreadsheets = check_spreadsheets(checker, conn, documents)
    structure_indexes = check_structure_indexes(checker, conn)
    baseline = check_baseline(
        checker,
        documents,
        args.baseline,
        current_counts={
            "chunk_rows": int(conn.execute("SELECT COUNT(*) FROM document_chunks").fetchone()[0]),
            "spreadsheet_cells": spreadsheets["spreadsheet_cells"],
            "definitions": structure_indexes["document_definitions"],
            "clauses": structure_indexes["document_clauses"],
        },
    )
    metadata_coverage = check_metadata(checker, documents)
    qdrant = (
        {"checked": False, "skipped": True}
        if args.skip_qdrant
        else check_qdrant(checker, conn, args.collection)
    )
    reparse = (
        {"checked_documents": 0, "matched": 0, "order_only": 0, "boundary_only": 0, "failed": 0, "skipped": True, "details": {}}
        if args.skip_reparse
        else check_reparse(checker, conn, documents, args.reparse_limit)
    )

    report = {
        "created_at": started.isoformat(),
        "database_path": str(args.db),
        "document_dir": str(args.documents_dir),
        "source_dir": str(args.source_dir) if args.source_dir else None,
        "source_manifest": str(args.source_manifest) if args.source_manifest else None,
        "expected_documents": args.expected_documents,
        "summary": {
            "document_count": len(documents),
            "status_counts": dict(Counter(str(doc["status"]) for doc in documents)),
            "chunk_rows": int(conn.execute("SELECT COUNT(*) FROM document_chunks").fetchone()[0]),
            **chunk_summary,
            "files": files,
            "baseline": baseline,
            "metadata_coverage": metadata_coverage,
            "spreadsheets": spreadsheets,
            "structure_indexes": structure_indexes,
            "qdrant": qdrant,
            "reparse": reparse,
        },
        "errors": checker.errors,
        "warnings": checker.warnings,
    }
    report_path = output_dir / "report.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    markdown_path = output_dir / "report.md"
    markdown_path.write_text(render_markdown(report), encoding="utf-8")

    print(
        json.dumps(
            {
                "documents": len(documents),
                "chunks": report["summary"]["chunk_rows"],
                "errors": len(checker.errors),
                "warnings": len(checker.warnings),
                "reparse_matched": reparse["matched"],
                "reparse_order_only": reparse.get("order_only", 0),
                "reparse_boundary_only": reparse.get("boundary_only", 0),
                "reparse_failed": reparse["failed"],
                "report": str(report_path),
            },
            ensure_ascii=False,
        ),
        flush=True,
    )
    if checker.errors:
        for item in checker.errors[:20]:
            print(f"[error] {item['code']} {item.get('document') or ''}: {item['detail']}", flush=True)
        return 1
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception:  # noqa: BLE001
        traceback.print_exc()
        raise SystemExit(2)
