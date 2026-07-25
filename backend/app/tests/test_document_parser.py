import zipfile

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from backend.app.core.database import Base
from backend.app.models.document import Document, DocumentChunk, SpreadsheetCell
from backend.app.services.chunk_service import build_chunks_from_parsed, count_tokens
from backend.app.services.document_types import ParsedBlock, ParsedDocument
from backend.app.services.document_parser import DocumentParseError, available_loader_names, parse_document
from backend.app.services.document_processing_service import ensure_document_chunks
from backend.app.services.loader_evaluation import evaluate_document_loaders
from backend.app.services.spreadsheet_cell_index_service import rebuild_spreadsheet_cell_index


def test_parse_markdown_returns_structured_blocks(tmp_path) -> None:
    path = tmp_path / "rules.md"
    path.write_text("# 监管填报说明\n\n## 资产合计\n资产合计应等于各项资产分项金额合计。\n", encoding="utf-8")

    parsed = parse_document(path)

    assert parsed.metadata["parser_version"] == "structured-v4-formula"
    assert parsed.metadata["loader_name"] == "markdown"
    assert [block.block_type for block in parsed.blocks] == ["heading", "heading", "paragraph"]
    assert parsed.blocks[2].section_title == "资产合计"


def test_parse_csv_builds_cell_level_rows(tmp_path) -> None:
    path = tmp_path / "监管统计.csv"
    path.write_text("指标,2025年\n资本充足率,12.5\n", encoding="utf-8")

    parsed = parse_document(path)

    row = next(block for block in parsed.blocks if block.metadata.get("table_chunk_role") == "row")
    assert row.metadata["spreadsheet_table"] is True
    assert row.metadata["period"]["year"] == 2025
    assert row.metadata["cells"][1]["coordinate"] == "B2"
    assert row.metadata["cells"][1]["normalized_value"] == 12.5


def test_parse_csv_infers_month_period_and_row_unit(tmp_path) -> None:
    path = tmp_path / "自制月报.csv"
    path.write_text("指标,2026年1月,单位\n资产合计,7000,万元\n", encoding="utf-8")

    parsed = parse_document(path)

    row = next(block for block in parsed.blocks if block.metadata.get("table_chunk_role") == "row")
    value_cell = next(cell for cell in row.metadata["cells"] if cell["coordinate"] == "B2")
    assert row.metadata["period"] == {"year": 2026, "month": 1}
    assert value_cell["normalized_value"] == 7000
    assert value_cell["unit"] == "万元"


def test_force_rebuild_refreshes_spreadsheet_cell_index(tmp_path) -> None:
    path = tmp_path / "自制月报.csv"
    path.write_text("指标,2026年1月,单位\n资产合计,7000,万元\n", encoding="utf-8")
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)

    with Session(engine) as db:
        document = Document(
            document_id="DOC-TEST-REBUILD",
            filename=path.name,
            filename_norm=path.name,
            file_type="csv",
            size=path.stat().st_size,
            file_sha256="b" * 64,
            storage_path=str(path),
            status="indexed",
            index_version="old",
            version_status="unknown",
            metadata_status="inferred",
            document_metadata="{}",
            chunk_count=1,
        )
        stale_chunk = DocumentChunk(
            chunk_id="DOC-TEST-REBUILD-CHUNK-STALE",
            document_id=document.document_id,
            text="旧表格行证据",
            embedding_text="旧表格行证据",
            chunk_metadata='{"table_chunk_role":"row","period":{},"cells":[{"coordinate":"B2","value":"7000","normalized_value":7000}]}',
            token_count=4,
            index_status="indexed",
            index_version="old",
            title=path.stem,
            source_file=path.name,
        )
        db.add_all([document, stale_chunk])
        db.flush()
        rebuild_spreadsheet_cell_index(db, document, [stale_chunk])
        assert db.query(SpreadsheetCell).one().year is None
        stale_chunk_id = stale_chunk.chunk_id
        document_id = document.document_id

        rebuilt = ensure_document_chunks(db, document, force_rebuild=True)

        assert all(chunk.chunk_id != stale_chunk_id for chunk in rebuilt)
        cell = (
            db.query(SpreadsheetCell)
            .filter(SpreadsheetCell.document_id == document_id, SpreadsheetCell.coordinate == "B2")
            .one()
        )
        assert cell.year == 2026
        assert cell.month == 1
        assert cell.unit == "万元"
        assert cell.numeric_value == 7000


def test_parse_jsonl_requires_content_and_rejects_qa_answers(tmp_path) -> None:
    valid = tmp_path / "rules.jsonl"
    valid.write_text('{"title":"规则一","text":"资本充足率应符合监管要求。"}\n', encoding="utf-8")
    parsed = parse_document(valid)
    assert parsed.blocks[0].section_title == "规则一"

    qa = tmp_path / "qa.jsonl"
    qa.write_text('{"question":"答案是什么","answer":"A","evidence":"标准答案"}\n', encoding="utf-8")
    with pytest.raises(DocumentParseError, match="QA/答案数据"):
        parse_document(qa)


def test_parse_normal_paragraph_articles_adds_article_locator(tmp_path) -> None:
    path = tmp_path / "rules.txt"
    path.write_text("第一章 总则\n\n第一条 银行业金融机构应当依法报送。\n第二条 不得迟报。", encoding="utf-8")

    parsed = parse_document(path)

    articles = [block for block in parsed.blocks if block.metadata.get("structure_type") == "article"]
    assert [block.metadata["article_number"] for block in articles] == ["第一条", "第二条"]
    assert articles[0].text.startswith("第一条")


def test_parse_html_preserves_headings_paragraphs_and_tables(tmp_path) -> None:
    path = tmp_path / "rule.html"
    path.write_text(
        "<html><head><title>监管规则</title></head><body><h1>第一章</h1><p>正文依据。</p>"
        "<table><tr><th>指标</th><th>值</th></tr><tr><td>资本</td><td>12</td></tr></table></body></html>",
        encoding="utf-8",
    )

    parsed = parse_document(path)

    assert parsed.metadata["html_title"] == "监管规则"
    assert any(block.block_type == "heading" and block.text == "第一章" for block in parsed.blocks)
    assert any(block.block_type == "table" and "资本" in block.text for block in parsed.blocks)


def test_loader_order_is_configured_by_file_type(tmp_path) -> None:
    assert available_loader_names(tmp_path / "rules.txt") == ["text", "unstructured"]
    assert available_loader_names(tmp_path / "rules.md") == ["markdown", "unstructured"]
    assert available_loader_names(tmp_path / "rules.doc") == ["libreoffice-doc", "antiword-doc"]
    assert available_loader_names(tmp_path / "rules.docx") == ["python-docx", "docling", "unstructured"]
    assert available_loader_names(tmp_path / "rules.pdf") == ["pymupdf4llm", "docling", "unstructured", "pypdf"]
    assert available_loader_names(tmp_path / "rules.xls") == ["spreadsheet-xls"]
    assert available_loader_names(tmp_path / "rules.xlsx") == ["spreadsheet-xlsx"]


def test_markdown_table_block_carries_structured_metadata(tmp_path) -> None:
    path = tmp_path / "rules.md"
    path.write_text(
        "## 资产指标\n\n| 字段 | 口径 |\n| --- | --- |\n| 资产合计 | 资产分项合计 |\n",
        encoding="utf-8",
    )

    parsed = parse_document(path)
    table_block = next(block for block in parsed.blocks if block.block_type == "table")

    assert table_block.metadata["headers"] == ["字段", "口径"]
    assert table_block.metadata["rows"][0]["row_index"] == 1
    assert table_block.metadata["rows"][0]["cells"] == {"字段": "资产合计", "口径": "资产分项合计"}
    assert "| 字段 | 口径 |" in table_block.metadata["raw_table_text"]


def test_chunks_inherit_section_from_structured_blocks(tmp_path) -> None:
    path = tmp_path / "rules.md"
    path.write_text("# 监管填报说明\n\n## 资产合计\n资产合计应等于各项资产分项金额合计。\n", encoding="utf-8")
    parsed = parse_document(path)

    chunks = build_chunks_from_parsed(document_id="DOC-TEST-0001", parsed=parsed)

    assert len(chunks) == 1
    assert chunks[0].section_title == "资产合计"
    assert "资产合计应等于各项资产分项金额合计" in chunks[0].text
    assert "章节路径：监管填报说明 > 资产合计" in chunks[0].embedding_text
    assert "章节：资产合计" in chunks[0].embedding_text
    assert "内容类型：正文" in chunks[0].embedding_text
    assert chunks[0].token_count > 0


def test_contextual_embedding_text_includes_document_and_location_metadata() -> None:
    parsed = ParsedDocument(
        text="3.1 资产合计差异处理\n\n若差异来自外币折算，应保留汇率日期。",
        metadata={"source_format": "pdf"},
        blocks=[
            ParsedBlock(
                text="3.1 资产合计差异处理",
                block_type="heading",
                order_index=1,
                page_number=7,
                section_title="3.1 资产合计差异处理",
                level=2,
            ),
            ParsedBlock(
                text="若差异来自外币折算，应保留汇率日期。",
                block_type="paragraph",
                order_index=2,
                page_number=7,
                section_title="3.1 资产合计差异处理",
            ),
        ],
    )

    chunks = build_chunks_from_parsed(
        document_id="DOC-TEST-0001",
        parsed=parsed,
        source_file="G01资产负债统计表填报说明.pdf",
    )

    assert len(chunks) == 1
    assert chunks[0].text == "3.1 资产合计差异处理\n\n若差异来自外币折算，应保留汇率日期。"
    assert "来源文件：G01资产负债统计表填报说明.pdf" in chunks[0].embedding_text
    assert "文档格式：pdf" in chunks[0].embedding_text
    assert "章节路径：3.1 资产合计差异处理" in chunks[0].embedding_text
    assert "条款号：3.1" in chunks[0].embedding_text
    assert "父条款号：3" in chunks[0].embedding_text
    assert "页码：7" in chunks[0].embedding_text
    assert chunks[0].embedding_text.endswith(chunks[0].text)


def test_chunks_keep_section_hierarchy_and_neighbor_ids(tmp_path) -> None:
    path = tmp_path / "rules.md"
    path.write_text(
        "## 3. 资产合计与校验关系\n"
        "资产合计应等于各项资产分项金额之和。\n\n"
        "### 3.1 资产合计差异处理\n"
        "若差异来自外币折算，应保留汇率日期、折算规则和原币金额来源。\n",
        encoding="utf-8",
    )
    parsed = parse_document(path)

    chunks = build_chunks_from_parsed(document_id="DOC-TEST-0001", parsed=parsed)

    assert len(chunks) == 2
    assert chunks[0].section_number == "3"
    assert chunks[0].parent_section_number is None
    assert chunks[0].section_path == ["3. 资产合计与校验关系"]
    assert chunks[0].next_chunk_id == chunks[1].chunk_id
    assert chunks[1].section_number == "3.1"
    assert chunks[1].parent_section_number == "3"
    assert chunks[1].section_path == ["3. 资产合计与校验关系", "3.1 资产合计差异处理"]
    assert chunks[1].previous_chunk_id == chunks[0].chunk_id


def test_chunk_splitting_respects_max_tokens_and_overlap() -> None:
    text = "资产质量" * 900
    parsed = ParsedDocument(
        text=text,
        blocks=[ParsedBlock(text=text, block_type="paragraph", order_index=1, section_title="资产质量")],
    )

    chunks = build_chunks_from_parsed(document_id="DOC-TEST-0001", parsed=parsed)

    assert len(chunks) > 1
    assert all(chunk.token_count <= 800 for chunk in chunks)
    assert chunks[0].text[-10:] in chunks[1].text


def test_long_paragraph_prefers_sentence_boundaries() -> None:
    sentence_a = "资产合计应等于各项资产分项金额合计。"
    sentence_b = "填报人员应保留口径说明和数据来源。"
    text = (sentence_a + sentence_b) * 90
    parsed = ParsedDocument(
        text=text,
        blocks=[ParsedBlock(text=text, block_type="paragraph", order_index=1, section_title="资产合计")],
    )

    chunks = build_chunks_from_parsed(document_id="DOC-TEST-0001", parsed=parsed)

    assert len(chunks) > 1
    assert all(count_tokens(chunk.text) <= 800 for chunk in chunks)
    assert all(chunk.text.endswith("。") for chunk in chunks)
    assert not any(chunk.text.startswith("合计。") for chunk in chunks[1:])


def test_semantic_break_splits_topic_shift(monkeypatch) -> None:
    paragraphs = [
        "资产合计应等于各项资产分项金额合计。" * 80,
        "资产分类应按照制度规定保持口径一致。" * 80,
        "资本充足率指标应按照监管口径计算。" * 80,
    ]
    parsed = ParsedDocument(
        text="\n\n".join(paragraphs),
        blocks=[
            ParsedBlock(
                text=paragraph,
                block_type="paragraph",
                order_index=index,
                section_title="监管指标",
            )
            for index, paragraph in enumerate(paragraphs, start=1)
        ],
    )

    monkeypatch.setattr(
        "backend.app.services.chunk_service.embed_for_semantic_split",
        lambda texts: [[1.0, 0.0], [0.95, 0.05], [0.0, 1.0]],
    )

    chunks = build_chunks_from_parsed(document_id="DOC-TEST-0001", parsed=parsed)

    assert len(chunks) >= 2
    assert any("资本充足率" in chunk.text for chunk in chunks)
    assert all(count_tokens(chunk.text) <= 800 for chunk in chunks)


def test_semantic_chunking_falls_back_when_embedding_is_unavailable(monkeypatch) -> None:
    from backend.app.services.embedding_service import EmbeddingServiceError

    paragraphs = [
        "资产合计应等于各项资产分项金额合计。" * 80,
        "资本充足率指标应按照监管口径计算。" * 80,
    ]
    parsed = ParsedDocument(
        text="\n\n".join(paragraphs),
        blocks=[
            ParsedBlock(
                text=paragraph,
                block_type="paragraph",
                order_index=index,
                section_title="监管指标",
            )
            for index, paragraph in enumerate(paragraphs, start=1)
        ],
    )
    monkeypatch.setattr(
        "backend.app.services.chunk_service.embed_for_semantic_split",
        lambda texts: (_ for _ in ()).throw(EmbeddingServiceError("model unavailable")),
    )

    chunks = build_chunks_from_parsed(document_id="DOC-TEST-0001", parsed=parsed)

    assert chunks
    assert any("资本充足率" in chunk.text for chunk in chunks)
    assert all(count_tokens(chunk.text) <= 800 for chunk in chunks)


def test_table_block_becomes_independent_chunk() -> None:
    parsed = ParsedDocument(
        text="监管填报说明\n\n字段 | 口径\n资产合计 | 资产分项合计",
        blocks=[
            ParsedBlock(text="监管填报说明", block_type="heading", order_index=1, section_title="监管填报说明"),
            ParsedBlock(text="字段 | 口径\n资产合计 | 资产分项合计", block_type="table", order_index=2, section_title="监管填报说明"),
        ],
    )

    chunks = build_chunks_from_parsed(document_id="DOC-TEST-0001", parsed=parsed)

    assert len(chunks) == 2
    assert chunks[0].text == "表格摘要：监管填报说明。表头：字段、口径。共1行数据。"
    assert chunks[0].metadata["table_chunk_role"] == "summary"
    assert chunks[1].text == "表格行证据：在《监管填报说明》中，字段为“资产合计”时，口径为“资产分项合计”。"
    assert chunks[1].chunk_type == "table"
    assert chunks[1].metadata["row_cells"] == {"字段": "资产合计", "口径": "资产分项合计"}
    assert "内容类型：表格" in chunks[1].embedding_text
    assert "表头：字段、口径" in chunks[1].embedding_text
    assert "行数据：字段=资产合计；口径=资产分项合计" in chunks[1].embedding_text


def test_short_paragraph_before_table_merges_into_table_chunk() -> None:
    parsed = ParsedDocument(
        text="2.1 不应纳入普惠小微统计的情形\n\n以下情形在本模拟文档中不应纳入普惠小微贷款统计：\n\n| 情形 | 处理口径 |",
        blocks=[
            ParsedBlock(text="2.1 不应纳入普惠小微统计的情形", block_type="heading", order_index=1, section_title="2.1 不应纳入普惠小微统计的情形"),
            ParsedBlock(text="以下情形在本模拟文档中不应纳入普惠小微贷款统计：", block_type="paragraph", order_index=2, section_title="2.1 不应纳入普惠小微统计的情形"),
            ParsedBlock(text="| 情形 | 处理口径 |\n| --- | --- |\n| 借款用途为个人住房装修 | 不纳入普惠小微贷款 |", block_type="table", order_index=3, section_title="2.1 不应纳入普惠小微统计的情形"),
        ],
    )

    chunks = build_chunks_from_parsed(document_id="DOC-TEST-0001", parsed=parsed)

    assert len(chunks) == 2
    assert all(chunk.chunk_type == "table" for chunk in chunks)
    assert chunks[0].metadata["table_chunk_role"] == "summary"
    assert chunks[1].metadata["table_chunk_role"] == "row"
    assert chunks[1].text.startswith("2.1 不应纳入普惠小微统计的情形")
    assert "以下情形在本模拟文档中不应纳入普惠小微贷款统计" in chunks[1].text
    assert "表格行证据：在《2.1 不应纳入普惠小微统计的情形》中，情形为“借款用途为个人住房装修”时，处理口径为“不纳入普惠小微贷款”。" in chunks[1].text


def test_large_table_splits_by_complete_rows_with_repeated_header() -> None:
    rows = ["| 情形 | 处理口径 |", "| --- | --- |"]
    rows.extend(f"| 情形{i} | 处理口径{i}，应保留完整行，不得截断。 |" for i in range(1, 90))
    parsed = ParsedDocument(
        text="\n".join(rows),
        blocks=[
            ParsedBlock(
                text="\n".join(rows),
                block_type="table",
                order_index=1,
                section_title="普惠小微统计口径",
            )
        ],
    )

    chunks = build_chunks_from_parsed(document_id="DOC-TEST-0001", parsed=parsed)

    assert len(chunks) == 90
    assert all(chunk.chunk_type == "table" for chunk in chunks)
    assert chunks[0].text == "表格摘要：普惠小微统计口径。表头：情形、处理口径。共89行数据。"
    assert chunks[1].text == "表格行证据：在《普惠小微统计口径》中，情形为“情形1”时，处理口径为“处理口径1，应保留完整行，不得截断。”。"
    assert chunks[-1].text == "表格行证据：在《普惠小微统计口径》中，情形为“情形89”时，处理口径为“处理口径89，应保留完整行，不得截断。”。"


def test_large_table_with_context_repeats_context_and_header() -> None:
    rows = ["| 情形 | 处理口径 |", "| --- | --- |"]
    rows.extend(f"| 情形{i} | 处理口径{i}，应保留完整行，不得截断。 |" for i in range(1, 90))
    parsed = ParsedDocument(
        text="\n".join(rows),
        blocks=[
            ParsedBlock(text="2.1 不应纳入普惠小微统计的情形", block_type="heading", order_index=1, section_title="2.1 不应纳入普惠小微统计的情形"),
            ParsedBlock(text="以下情形在本模拟文档中不应纳入普惠小微贷款统计：", block_type="paragraph", order_index=2, section_title="2.1 不应纳入普惠小微统计的情形"),
            ParsedBlock(
                text="\n".join(rows),
                block_type="table",
                order_index=3,
                section_title="2.1 不应纳入普惠小微统计的情形",
            ),
        ],
    )

    chunks = build_chunks_from_parsed(document_id="DOC-TEST-0001", parsed=parsed)

    assert len(chunks) == 90
    assert all(chunk.text.startswith("2.1 不应纳入普惠小微统计的情形") for chunk in chunks)
    assert all("以下情形在本模拟文档中不应纳入普惠小微贷款统计" in chunk.text for chunk in chunks)
    assert chunks[0].metadata["table_chunk_role"] == "summary"
    assert all("表格行证据：在《2.1 不应纳入普惠小微统计的情形》中，情形为" in chunk.text for chunk in chunks[1:])


def test_docx_table_exports_standard_markdown(tmp_path) -> None:
    from docx import Document as DocxDocument

    path = tmp_path / "rules.docx"
    document = DocxDocument()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "情形"
    table.cell(0, 1).text = "处理口径"
    table.cell(1, 0).text = "同一合同拆分为多笔"
    table.cell(1, 1).text = "合并判断"
    document.save(path)

    parsed = parse_document(path)

    assert parsed.blocks[0].block_type == "table"
    assert parsed.blocks[0].text.splitlines() == [
        "| 情形 | 处理口径 |",
        "| --- | --- |",
        "| 同一合同拆分为多笔 | 合并判断 |",
    ]


def test_docx_loader_preserves_paragraph_and_table_order(tmp_path) -> None:
    from docx import Document as DocxDocument

    path = tmp_path / "rules.docx"
    document = DocxDocument()
    document.add_paragraph("一、填报说明")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "口径"
    table.cell(1, 0).text = "资产合计"
    table.cell(1, 1).text = "资产分项合计"
    document.add_paragraph("二、留痕要求")
    document.save(path)

    parsed = parse_document(path)

    assert [block.block_type for block in parsed.blocks] == ["paragraph", "table", "paragraph"]
    assert parsed.blocks[0].text == "一、填报说明"
    assert "资产合计" in parsed.blocks[1].text
    assert parsed.blocks[2].text == "二、留痕要求"


def test_docx_loader_preserves_omml_formula_context_and_metadata(tmp_path) -> None:
    from docx import Document as DocxDocument

    path = tmp_path / "formula_rules.docx"
    document = DocxDocument()
    document.add_paragraph("Capital requirement is calculated as FORMULA_PLACEHOLDER where A and B are exposures.")
    document.save(path)
    _replace_docx_text_with_omml(
        path,
        "FORMULA_PLACEHOLDER",
        """
        <m:oMath>
          <m:r><m:t>K</m:t></m:r>
          <m:r><m:t>=</m:t></m:r>
          <m:f>
            <m:num>
              <m:r><m:t>A</m:t></m:r>
              <m:r><m:t>+</m:t></m:r>
              <m:r><m:t>B</m:t></m:r>
            </m:num>
            <m:den><m:r><m:t>C</m:t></m:r></m:den>
          </m:f>
        </m:oMath>
        """,
    )

    parsed = parse_document(path)
    block = parsed.blocks[0]

    assert "[公式] K=(A+B)/(C)" in block.text
    assert "Capital requirement" in block.text
    assert "where A and B are exposures" in block.text
    assert block.metadata["contains_formula"] is True
    assert block.metadata["formula_count"] == 1
    assert block.metadata["formulas"][0]["source_type"] == "omml"
    assert block.metadata["formulas"][0]["text"] == "K=(A+B)/(C)"


def test_docx_loader_preserves_omml_formula_inside_table_cell(tmp_path) -> None:
    from docx import Document as DocxDocument

    path = tmp_path / "table_formula_rules.docx"
    document = DocxDocument()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Formula"
    table.cell(1, 0).text = "LGD"
    table.cell(1, 1).text = "FORMULA_PLACEHOLDER"
    document.save(path)
    _replace_docx_text_with_omml(
        path,
        "FORMULA_PLACEHOLDER",
        """
        <m:oMath>
          <m:sSup>
            <m:e><m:r><m:t>LGD</m:t></m:r></m:e>
            <m:sup><m:r><m:t>*</m:t></m:r></m:sup>
          </m:sSup>
          <m:r><m:t>=</m:t></m:r>
          <m:r><m:t>LGD</m:t></m:r>
          <m:sSub><m:e><m:r><m:t>s</m:t></m:r></m:e><m:sub><m:r><m:t>1</m:t></m:r></m:sub></m:sSub>
        </m:oMath>
        """,
    )

    parsed = parse_document(path)
    block = parsed.blocks[0]

    assert block.block_type == "table"
    assert "[公式] LGD^*=LGDs_1" in block.text
    assert block.metadata["contains_formula"] is True
    assert block.metadata["formula_source_type"] == "omml"


def test_pdf_text_formula_is_annotated_without_ocr(tmp_path) -> None:
    import pymupdf

    path = tmp_path / "formula.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Capital formula")
    page.insert_text((72, 96), "K = A + B")
    page.insert_text((72, 120), "A and B are risk components.")
    document.save(path)
    document.close()

    parsed = parse_document(path)
    formula_blocks = [block for block in parsed.blocks if block.metadata.get("contains_formula")]

    assert formula_blocks
    assert any("K=A+B" in item["text"] for block in formula_blocks for item in block.metadata["formulas"])


def test_pdf_non_text_formula_is_not_misreported(tmp_path) -> None:
    import pymupdf

    path = tmp_path / "image_formula.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "The formula is shown in the image below.")
    page.draw_rect((72, 96, 180, 130))
    document.save(path)
    document.close()

    parsed = parse_document(path)

    assert not any(block.metadata.get("contains_formula") for block in parsed.blocks)


def _replace_docx_text_with_omml(path, placeholder: str, omml_xml: str) -> None:
    original_bytes = path.read_bytes()
    tmp_path = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "word/document.xml":
                xml = data.decode("utf-8")
                assert placeholder in xml
                xml = xml.replace(
                    placeholder,
                    f"</w:t></w:r>{omml_xml.strip()}<w:r><w:t>",
                    1,
                )
                data = xml.encode("utf-8")
            target.writestr(item, data)
    path.write_bytes(tmp_path.read_bytes())
    tmp_path.unlink()
    assert path.read_bytes() != original_bytes


def test_office_conversion_rejects_oversized_output(tmp_path, monkeypatch) -> None:
    from backend.app.services import office_conversion

    output_dir = tmp_path / "regumate-office-test"
    output_dir.mkdir()
    converted = output_dir / "oversized.docx"
    converted.write_bytes(b"123456789")
    monkeypatch.setattr(office_conversion, "OFFICE_CONVERSION_MAX_BYTES", 8)

    with pytest.raises(office_conversion.OfficeConversionError, match="exceeds 8 bytes"):
        office_conversion._validate_conversion_output(converted, output_dir)

    assert not output_dir.exists()


def test_legacy_doc_loader_uses_libreoffice_conversion(tmp_path, monkeypatch) -> None:
    from docx import Document as DocxDocument

    converted_path = tmp_path / "converted.docx"
    document = DocxDocument()
    document.add_paragraph("监管制度正文")
    document.save(converted_path)
    legacy_path = tmp_path / "legacy.doc"
    legacy_path.write_bytes(b"legacy-binary-placeholder")

    from backend.app.services.office_conversion import OfficeConversionResult

    monkeypatch.setattr(
        "backend.app.services.office_conversion.convert_with_libreoffice_detailed",
        lambda file_path, target_extension: OfficeConversionResult(
            path=converted_path, backend="libreoffice", elapsed_ms=12.5
        ),
    )

    parsed = parse_document(legacy_path)

    assert parsed.metadata["loader_name"] == "libreoffice-doc"
    assert parsed.metadata["converted_from"] == "doc"
    assert parsed.metadata["parser_backend"] == "libreoffice"
    assert parsed.metadata["degraded"] is False
    assert "监管制度正文" in parsed.text


def test_legacy_doc_falls_back_to_antiword(tmp_path, monkeypatch) -> None:
    import subprocess

    legacy_path = tmp_path / "legacy.doc"
    legacy_path.write_bytes(b"legacy-binary-placeholder")

    monkeypatch.setattr(
        "backend.app.services.office_conversion.convert_with_libreoffice_detailed",
        lambda file_path, target_extension: (_ for _ in ()).throw(
            __import__(
                "backend.app.services.office_conversion", fromlist=["OfficeConversionError"]
            ).OfficeConversionError("forced failure")
        ),
    )
    monkeypatch.setattr("shutil.which", lambda name: "antiword" if name == "antiword" else None)
    monkeypatch.setattr(
        subprocess,
        "run",
        lambda *args, **kwargs: subprocess.CompletedProcess(
            args=args[0], returncode=0, stdout="监管制度降级正文".encode("utf-8"), stderr=b""
        ),
    )

    parsed = parse_document(legacy_path)

    assert parsed.metadata["loader_name"] == "antiword-doc"
    assert parsed.metadata["parser_backend"] == "antiword"
    assert parsed.metadata["degraded"] is True
    assert "监管制度降级正文" in parsed.text


def test_xlsx_parser_builds_spreadsheet_semantic_blocks(tmp_path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "2024年一季度全国各地区原保险保费收入情况表.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "2024年一季度"
    sheet.merge_cells("A1:D1")
    sheet["A1"] = "2024年一季度全国各地区原保险保费收入情况表"
    sheet["A2"] = "单位：亿元"
    sheet.append(["地区", "指标", "本年累计", "本年累计"])
    sheet.append(["", "", "原保险保费收入", "同比增长"])
    sheet.append(["全国合计", "原保险保费收入", 123.45, "6.5%"])
    sheet.append(["北京", "原保险保费收入", 11.2, "5.0%"])
    workbook.save(path)

    parsed = parse_document(path)
    chunks = build_chunks_from_parsed(document_id="DOC-TEST-0001", parsed=parsed, source_file=path.name)
    row_block = next(block for block in parsed.blocks if block.metadata.get("table_chunk_role") == "row")
    cell = next(item for item in row_block.metadata["cells"] if item["coordinate"] == "C5")

    assert parsed.metadata["loader_name"] == "spreadsheet-xlsx"
    assert parsed.metadata["source_format"] == "xlsx"
    assert parsed.metadata["inferred_year"] == 2024
    assert row_block.metadata["spreadsheet_table"] is True
    assert row_block.metadata["sheet_name"] == "2024年一季度"
    assert row_block.metadata["table_title"] == "2024年一季度全国各地区原保险保费收入情况表"
    assert row_block.metadata["unit"] == "亿元"
    assert row_block.metadata["period"]["quarter"] == 1
    assert row_block.metadata["row_label"] == "全国合计 / 原保险保费收入"
    assert cell["column_label"] == "本年累计 / 原保险保费收入"
    assert cell["normalized_value"] == 123.45
    assert cell["unit"] == "亿元"
    assert chunks[1].metadata["cells"][2]["coordinate"] == "C5"
    assert "工作表：2024年一季度" in chunks[1].embedding_text
    assert "单位：亿元" in chunks[1].embedding_text


def test_xls_loader_uses_libreoffice_conversion_path(tmp_path, monkeypatch) -> None:
    from openpyxl import Workbook

    converted_path = tmp_path / "converted.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "报表"
    sheet.append(["统计表"])
    sheet.append(["单位：亿元"])
    sheet.append(["地区", "指标", "数值"])
    sheet.append(["全国合计", "资产合计", 100])
    workbook.save(converted_path)
    xls_path = tmp_path / "监管报表.xls"
    xls_path.write_bytes(b"legacy-xls-placeholder")

    monkeypatch.setattr(
        "backend.app.services.spreadsheet_parser.convert_with_libreoffice",
        lambda file_path, target_extension: converted_path,
    )

    parsed = parse_document(xls_path)

    assert parsed.metadata["loader_name"] == "spreadsheet-xls-libreoffice"
    assert parsed.metadata["source_format"] == "xls"
    assert parsed.metadata["converted_from"] == "xls"
    assert any(block.metadata.get("source_format") == "xls" for block in parsed.blocks)


def test_xls_loader_falls_back_when_converted_workbook_is_invalid(tmp_path, monkeypatch) -> None:
    from backend.app.services.document_types import LoaderResult, ParsedBlock
    from backend.app.services.spreadsheet_parser import load_xls_workbook

    converted_path = tmp_path / "broken.xlsx"
    converted_path.write_bytes(b"not-a-valid-ooxml-package")
    xls_path = tmp_path / "legacy.xls"
    xls_path.write_bytes(b"legacy-biff-placeholder")
    fallback = LoaderResult(
        blocks=[ParsedBlock(text="xlrd fallback", block_type="table", order_index=1)],
        loader_name="spreadsheet-xls-xlrd",
    )

    monkeypatch.setattr(
        "backend.app.services.spreadsheet_parser.convert_with_libreoffice",
        lambda file_path, target_extension: converted_path,
    )
    monkeypatch.setattr(
        "backend.app.services.spreadsheet_parser._load_xls_with_xlrd",
        lambda file_path: fallback,
    )

    result = load_xls_workbook(xls_path)

    assert result.loader_name == "spreadsheet-xls-xlrd"
    assert result.blocks[0].text == "xlrd fallback"
    assert tmp_path.exists()


def test_pdf_defaults_to_pymupdf4llm_loader(tmp_path) -> None:
    import pymupdf

    path = tmp_path / "rules.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Total assets equal the sum of asset items.")
    document.save(path)
    document.close()

    parsed = parse_document(path)

    assert parsed.metadata["loader_name"] in {"pymupdf4llm", "pypdf"}
    assert parsed.blocks[0].page_number == 1
    assert "Total assets" in parsed.text


def test_loader_evaluation_compares_pdf_loaders(tmp_path) -> None:
    import pymupdf

    path = tmp_path / "rules.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "监管填报说明")
    document.save(path)
    document.close()

    evaluations = evaluate_document_loaders(path)

    names = [evaluation.loader_name for evaluation in evaluations]
    assert names == ["pymupdf4llm", "docling", "unstructured", "pypdf"]
    assert any(evaluation.ok and evaluation.block_count > 0 for evaluation in evaluations)


def test_formula_metadata_is_scoped_to_its_own_chunk() -> None:
    parsed = ParsedDocument(
        text="第一段\nd=SD*Notional\n第三段",
        blocks=[
            ParsedBlock(text="第一章 计量规则", block_type="heading", order_index=1, level=1),
            ParsedBlock(text="第一段普通说明。", block_type="paragraph", order_index=2, section_title="第一章 计量规则"),
            ParsedBlock(
                text="[公式] d=SD*Notional",
                block_type="paragraph",
                order_index=3,
                section_title="第一章 计量规则",
                metadata={
                    "contains_formula": True,
                    "formula_count": 1,
                    "formulas": [{"text": "d=SD*Notional", "source_type": "omml", "order_index": 1}],
                    "formula_source_type": "omml",
                },
            ),
            ParsedBlock(text="第三段普通说明。", block_type="paragraph", order_index=4, section_title="第一章 计量规则"),
        ],
        metadata={"source_format": "docx"},
    )

    chunks = build_chunks_from_parsed("DOC-FORMULA-SCOPE", parsed)
    formula_chunks = [chunk for chunk in chunks if (chunk.metadata or {}).get("contains_formula")]

    assert len(formula_chunks) == 1
    assert formula_chunks[0].text == "[公式] d=SD*Notional"
    assert formula_chunks[0].metadata["formulas"][0]["text"] == "d=SD*Notional"
    assert all(
        not (chunk.metadata or {}).get("contains_formula")
        for chunk in chunks
        if chunk is not formula_chunks[0]
    )
