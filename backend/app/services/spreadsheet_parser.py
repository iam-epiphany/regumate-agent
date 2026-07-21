from __future__ import annotations

from dataclasses import dataclass, replace
from datetime import date, datetime
from pathlib import Path
import re
from typing import Any

from backend.app.services.document_parser import DocumentParseError
from backend.app.core.config import MAX_SPREADSHEET_LOGICAL_CELLS
from backend.app.services.document_types import LoaderResult, ParsedBlock
from backend.app.services.office_conversion import (
    OfficeConversionError,
    cleanup_conversion_output,
    convert_with_libreoffice,
)


@dataclass(frozen=True)
class SpreadsheetCell:
    row: int
    column: int
    coordinate: str
    value: Any
    display_value: str
    normalized_value: Any
    data_type: str
    is_formula: bool = False
    formula: str | None = None
    merged_from: str | None = None


def load_xlsx_workbook(file_path: Path, *, loader_name: str = "spreadsheet-xlsx") -> LoaderResult:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise DocumentParseError("缺少 openpyxl 依赖，无法解析 xlsx") from exc

    try:
        value_workbook = load_workbook(str(file_path), data_only=True, read_only=False)
        formula_workbook = load_workbook(str(file_path), data_only=False, read_only=False)
    except Exception as exc:
        raise DocumentParseError("xlsx 表格解析失败") from exc

    try:
        logical_cells = sum(
            max_row * max_col
            for sheet in value_workbook.worksheets
            for _, max_row, _, max_col in [_used_bounds(sheet)]
        )
        if logical_cells > MAX_SPREADSHEET_LOGICAL_CELLS:
            raise DocumentParseError(
                "Excel 逻辑单元格数量超过处理上限"
                f"（{logical_cells} > {MAX_SPREADSHEET_LOGICAL_CELLS}）"
            )
        blocks: list[ParsedBlock] = []
        workbook_metadata = _workbook_metadata(file_path)
        for sheet_index, value_sheet in enumerate(value_workbook.worksheets, start=1):
            formula_sheet = formula_workbook[value_sheet.title]
            blocks.extend(_sheet_blocks(file_path, value_sheet, formula_sheet, sheet_index, workbook_metadata))
        return LoaderResult(blocks=blocks, loader_name=loader_name, metadata=workbook_metadata)
    finally:
        value_workbook.close()
        formula_workbook.close()


def load_xls_workbook(file_path: Path) -> LoaderResult:
    try:
        converted = convert_with_libreoffice(file_path, "xlsx")
    except OfficeConversionError:
        return _load_xls_with_xlrd(file_path)
    try:
        result = load_xlsx_workbook(converted, loader_name="spreadsheet-xls-libreoffice")
    except DocumentParseError:
        # Some legacy BIFF/WPS workbooks make LibreOffice report a successful
        # conversion while the produced OOXML package is incomplete. In that
        # case, fall back to xlrd instead of treating the source as unreadable.
        return _load_xls_with_xlrd(file_path)
    finally:
        cleanup_conversion_output(converted)
    original_metadata = {
        **_workbook_metadata(file_path),
        "converted_from": "xls",
        "conversion_loader": "libreoffice",
    }
    return LoaderResult(
        blocks=[
            replace(
                block,
                metadata={
                    **block.metadata,
                    **original_metadata,
                },
            )
            for block in result.blocks
        ],
        loader_name=result.loader_name,
        metadata={**result.metadata, **original_metadata},
    )


def _load_xls_with_xlrd(file_path: Path) -> LoaderResult:
    try:
        import xlrd
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise DocumentParseError("缺少 xlrd 依赖，且 LibreOffice 转换失败，无法解析 xls") from exc

    try:
        book = xlrd.open_workbook(str(file_path), formatting_info=False)
    except Exception as exc:
        raise DocumentParseError("xls 表格解析失败") from exc

    workbook_metadata = _workbook_metadata(file_path)
    blocks: list[ParsedBlock] = []
    for sheet_index in range(book.nsheets):
        sheet = book.sheet_by_index(sheet_index)
        rows: list[list[SpreadsheetCell]] = []
        for row_index in range(sheet.nrows):
            row: list[SpreadsheetCell] = []
            for col_index in range(sheet.ncols):
                cell = sheet.cell(row_index, col_index)
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        value = xlrd.xldate_as_datetime(value, book.datemode)
                    except Exception:
                        pass
                coordinate = f"{get_column_letter(col_index + 1)}{row_index + 1}"
                display = _display_value(value)
                row.append(
                    SpreadsheetCell(
                        row=row_index + 1,
                        column=col_index + 1,
                        coordinate=coordinate,
                        value=value,
                        display_value=display,
                        normalized_value=_normalized_value(value),
                        data_type=str(cell.ctype),
                    )
                )
            rows.append(row)
        blocks.extend(
            _blocks_from_cell_rows(
                file_path=file_path,
                sheet_name=sheet.name,
                sheet_index=sheet_index + 1,
                sheet_state="visible",
                rows=rows,
                workbook_metadata=workbook_metadata,
            )
        )
    return LoaderResult(blocks=blocks, loader_name="spreadsheet-xls-xlrd", metadata=workbook_metadata)


def _sheet_blocks(
    file_path: Path,
    value_sheet: Any,
    formula_sheet: Any,
    sheet_index: int,
    workbook_metadata: dict[str, Any],
) -> list[ParsedBlock]:
    from openpyxl.utils import get_column_letter

    merged_sources = _merged_sources(value_sheet)
    rows: list[list[SpreadsheetCell]] = []
    _, max_row, _, max_col = _used_bounds(value_sheet)
    for row_index in range(1, max_row + 1):
        row: list[SpreadsheetCell] = []
        for col_index in range(1, max_col + 1):
            coordinate = f"{get_column_letter(col_index)}{row_index}"
            source_coordinate = merged_sources.get(coordinate)
            source = value_sheet[source_coordinate] if source_coordinate else value_sheet[coordinate]
            formula_source = formula_sheet[source_coordinate] if source_coordinate else formula_sheet[coordinate]
            value = source.value
            formula_value = formula_source.value
            is_formula = isinstance(formula_value, str) and formula_value.startswith("=")
            display = _display_value(value)
            row.append(
                SpreadsheetCell(
                    row=row_index,
                    column=col_index,
                    coordinate=coordinate,
                    value=value,
                    display_value=display,
                    normalized_value=_normalized_value(value),
                    data_type=str(source.data_type),
                    is_formula=is_formula,
                    formula=formula_value if is_formula else None,
                    merged_from=source_coordinate if source_coordinate != coordinate else None,
                )
            )
        rows.append(row)
    return _blocks_from_cell_rows(
        file_path=file_path,
        sheet_name=value_sheet.title,
        sheet_index=sheet_index,
        sheet_state=getattr(value_sheet, "sheet_state", "visible"),
        rows=rows,
        workbook_metadata=workbook_metadata,
    )


def _used_bounds(sheet: Any) -> tuple[int, int, int, int]:
    """Use materialized cells and merged ranges, not an untrusted worksheet dimension."""

    coordinates = list(getattr(sheet, "_cells", {}).keys())
    min_row = min((row for row, _ in coordinates), default=1)
    max_row = max((row for row, _ in coordinates), default=0)
    min_col = min((column for _, column in coordinates), default=1)
    max_col = max((column for _, column in coordinates), default=0)
    for merged_range in getattr(sheet, "merged_cells", []).ranges:
        min_row = min(min_row, merged_range.min_row)
        max_row = max(max_row, merged_range.max_row)
        min_col = min(min_col, merged_range.min_col)
        max_col = max(max_col, merged_range.max_col)
    return min_row, max_row, min_col, max_col


def _blocks_from_cell_rows(
    *,
    file_path: Path,
    sheet_name: str,
    sheet_index: int,
    sheet_state: str,
    rows: list[list[SpreadsheetCell]],
    workbook_metadata: dict[str, Any],
) -> list[ParsedBlock]:
    rows = _trim_empty_edges(rows)
    if not rows:
        return []

    table_title = _infer_table_title(file_path, sheet_name, rows)
    unit = _infer_unit(rows)
    period = _infer_period(" ".join([file_path.name, sheet_name, table_title]))
    data_start_index = _infer_data_start(rows)
    header_indexes = _infer_header_indexes(rows, data_start_index)
    headers = _column_headers(rows, header_indexes)
    row_header_columns = _infer_row_header_columns(rows, data_start_index, headers)
    data_rows = rows[data_start_index:] if data_start_index is not None else []
    table_id = f"{file_path.stem}-SHEET-{sheet_index:02d}"

    common_metadata = {
        **workbook_metadata,
        "spreadsheet_table": True,
        "sheet_name": sheet_name,
        "sheet_index": sheet_index,
        "sheet_state": sheet_state,
        "table_id": table_id,
        "table_title": table_title,
        "unit": unit,
        "period": period,
        "table_headers": headers,
        "row_header_columns": row_header_columns,
    }
    summary_text = _summary_text(table_title, sheet_name, unit, period, headers, len(data_rows))
    blocks = [
        ParsedBlock(
            text=summary_text,
            block_type="table",
            order_index=1,
            section_title=table_title,
            metadata={
                **common_metadata,
                "table_chunk_role": "summary",
                "row_index": None,
                "raw_table_preview": _preview_rows(rows),
            },
        )
    ]

    row_context: dict[int, str] = {}
    for row_number, row in enumerate(data_rows, start=(data_start_index or 0) + 1):
        row_label = _row_label(row, row_header_columns, row_context)
        cell_records = _cell_records(row, headers, row_label, unit)
        if not cell_records:
            continue
        row_cells: dict[str, str] = {}
        for record in cell_records:
            label = str(record.get("column_label") or "").strip()
            value = str(record.get("value") or "").strip()
            if not label or not value:
                continue
            # Malformed or merged headers can legitimately produce duplicate
            # labels.  Keep every value traceable instead of silently letting
            # the last column overwrite the earlier cells.
            key = label if label not in row_cells else f"{label} [{record['coordinate']}]"
            row_cells[key] = value
        row_text = _row_text(table_title, sheet_name, row_label, row_cells, unit)
        blocks.append(
            ParsedBlock(
                text=row_text,
                block_type="table",
                order_index=len(blocks) + 1,
                section_title=table_title,
                metadata={
                    **common_metadata,
                    "table_chunk_role": "row",
                    "row_index": row_number,
                    "row_label": row_label,
                    "row_cells": row_cells,
                    "cells": cell_records,
                    "raw_table_preview": row_text[:500],
                },
            )
        )
    return blocks


def _workbook_metadata(file_path: Path) -> dict[str, Any]:
    text = file_path.stem
    return {
        "source_title": _source_title(file_path),
        "source_format": file_path.suffix.lower().lstrip("."),
        "inferred_year": _first_int(re.findall(r"(20\d{2})年?", text)),
        "inferred_month": _first_int(re.findall(r"(?:20\d{2}年)?(\d{1,2})月", text)),
        "inferred_quarter": _infer_quarter(text),
        "report_type": _infer_report_type(text),
        "business_domain": _infer_business_domain(text),
    }


def _merged_sources(sheet: Any) -> dict[str, str]:
    sources: dict[str, str] = {}
    for merged_range in getattr(sheet, "merged_cells", []).ranges:
        anchor = merged_range.start_cell.coordinate
        for row in sheet.iter_rows(
            min_row=merged_range.min_row,
            max_row=merged_range.max_row,
            min_col=merged_range.min_col,
            max_col=merged_range.max_col,
        ):
            for cell in row:
                sources[cell.coordinate] = anchor
    return sources


def _trim_empty_edges(rows: list[list[SpreadsheetCell]]) -> list[list[SpreadsheetCell]]:
    non_empty = [index for index, row in enumerate(rows) if any(cell.display_value for cell in row)]
    if not non_empty:
        return []
    start, end = min(non_empty), max(non_empty)
    col_indexes = [
        cell.column
        for row in rows[start : end + 1]
        for cell in row
        if cell.display_value
    ]
    min_col, max_col = min(col_indexes), max(col_indexes)
    return [[cell for cell in row if min_col <= cell.column <= max_col] for row in rows[start : end + 1]]


def _infer_table_title(file_path: Path, sheet_name: str, rows: list[list[SpreadsheetCell]]) -> str:
    for row in rows[:5]:
        values = [cell.display_value for cell in row if cell.display_value]
        if len(values) == 1 and not _looks_like_unit(values[0]):
            return values[0]
        for value in values:
            if len(value) <= 80 and any(term in value for term in ["表", "情况", "统计", "规则", "名单"]):
                return value
    return _source_title(file_path) or sheet_name


def _infer_unit(rows: list[list[SpreadsheetCell]]) -> str | None:
    for row in rows[:10]:
        for cell in row:
            text = cell.display_value
            match = re.search(r"单位\s*[:：]\s*([^；;，,。]+)", text)
            if match:
                return match.group(1).strip()
    return None


def _infer_period(text: str) -> dict[str, int | str | None]:
    return {
        "year": _first_int(re.findall(r"(20\d{2})年?", text)),
        "month": _first_int(re.findall(r"(?:20\d{2}年)?(\d{1,2})月", text)),
        "quarter": _infer_quarter(text),
        "raw": _period_raw(text),
    }


def _infer_data_start(rows: list[list[SpreadsheetCell]]) -> int | None:
    for index, row in enumerate(rows):
        if index == 0:
            continue
        non_empty = [cell for cell in row if cell.display_value]
        if len(non_empty) < 2:
            continue
        numeric_count = sum(1 for cell in non_empty if isinstance(cell.normalized_value, (int, float)))
        text_count = len(non_empty) - numeric_count
        prior_header_cells = max(
            (
                sum(
                    1
                    for cell in rows[prior_index]
                    if cell.display_value and not isinstance(cell.normalized_value, (int, float))
                )
                for prior_index in range(max(0, index - 3), index)
            ),
            default=0,
        )
        if numeric_count >= 1 and text_count >= 1 and prior_header_cells >= 2:
            return index
        # Text-only multi-row headers are not data.  The previous fallback
        # returned the second header row as the first data row, which collapsed
        # every column label to the workbook title on several contest sheets.
    return None


def _infer_header_indexes(rows: list[list[SpreadsheetCell]], data_start_index: int | None) -> list[int]:
    if data_start_index is None:
        return []
    candidates: list[int] = []
    for index in range(max(0, data_start_index - 3), data_start_index):
        row = rows[index]
        values = [cell.display_value for cell in row if cell.display_value]
        if len(values) >= 2 and not any(_looks_like_unit(value) for value in values):
            candidates.append(index)
    return candidates or [max(data_start_index - 1, 0)]


def _column_headers(rows: list[list[SpreadsheetCell]], header_indexes: list[int]) -> list[str]:
    if not rows:
        return []
    width = max(len(row) for row in rows)
    headers: list[str] = []
    for col_index in range(width):
        parts: list[str] = []
        for row_index in header_indexes:
            if col_index >= len(rows[row_index]):
                continue
            value = rows[row_index][col_index].display_value
            if value and value not in parts and not _looks_like_unit(value):
                parts.append(value)
        headers.append(" / ".join(parts) if parts else f"列{col_index + 1}")
    return headers


def _infer_row_header_columns(
    rows: list[list[SpreadsheetCell]],
    data_start_index: int | None,
    headers: list[str],
) -> list[int]:
    if data_start_index is None:
        return [1]
    sample = rows[data_start_index : min(len(rows), data_start_index + 8)]
    columns: list[int] = []
    for col_index in range(min(3, len(headers))):
        text_values = 0
        numeric_values = 0
        for row in sample:
            if col_index >= len(row) or not row[col_index].display_value:
                continue
            if isinstance(row[col_index].normalized_value, (int, float)):
                numeric_values += 1
            else:
                text_values += 1
        if text_values > 0 and numeric_values == 0:
            columns.append(col_index + 1)
    return columns or [1]


def _row_label(row: list[SpreadsheetCell], row_header_columns: list[int], row_context: dict[int, str]) -> str:
    labels: list[str] = []
    for column in row_header_columns:
        cell = next((item for item in row if item.column == column), None)
        value = cell.display_value if cell else ""
        if value:
            row_context[column] = value
        elif column in row_context:
            value = row_context[column]
        if value and value not in labels:
            labels.append(value)
    if labels:
        return " / ".join(labels)
    first = next((cell.display_value for cell in row if cell.display_value), "")
    return first or f"第{row[0].row}行"


def _cell_records(
    row: list[SpreadsheetCell],
    headers: list[str],
    row_label: str,
    unit: str | None,
) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for index, cell in enumerate(row):
        if not cell.display_value:
            continue
        column_label = headers[index] if index < len(headers) else f"列{index + 1}"
        records.append(
            {
                "coordinate": cell.coordinate,
                "row": cell.row,
                "column": cell.column,
                "row_label": row_label,
                "column_label": column_label,
                "column_path": column_label.split(" / "),
                "value": cell.display_value,
                "normalized_value": cell.normalized_value,
                "unit": unit,
                "data_type": cell.data_type,
                "is_formula": cell.is_formula,
                "formula": cell.formula,
                "merged_from": cell.merged_from,
            }
        )
    return records


def _summary_text(
    table_title: str,
    sheet_name: str,
    unit: str | None,
    period: dict[str, Any],
    headers: list[str],
    row_count: int,
) -> str:
    header_text = "、".join(header for header in headers if header) or "未识别表头"
    period_text = period.get("raw") or "未识别期间"
    unit_text = unit or "未识别单位"
    return (
        f"表格摘要：{table_title}。工作表：{sheet_name}。期间：{period_text}。"
        f"单位：{unit_text}。表头：{header_text}。共{row_count}行数据。"
    )


def _row_text(
    table_title: str,
    sheet_name: str,
    row_label: str,
    row_cells: dict[str, str],
    unit: str | None,
) -> str:
    cells = "；".join(f"{key}=“{value}”" for key, value in row_cells.items())
    unit_text = f"单位：{unit}。" if unit else ""
    return f"表格行证据：在《{table_title}》工作表“{sheet_name}”中，行标签为“{row_label}”。{unit_text}{cells}。"


def _preview_rows(rows: list[list[SpreadsheetCell]], limit: int = 6) -> str:
    lines = []
    for row in rows[:limit]:
        lines.append(" | ".join(cell.display_value for cell in row if cell.display_value))
    return "\n".join(line for line in lines if line)


def _display_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S").rstrip(" 00:00:00")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.10g}"
    return str(value).strip()


def _normalized_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    text = str(value).strip().replace(",", "")
    if not text or text in {"-", "--", "—"}:
        return None
    percent = text.endswith("%")
    numeric_text = text[:-1] if percent else text
    try:
        number = float(numeric_text)
    except ValueError:
        return str(value).strip()
    return number / 100 if percent else number


def _source_title(file_path: Path) -> str:
    stem = file_path.stem
    if "_" in stem:
        return stem.split("_", maxsplit=1)[-1]
    return stem


def _infer_report_type(text: str) -> str | None:
    candidates = [
        "全国各地区原保险保费收入情况表",
        "人身险公司经营情况表",
        "财产保险公司经营情况表",
        "财产险公司经营情况表",
        "保险业经营情况表",
        "保险业资金运用情况表",
        "商业银行主要监管指标情况表",
        "商业银行主要指标分机构类情况表",
        "银行业总资产、总负债",
    ]
    return next((item for item in candidates if item in text), None)


def _infer_business_domain(text: str) -> str | None:
    if "保险" in text or "保费" in text:
        return "insurance"
    if "银行" in text or "商业银行" in text:
        return "banking"
    return None


def _infer_quarter(text: str) -> int | None:
    mapping = {"一": 1, "二": 2, "三": 3, "四": 4}
    match = re.search(r"([一二三四1-4])季度|([1-4])季", text)
    if not match:
        return None
    value = match.group(1) or match.group(2)
    return mapping.get(value, int(value) if value.isdigit() else None)


def _period_raw(text: str) -> str | None:
    patterns = [
        r"20\d{2}年\d{1,2}月",
        r"20\d{2}年[一二三四1-4]季度",
        r"20\d{2}年[一二三四]季度",
        r"20\d{2}年",
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return match.group(0)
    return None


def _first_int(values: list[str]) -> int | None:
    for value in values:
        try:
            return int(value)
        except (TypeError, ValueError):
            continue
    return None


def _looks_like_unit(text: str) -> bool:
    return bool(re.search(r"单位\s*[:：]", text))
